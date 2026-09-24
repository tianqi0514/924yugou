"""Exercise the project-original evidence gate in report_platform_test only.

The uploaded DOCX is created in memory from synthetic QA assertions. A passing
gate proves that source-position binding and computed-rule provenance work; it
does not independently establish the assertions as commercial facts.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import os
import sys
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile


HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
TEST_URL = "postgresql+psycopg://report:local_development_only@127.0.0.1:55432/report_platform_test"


def body(response, expected: int) -> dict:
    if response.status_code != expected:
        raise RuntimeError(f"{response.request.method} {response.request.url.path}: "
                           f"expected {expected}, got {response.status_code}: {response.text[:500]}")
    return response.json()


def synthetic_docx() -> bytes:
    from docx import Document

    document = Document()
    document.add_heading("QA 合成输入说明", level=1)
    document.add_paragraph("仅供软件流程验证；以下数值不是独立核实的商业证据。")
    document.add_paragraph("首年客户需求 300000 套。")
    document.add_paragraph("规划合格能力 254016 套。")
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--execute-test-db", action="store_true", required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    existing = os.environ.get("DATABASE_URL")
    if existing and not existing.endswith("/report_platform_test"):
        raise RuntimeError("DATABASE_URL does not target report_platform_test; refusing to run")
    os.environ["DATABASE_URL"] = existing or TEST_URL
    sys.path.insert(0, str(ROOT / "backend"))

    from fastapi.testclient import TestClient
    from app.db import engine
    from app.main import app
    import app.main as main_module

    if engine.url.database != "report_platform_test":
        raise RuntimeError("Refusing to write outside report_platform_test")
    # Keep the uploaded QA file in this evaluation directory, away from app data.
    main_module.STORAGE = HERE / "_test_storage"

    with TestClient(app) as client:
        project = body(client.post("/api/projects", json={"name": "QA-合成原件绑定闸门验证"}), 201)
        pid = project["id"]
        builtin = next(item for item in body(client.get("/api/projects"), 200) if item["is_builtin"])
        corpus = body(client.get(f"/api/projects/{builtin['id']}/corpus/summary"), 200)
        body(client.put(f"/api/projects/{pid}/writing/reference", json={
            "corpus_id": corpus["corpus_id"], "corpus_version": corpus["version"],
            "target_report_type": "feasibility", "accepted": True,
        }), 200)
        body(client.post(f"/api/projects/{pid}/writing/setup"), 200)
        body(client.put(f"/api/projects/{pid}/writing/bindings", json={"bindings": {
            "N017": "first_year_demand", "N034": "qualified_capacity", "N035": "planned_sales",
            "supplier_name": "supplier_name", "N080": "equipment_unit_price",
        }}), 200)
        for fact_key, value in (("first_year_demand", "300000"), ("qualified_capacity", "254016")):
            payload = {"fact_key": fact_key, "value": value, "source": "QA合成原件；非商业真实性认证",
                       "reason": "验证可定位原文与规则计算"}
            preview = body(client.post(f"/api/projects/{pid}/changes/preview", json=payload), 200)
            body(client.post(f"/api/projects/{pid}/changes/commit", json={
                **payload, "base_version": preview["base_version"], "preview_token": preview["preview_token"],
            }), 200)
        facts = body(client.get(f"/api/projects/{pid}/facts"), 200)["facts"]
        values = {fact["key"]: fact["value"] for fact in facts}
        if values.get("planned_sales") != "254016":
            raise RuntimeError("Project min rule did not compute 254016")

        document_bytes = synthetic_docx()
        uploaded = body(client.post(f"/api/projects/{pid}/documents", files={
            "file": ("QA-合成输入.docx", document_bytes,
                     "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
        }), 201)
        did = uploaded["id"]
        detail = body(client.get(f"/api/projects/{pid}/documents/{did}"), 200)
        segments = detail["segments"]
        selected = {
            "first_year_demand": next(item["ref"] for item in segments if "300000" in item["text"]),
            "qualified_capacity": next(item["ref"] for item in segments if "254016" in item["text"]),
        }
        derived_direct_bind = client.post(f"/api/projects/{pid}/facts/planned_sales/evidence/bind",
                                          json={"document_id": did, "source_refs": [selected["qualified_capacity"]]})
        if derived_direct_bind.status_code != 409:
            raise RuntimeError("Computed fact unexpectedly accepted a direct document binding")
        bindings = {}
        for fact_key, ref in selected.items():
            bindings[fact_key] = body(client.post(f"/api/projects/{pid}/facts/{fact_key}/evidence/bind",
                                                  json={"document_id": did, "source_refs": [ref]}), 200)
        report = body(client.post(f"/api/projects/{pid}/reports", json={"title": "S4 合成证据闸门验证"}), 201)
        rid = report["id"]
        package = body(client.get(f"/api/projects/{pid}/writing/packages/S4"), 200)
        approved = [item["item_id"] for item in package["items"] if item["decision"] == "selectable"]
        request = {"section_id": "S4", "approved_item_ids": approved, "mode": "guided"}
        preview = body(client.post(f"/api/projects/{pid}/reports/{rid}/writing/preview", json=request), 200)
        body(client.post(f"/api/projects/{pid}/reports/{rid}/writing/commit", json={**request, **preview}), 200)
        persisted = body(client.get(f"/api/projects/{pid}/reports/{rid}"), 200)
        reviewed = client.post(f"/api/projects/{pid}/reports/{rid}/review")
        reviewed_body = body(reviewed, 200) if reviewed.status_code == 200 else {"detail": reviewed.text[:500]}
        exported = client.get(f"/api/projects/{pid}/reports/{rid}/export")
        audit = None
        docx_text = pdf_text = ""
        if exported.status_code == 200:
            from docx import Document
            import pymupdf

            with ZipFile(BytesIO(exported.content)) as archive:
                audit = json.loads(archive.read("audit.json"))
                docx_text = "\n".join(p.text for p in Document(BytesIO(archive.read("report.docx"))).paragraphs)
                pdf = pymupdf.open(stream=archive.read("report.pdf"), filetype="pdf")
                pdf_text = "\n".join(page.get_text() for page in pdf)
                pdf.close()
        result = {
            "schema_version": "writing-support-evidence-gate/1.0",
            "test_database": engine.url.database,
            "synthetic_qa_only": True,
            "project_id": pid, "report_id": rid, "document_id": did,
            "document_sha256": hashlib.sha256(document_bytes).hexdigest(),
            "source_segments": [{"ref": part["ref"], "text": part["text"]} for part in segments],
            "facts": [{"key": fact["key"], "value": fact["value"], "revision": fact["revision"],
                       "status": fact["status"]} for fact in facts],
            "computed_direct_bind_status": derived_direct_bind.status_code,
            "fact_bindings": bindings,
            "persisted_report_version": persisted["version"],
            "review_status": reviewed.status_code, "review_body": reviewed_body,
            "export_status": exported.status_code,
            "export_error": None if exported.status_code == 200 else exported.text[:500],
            "audit": audit,
            "docx_text": docx_text, "pdf_text": pdf_text,
        }
    args.output.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"Synthetic source gate: review={result['review_status']}, export={result['export_status']}, "
          f"audit={audit.get('delivery_status') if audit else None}")


if __name__ == "__main__":
    main()
