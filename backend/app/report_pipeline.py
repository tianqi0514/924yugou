"""Small, inspectable report AST, impact, gate and export helpers."""

from __future__ import annotations

import hashlib
import io
import json
import re
from difflib import SequenceMatcher
from decimal import Decimal
from html import escape
from urllib.parse import urlparse
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Cm, Pt
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from .model_settings import chat_json


ALLOWED_BLOCKS = {"p", "h1", "h2", "h3", "blockquote", "table"}
TEXT_MARKS = {"bold", "italic", "underline", "strikethrough"}
EXPORT_RENDER_VERSION = "analysis-basis-v1"


def _valid_url(url: str) -> bool:
    try:
        parts = urlparse(url)
        return parts.scheme in {"https", "http"} and bool(parts.netloc) and len(url) <= 2_000
    except ValueError:
        return False


def _validate_inline(children: object) -> set[str]:
    if not isinstance(children, list) or not children or len(children) > 2_000:
        raise ValueError("段落文字结构不正确")
    cited: set[str] = set()
    for child in children:
        if not isinstance(child, dict):
            raise ValueError("段落文字结构不正确")
        kind = child.get("type")
        if kind == "a":
            if set(child) - {"type", "url", "target", "children"} or not isinstance(child.get("url"), str) or not _valid_url(child["url"]):
                raise ValueError("链接地址只支持 http 或 https")
            if child.get("target") not in (None, "_blank"):
                raise ValueError("链接打开方式不受支持")
            if not isinstance(child.get("children"), list) or any(not isinstance(leaf, dict) or "type" in leaf for leaf in child["children"]):
                raise ValueError("链接文字结构不正确")
            _validate_inline(child["children"])
        elif kind == "fact_ref":
            key, display = child.get("fact_key"), child.get("display")
            if set(child) - {"type", "fact_key", "display", "children"} or not isinstance(key, str) or not key or len(key) > 160 or not isinstance(display, str) or not display or len(display) > 300:
                raise ValueError("事实引用结构不正确")
            if child.get("children") != [{"text": ""}]:
                raise ValueError("事实引用不可直接修改")
            cited.add(key)
        elif kind is None:
            if set(child) - ({"text"} | TEXT_MARKS) or not isinstance(child.get("text"), str):
                raise ValueError("段落文字结构不正确")
            if any(mark in child and child[mark] is not True for mark in TEXT_MARKS):
                raise ValueError("文字格式结构不正确")
        else:
            raise ValueError("正文包含不支持的行内类型")
    return cited


def _validate_metadata(node: dict) -> set[str]:
    node_id = node.get("id")
    if node_id is not None and (not isinstance(node_id, str) or not node_id or len(node_id) > 100):
        raise ValueError("段落 ID 不正确")
    keys = node.get("fact_keys", [])
    if not isinstance(keys, list) or any(not isinstance(key, str) or not key or len(key) > 160 for key in keys) or len(set(keys)) != len(keys):
        raise ValueError("事实引用结构不正确")
    if "origin" in node and node["origin"] not in {"model", "manual", "guided"}:
        raise ValueError("正文来源类型不正确")
    section_id = node.get("section_id")
    if section_id is not None and (not isinstance(section_id, str) or not section_id or len(section_id) > 80):
        raise ValueError("章节 ID 不正确")
    refs = node.get("source_refs", [])
    if not isinstance(refs, list) or len(refs) > 40:
        raise ValueError("语料来源结构不正确")
    required = {"corpus_id", "corpus_version", "category_id", "artifact_id", "record_id"}
    allowed = required | {"semantic_id", "use"}
    for ref in refs:
        if (not isinstance(ref, dict) or not required.issubset(ref) or set(ref) - allowed
                or any(not isinstance(ref[key], str) or not ref[key] or len(ref[key]) > 180 for key in required)
                or ("semantic_id" in ref and (not isinstance(ref["semantic_id"], str) or len(ref["semantic_id"]) > 160))
                or ("use" in ref and (not isinstance(ref["use"], str) or len(ref["use"]) > 80))):
            raise ValueError("语料来源结构不正确")
    if refs and (section_id is None or node_id is None):
        raise ValueError("引用语料的段落必须有稳定 ID 和章节 ID")
    project_rules = node.get("project_rule_refs", [])
    if not isinstance(project_rules, list) or len(project_rules) > 20:
        raise ValueError("项目规则引用结构不正确")
    for rule in project_rules:
        if (not isinstance(rule, dict) or set(rule) != {"rule_id", "expression", "target_key", "deps", "input_fact_revisions", "target_fact_revision"}
                or any(not isinstance(rule[key], str) or not rule[key] for key in ("rule_id", "expression", "target_key"))
                or not isinstance(rule["deps"], list) or any(not isinstance(key, str) or not key for key in rule["deps"])
                or not isinstance(rule["input_fact_revisions"], list)
                or any(not isinstance(row, dict) or set(row) != {"fact_key", "revision", "value"}
                       or not isinstance(row["fact_key"], str) or not isinstance(row["revision"], int)
                       or not isinstance(row["value"], str) for row in rule["input_fact_revisions"])
                or not isinstance(rule["target_fact_revision"], int)):
            raise ValueError("项目规则引用结构不正确")
    analysis_refs = node.get("analysis_refs", [])
    if not isinstance(analysis_refs, list) or len(analysis_refs) > 30:
        raise ValueError("推演引用结构不正确")
    for ref in analysis_refs:
        if (not isinstance(ref, dict) or set(ref) != {"run_id", "result_key", "value", "unit"}
                or any(not isinstance(ref[key], str) or not ref[key] or len(ref[key]) > 100
                       for key in ("run_id", "result_key", "value"))
                or not isinstance(ref["unit"], str) or len(ref["unit"]) > 80):
            raise ValueError("推演引用结构不正确")
    if analysis_refs and (section_id is None or node_id is None):
        raise ValueError("推演引用必须属于稳定的章节和段落")
    return set(keys)


def _validate_text_block(node: dict, *, cell: bool = False) -> set[str]:
    allowed = {"type", "children", "id", "fact_keys", "origin", "section_id", "source_refs", "project_rule_refs", "analysis_refs"}
    if node.get("type") == "p":
        allowed |= {"listStyleType", "indent", "listStart"}
        if node.get("listStyleType") not in (None, "disc", "decimal"):
            raise ValueError("列表类型不受支持")
        if "indent" in node and (not isinstance(node["indent"], int) or not 1 <= node["indent"] <= 6):
            raise ValueError("列表缩进不受支持")
        if "listStart" in node and (not isinstance(node["listStart"], int) or not 1 <= node["listStart"] <= 1_000):
            raise ValueError("列表序号不正确")
    if set(node) - allowed:
        raise ValueError("正文包含不支持的段落属性")
    keys = _validate_metadata(node)
    cited = _validate_inline(node.get("children"))
    if not cell and not cited.issubset(keys):
        raise ValueError("事实引用未绑定到所在段落")
    if len(plain(node)) > 20_000:
        raise ValueError("单段文字过长")
    return keys | cited


def _validate_table(node: dict) -> set[str]:
    if set(node) - {"type", "children", "id", "fact_keys", "origin", "section_id", "source_refs", "project_rule_refs", "analysis_refs"}:
        raise ValueError("表格包含不支持的属性")
    keys = _validate_metadata(node)
    rows = node.get("children")
    if not isinstance(rows, list) or not 1 <= len(rows) <= 100:
        raise ValueError("表格行数不正确")
    width = None
    cited: set[str] = set()
    for row in rows:
        if not isinstance(row, dict) or row.get("type") != "tr" or set(row) - {"type", "children", "id"}:
            raise ValueError("表格行结构不正确")
        if "id" in row and (not isinstance(row["id"], str) or not row["id"]):
            raise ValueError("表格行 ID 不正确")
        cells = row.get("children")
        if not isinstance(cells, list) or not 1 <= len(cells) <= 8 or (width is not None and len(cells) != width):
            raise ValueError("表格列数不一致或超过限制")
        width = len(cells)
        for cell in cells:
            if not isinstance(cell, dict) or cell.get("type") not in {"td", "th"} or set(cell) - {"type", "children", "id"}:
                raise ValueError("表格单元格结构不正确")
            if "id" in cell and (not isinstance(cell["id"], str) or not cell["id"]):
                raise ValueError("表格单元格 ID 不正确")
            paragraphs = cell.get("children")
            if not isinstance(paragraphs, list) or not 1 <= len(paragraphs) <= 20:
                raise ValueError("表格单元格段落不正确")
            for paragraph in paragraphs:
                if not isinstance(paragraph, dict) or paragraph.get("type") != "p":
                    raise ValueError("表格单元格仅支持正文段落")
                cited |= _validate_text_block(paragraph, cell=True)
    if not cited.issubset(keys):
        raise ValueError("表格事实引用未绑定到表格")
    return keys


def source_display(source: str) -> str:
    """Turn a stored document locator into a readable citation without losing the original audit value."""
    match = re.fullmatch(r"document:[^#]+#p(\d+)-ocr(\d+)(?:\+[^ ]+)* · (.+)", source)
    if match:
        return f"{match[3]}，第 {match[1]} 页，OCR 识别片段 {match[2]}"
    match = re.fullmatch(r"document:[^#]+#p(\d+)-t(\d+)-r(\d+)(?:\+[^ ]+)* · (.+)", source)
    if match:
        return f"{match[4]}，第 {match[1]} 页，表 {match[2]}，第 {match[3]} 行"
    match = re.fullmatch(r"document:[^#]+#p(\d+)-s(\d+)(?:\+[^ ]+)* · (.+)", source)
    if match:
        return f"{match[3]}，第 {match[1]} 页，第 {match[2]} 段"
    match = re.fullmatch(r"document:[^#]+#d-p(\d+) · (.+)", source)
    if match:
        return f"{match[2]}，第 {match[1]} 段"
    match = re.fullmatch(r"document:[^#]+#d-t(\d+)-r(\d+) · (.+)", source)
    if match:
        return f"{match[3]}，表 {match[1]}，第 {match[2]} 行"
    return source


def paragraph_markup(node: dict) -> str:
    """Render supported inline nodes to ReportLab paragraph markup."""
    pieces = []
    for leaf in node.get("children", []):
        if leaf.get("type") == "a":
            pieces.append(f'<link href="{escape(leaf["url"], quote=True)}">{paragraph_markup(leaf)}</link>')
            continue
        if leaf.get("type") == "fact_ref":
            pieces.append(f'<b>{escape(leaf["display"])}</b>')
            continue
        value = escape(str(leaf.get("text", ""))).replace("\n", "<br/>")
        if leaf.get("bold"):
            value = f"<b>{value}</b>"
        if leaf.get("italic"):
            value = f"<i>{value}</i>"
        if leaf.get("underline"):
            value = f"<u>{value}</u>"
        if leaf.get("strikethrough"):
            value = f"<strike>{value}</strike>"
        pieces.append(value)
    return "".join(pieces)


def plain(node: dict) -> str:
    if "text" in node:
        return str(node["text"])
    if node.get("type") == "fact_ref":
        return str(node.get("display", ""))
    children = node.get("children", [])
    if node.get("type") == "table":
        return "\n".join(plain(child) for child in children if isinstance(child, dict))
    if node.get("type") == "tr":
        return "\t".join(plain(child) for child in children if isinstance(child, dict))
    return "".join(plain(child) for child in children if isinstance(child, dict))


def validate_content(content: list[dict]) -> None:
    if not isinstance(content, list) or not content or len(content) > 400:
        raise ValueError("报告正文必须包含 1 至 400 个段落")
    if len(json.dumps(content, ensure_ascii=False)) > 400_000:
        raise ValueError("报告正文超过当前版本容量")
    ids: set[str] = set()
    for node in content:
        if not isinstance(node, dict) or node.get("type") not in ALLOWED_BLOCKS:
            raise ValueError("正文包含不支持的段落类型")
        if node.get("type") == "table":
            _validate_table(node)
        else:
            _validate_text_block(node)
        node_id = node.get("id")
        if node_id:
            if node_id in ids:
                raise ValueError("报告段落 ID 重复")
            ids.add(node_id)


def content_hash(content: list[dict]) -> str:
    return hashlib.sha256(json.dumps(content, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode()).hexdigest()


def contains_fact_value(content: list[dict], key: str, value: str) -> bool:
    token = value.replace(",", "")
    for node in content:
        if key not in node.get("fact_keys", []):
            continue
        text = plain(node).replace(",", "")
        if re.search(rf"(?<![\d.]){re.escape(token)}(?![\d.])", text):
            return True
    return False


def numeric_tokens(value: str) -> set[str]:
    """Numbers visible in report prose, normalized for commas and leading zeroes."""
    tokens = re.findall(r"(?<![\d.])\d[\d,]*(?:\.\d+)?(?![\d.])", value)
    return {str(int(token.replace(",", ""))) if token.replace(",", "").isdigit()
            else token.replace(",", "") for token in tokens}


def _reference_display_numbers(node: dict, facts: dict, position: int) -> tuple[set[str], list[dict]]:
    """Accept a bound numeric token only when its visible value matches the fact or a known unit display."""
    allowed: set[str] = set()
    issues: list[dict] = []
    for child in node.get("children", []):
        if not isinstance(child, dict):
            continue
        if child.get("type") in {"table", "tr", "td", "th", "p", "a"}:
            nested, problems = _reference_display_numbers(child, facts, position)
            allowed |= nested
            issues.extend(problems)
        if child.get("type") != "fact_ref":
            continue
        key, display = child["fact_key"], child["display"]
        fact = facts.get(key)
        if fact is None or fact.value_text is None or fact.data_type not in {"integer", "decimal"}:
            continue
        tokens = numeric_tokens(display)
        if len(tokens) != 1:
            issues.append({"code": "FACT_REF_DISPLAY_MISMATCH", "message": f"第 {position} 段事实 {key} 的显示值无效",
                           "severity": "block", "fact_key": key, "position": position})
            continue
        displayed = Decimal(next(iter(tokens)))
        actual = Decimal(fact.value_text)
        if displayed == actual or (getattr(fact, "unit", "") == "万元/条" and displayed == actual * Decimal("10000")):
            allowed |= tokens
        else:
            issues.append({"code": "FACT_REF_DISPLAY_MISMATCH", "message": f"第 {position} 段事实 {key} 的显示值与当前项目不符",
                           "severity": "block", "fact_key": key, "position": position})
    return allowed, issues


def report_fact_impacts(content: list[dict], bound_facts: dict, current_facts: dict) -> list[dict]:
    """Locate paragraphs that still need a visible fact-reference update."""
    impacts = []
    for position, node in enumerate(content, 1):
        for key in dict.fromkeys(node.get("fact_keys", [])):
            fact = current_facts.get(key)
            old = bound_facts.get(key)
            new = {"value": fact.value_text, "revision": fact.revision} if fact else None
            if old != new:
                # A newly inserted chapter can already show the current value
                # while another paragraph still uses the older shared snapshot.
                # Show the user the paragraph that actually needs editing.
                if (old and new and old.get("value") != new.get("value")
                        and new.get("value") is not None
                        and contains_fact_value([node], key, new["value"])
                        and (old.get("value") is None
                             or not contains_fact_value([node], key, old["value"]))):
                    continue
                impacts.append({"position": position, "text": plain(node), "fact_key": key,
                                "before": old, "after": new})
    return impacts


def change_impact(before: list[dict], after: list[dict]) -> list[dict]:
    def signature(block: dict) -> str:
        visible = {key: value for key, value in block.items() if key not in {"id", "origin"}}
        return json.dumps(visible, ensure_ascii=False, sort_keys=True)

    changes = []
    matcher = SequenceMatcher(a=[signature(block) for block in before],
                              b=[signature(block) for block in after], autojunk=False)
    for kind, old_start, old_end, new_start, new_end in matcher.get_opcodes():
        if kind == "equal":
            continue
        count = max(old_end - old_start, new_end - new_start)
        for offset in range(count):
            old = before[old_start + offset] if old_start + offset < old_end else {}
            new = after[new_start + offset] if new_start + offset < new_end else {}
            position = (new_start + offset + 1) if new else (old_start + offset + 1)
            changes.append({"position": position, "before": plain(old), "after": plain(new),
                            "fact_keys": sorted(set(old.get("fact_keys", [])) | set(new.get("fact_keys", [])))})
    return changes


def gate(content: list[dict], reviewed_hash: str | None, bound_facts: dict, current_facts: dict,
         source_numbers: dict[int, set[str]] | None = None) -> list[dict]:
    issues = []
    if not any(plain(node).strip() for node in content):
        issues.append({"code": "EMPTY", "message": "报告正文为空", "severity": "block"})
    if content_hash(content) != reviewed_hash:
        issues.append({"code": "UNREVIEWED", "message": "正文保存后尚未人工核对", "severity": "block"})
    used = {key for node in content for key in node.get("fact_keys", [])}
    if not used and not any(node.get("analysis_refs") or node.get("source_refs") for node in content):
        issues.append({"code": "NO_FACT_LINK", "message": "正文没有绑定项目事实；请插入已确认事实并核对", "severity": "block"})
    for key in sorted(used):
        fact = current_facts.get(key)
        snapshot = bound_facts.get(key)
        if fact is None or fact.value_text is None or fact.value_status not in ("PROVIDED", "COMPUTED"):
            issues.append({"code": "FACT_MISSING", "message": f"引用事实 {key} 不存在或尚未定义", "severity": "block", "fact_key": key})
        elif snapshot is None or snapshot != {"value": fact.value_text, "revision": fact.revision}:
            issues.append({"code": "FACT_CHANGED", "message": f"引用事实 {key} 已变化，请核对并更新正文", "severity": "block", "fact_key": key})
    for position, node in enumerate(content, 1):
        displayed_numbers, display_issues = _reference_display_numbers(node, current_facts, position)
        issues.extend(display_issues)
        for key in dict.fromkeys(node.get("fact_keys", [])):
            fact = current_facts.get(key)
            if fact is not None and fact.value_text is not None and fact.data_type in ("integer", "decimal") and not contains_fact_value([node], key, fact.value_text):
                issues.append({"code": "FACT_TEXT_STALE", "message": f"第 {position} 段未找到事实 {key} 的当前数值", "severity": "block", "fact_key": key, "position": position})
        if node.get("type") in ("p", "blockquote", "table"):
            visible_numbers = numeric_tokens(plain(node))
            provenance_numbers = (source_numbers or {}).get(position, set())
            if visible_numbers and not node.get("fact_keys") and not provenance_numbers:
                issues.append({"code": "UNCITED_NUMBER", "message": f"第 {position} 段含数字但未绑定项目事实", "severity": "block", "position": position})
            else:
                allowed_numbers = displayed_numbers | provenance_numbers | set().union(*(numeric_tokens(current_facts[key].value_text)
                                                for key in node.get("fact_keys", [])
                                                if key in current_facts and current_facts[key].value_text))
                extra = sorted(visible_numbers - allowed_numbers)
                if extra:
                    issues.append({"code": "UNSUPPORTED_NUMBER", "message": f"第 {position} 段含未由绑定事实支持的数字：{'、'.join(extra)}", "severity": "block", "position": position})
    return issues


def model_section(title: str, facts: list[dict]) -> list[dict]:
    allowed = {item["key"]: item for item in facts}
    prompt = (
        "为专业报告的指定章节起草 1 至 3 个中文段落。只能复述给定项目事实，不新增任何数字、因果、评价或外部依据。"
        "各段合计必须覆盖所有给定事实。每段标出实际使用的 fact_keys，每段至少使用一个；"
        "每个列出的 fact_key 都要在该段保留独立可定位的原值，若两个事实的值相同也要分别写出；"
        "数值只可加千分位，不换算、不四舍五入。"
        "只返回 JSON {\"paragraphs\":[{\"text\":\"...\",\"fact_keys\":[\"...\"]}]}。"
        "生成结果只是待审核草稿。\n"
        f"章节：{title}\n已确认事实：{json.dumps(facts, ensure_ascii=False)}"
    )
    try:
        body = chat_json("writing", [{"role": "system", "content": "你只起草由项目事实直接支持的待审正文。"},
                                     {"role": "user", "content": prompt}], max_tokens=1200)
    except ValueError as exc:
        raise ValueError(f"{exc}；现有报告没有改变") from exc
    paragraphs = body.get("paragraphs")
    if not isinstance(paragraphs, list) or not 1 <= len(paragraphs) <= 3:
        raise ValueError("模型正文结构不符合要求；现有报告没有改变")
    output = []
    number_pattern = re.compile(r"(?<![\d.,])[-+]?\d[\d,]*(?:\.\d+)?(?![\d.,])")

    def mentioned_numbers(value: str) -> set[Decimal]:
        output: set[Decimal] = set()
        for match in number_pattern.finditer(value):
            token = match.group()
            unsigned = token.lstrip("+-")
            if "," in unsigned and not re.fullmatch(r"\d{1,3}(?:,\d{3})+(?:\.\d+)?", unsigned):
                raise ValueError("模型正文数字格式不正确；现有报告没有改变")
            output.add(Decimal(token.replace(",", "")))
        return output

    allowed_numbers = set().union(*(mentioned_numbers(str(item["value"])) for item in facts))
    for item in paragraphs:
        if not isinstance(item, dict) or not isinstance(item.get("text"), str) or not item["text"].strip():
            raise ValueError("模型正文结构不符合要求；现有报告没有改变")
        keys = item.get("fact_keys")
        if not isinstance(keys, list) or not keys or any(key not in allowed for key in keys):
            raise ValueError("模型引用了未提供的项目事实；现有报告没有改变")
        numbers = mentioned_numbers(item["text"])
        if numbers - allowed_numbers:
            raise ValueError("模型正文出现未提供的数字；现有报告没有改变")
        output.append({"type": "p", "children": [{"text": item["text"].strip()}], "fact_keys": keys, "origin": "model"})
    return output


def _word_inline(paragraph, node: dict) -> None:
    """Write the small approved inline AST, including real Word hyperlinks."""
    for child in node.get("children", []):
        if child.get("type") == "a":
            relationship_id = paragraph.part.relate_to(child["url"], RT.HYPERLINK, is_external=True)
            hyperlink = OxmlElement("w:hyperlink")
            hyperlink.set(qn("r:id"), relationship_id)
            for leaf in child["children"]:
                run = OxmlElement("w:r")
                properties = OxmlElement("w:rPr")
                color = OxmlElement("w:color")
                color.set(qn("w:val"), "2D61DC")
                properties.append(color)
                underline = OxmlElement("w:u")
                underline.set(qn("w:val"), "single")
                properties.append(underline)
                for mark, tag in (("bold", "b"), ("italic", "i"), ("strikethrough", "strike")):
                    if leaf.get(mark):
                        properties.append(OxmlElement(f"w:{tag}"))
                run.append(properties)
                text = OxmlElement("w:t")
                text.text = leaf["text"]
                run.append(text)
                hyperlink.append(run)
            paragraph._p.append(hyperlink)
            continue
        if child.get("type") == "fact_ref":
            run = paragraph.add_run(child["display"])
            run.bold = True
            continue
        run = paragraph.add_run(child["text"])
        run.bold = bool(child.get("bold"))
        run.italic = bool(child.get("italic"))
        run.underline = bool(child.get("underline"))
        run.font.strike = bool(child.get("strikethrough"))


def _pdf_node_markup(node: dict) -> str:
    return paragraph_markup(node)


def _fact_basis(fact: dict, audit: dict) -> str:
    """Describe a computed fact by its project rule and inputs, not a blank source."""
    evidence = next((entry for entry in audit.get("evidence_bindings", [])
                     if entry.get("fact_key") == fact.get("key")), {})
    rule = evidence.get("project_rule") or {}

    def locator(binding: dict) -> str:
        filename = binding.get("document_filename")
        document_id = binding.get("document_id")
        if not filename or not document_id:
            return ""
        labels = binding.get("location_labels") or []
        if labels:
            return "；".join(f"{filename}，{label}" for label in labels)
        locations = []
        for reference in binding.get("source_refs", []):
            display = source_display(f"document:{document_id}#{reference} · {filename}")
            locations.append(display if not display.startswith("document:") else f"{filename}，片段 {reference}")
        return "；".join(locations)

    if rule:
        labels = {entry["key"]: entry.get("label") or entry["key"] for entry in audit.get("facts", [])}
        inputs = "、".join(
            f"{labels.get(entry.get('fact_key'), entry.get('fact_key'))} r{entry.get('fact_revision')}"
            + (f"（{locator(entry)}）" if entry.get("status") == "source_locator_reviewed" and locator(entry) else "")
            for entry in evidence.get("input_fact_refs", []) if entry.get("fact_revision") is not None
        )
        basis = f"本项目计算规则：{rule.get('expression', '')}（规则 ID：{rule.get('rule_id', '')}）"
        if inputs:
            basis += f"；输入修订：{inputs}"
        basis += "；输入原文位置已由操作者核对" if evidence.get("status") == "derived_from_reviewed_inputs" else "；输入证据待核对"
        return basis
    if evidence.get("status") == "source_locator_reviewed" and locator(evidence):
        return f"本项目原件：{locator(evidence)}（位置已核对）"
    if evidence:
        original = source_display(fact.get("source") or "")
        return f"项目原件位置待核对；原录入来源：{original}" if original else "项目原件位置待核对"
    return f"来源：{source_display(fact.get('source') or '')}"


def _analysis_basis(audit: dict) -> list[str]:
    """Render only cited run values and their explicit upstream calculation steps."""
    snapshot = audit.get("analysis_run") or {}
    results = snapshot.get("results") or {}
    inputs = snapshot.get("inputs") or {}
    traces = {step.get("target"): step for step in snapshot.get("trace", [])
              if step.get("status") == "COMPUTED"}
    cited = {ref.get("result_key") for block in audit.get("analysis_refs", [])
             for ref in block.get("refs", []) if ref.get("result_key") in results}
    if not cited:
        return []

    needed: set[str] = set()

    def collect(key: str) -> None:
        if key in needed:
            return
        needed.add(key)
        for upstream in traces.get(key, {}).get("inputs", {}):
            if upstream in results:
                collect(upstream)

    for key in cited:
        collect(key)
    labels = {key: row.get("label") or key for key, row in results.items()}
    origin_labels = {"project_fact": "项目事实快照", "historical_reference": "历史参考",
                     "scenario_assumption": "方案假设", "missing": "待补"}
    lines = []
    for field in snapshot.get("definitions", []):
        key = field.get("key")
        if key not in needed or key not in inputs:
            continue
        row = results[key]
        value = row.get("value")
        origin = origin_labels.get(inputs[key].get("origin"), "方案输入")
        lines.append(f"输入：{labels[key]} {value if value is not None else '未定义'}{row.get('unit') or ''}（{origin}）")
    for step in snapshot.get("trace", []):
        key = step.get("target")
        if key not in needed or step.get("status") != "COMPUTED" or key not in results:
            continue
        expression = re.sub(r"\b[A-Za-z_][A-Za-z_0-9]*\b",
                            lambda match: labels.get(match.group(), match.group()),
                            step.get("expression") or "")
        substitutions = "、".join(
            f"{labels.get(dependency, dependency)} {value if value is not None else '未定义'}"
            f"{results.get(dependency, {}).get('unit') or ''}"
            for dependency, value in step.get("inputs", {}).items())
        row = results[key]
        lines.append(f"计算：{labels[key]} {row.get('value')}{row.get('unit') or ''}；"
                     f"规则：{expression}；代入：{substitutions}")
    return lines


def export_bundle(title: str, content: list[dict], audit: dict, preview_label: str | None = None) -> bytes:
    """Render one frozen report snapshot to DOCX, PDF and a JSON audit record."""
    word = Document()
    section = word.sections[0]
    section.top_margin = section.bottom_margin = Cm(2.3)
    word.styles["Normal"].font.name = "宋体"
    word.styles["Normal"].font.size = Pt(11)
    if preview_label:
        marker = word.add_paragraph(preview_label)
        marker.style = word.styles["Subtitle"]
    word.add_heading(title, 0)
    body = content[1:] if content and content[0].get("type") == "h1" and plain(content[0]).strip() == title.strip() else content
    for node in body:
        if node["type"] == "table":
            rows = node["children"]
            table = word.add_table(rows=len(rows), cols=len(rows[0]["children"]))
            table.style = "Table Grid"
            for row_index, row in enumerate(rows):
                for column_index, cell in enumerate(row["children"]):
                    target = table.cell(row_index, column_index)
                    for paragraph_index, source in enumerate(cell["children"]):
                        paragraph = target.paragraphs[0] if paragraph_index == 0 else target.add_paragraph()
                        if source.get("listStyleType") == "disc":
                            paragraph.style = "List Bullet"
                        elif source.get("listStyleType") == "decimal":
                            paragraph.style = "List Number"
                        if source.get("indent", 1) > 1:
                            paragraph.paragraph_format.left_indent = Cm(0.7 * source["indent"])
                        _word_inline(paragraph, source)
            continue
        if node["type"] in ("h1", "h2", "h3"):
            paragraph = word.add_heading(level=int(node["type"][1]))
        else:
            style = "List Bullet" if node.get("listStyleType") == "disc" else "List Number" if node.get("listStyleType") == "decimal" else None
            paragraph = word.add_paragraph(style=style)
            if node["type"] == "blockquote":
                paragraph.paragraph_format.left_indent = Cm(0.7)
                paragraph.paragraph_format.right_indent = Cm(0.4)
            elif node.get("indent", 1) > 1:
                paragraph.paragraph_format.left_indent = Cm(0.7 * node["indent"])
        _word_inline(paragraph, node)
    basis_lines = _analysis_basis(audit)
    if basis_lines:
        word.add_heading("推演依据", level=2)
        for line in basis_lines:
            word.add_paragraph(line)
    if audit["facts"]:
        word.add_heading("事实依据", level=2)
        for fact in audit["facts"]:
            word.add_paragraph(f"{fact.get('label') or fact['key']}：{fact['value']}{fact['unit']}。{_fact_basis(fact, audit)}")
    docx_buffer = io.BytesIO()
    word.save(docx_buffer)

    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    styles = {key: ParagraphStyle(key, fontName="STSong-Light", fontSize=size, leading=size * 1.55,
                                  spaceAfter=12 if key == "title" else 9,
                                  alignment=1 if key == "title" else 0,
                                  textColor=colors.HexColor("#24344D"))
              for key, size in (("title", 20), ("h1", 16), ("h2", 14), ("h3", 12), ("p", 11))}
    styles["blockquote"] = ParagraphStyle("blockquote", parent=styles["p"], leftIndent=22, rightIndent=10,
                                            textColor=colors.HexColor("#53667F"))
    styles["basis"] = ParagraphStyle("basis", parent=styles["p"], fontSize=9.5, leading=14, spaceAfter=5)
    styles["table_cell"] = ParagraphStyle("table_cell", parent=styles["p"], fontSize=8, leading=12, spaceAfter=0)
    pdf_buffer = io.BytesIO()
    # STSong-Light renders the middle dot in the DOCX preview marker as a
    # missing-glyph box. Use a supported Chinese separator in the PDF header.
    pdf_preview_label = preview_label.replace(" · ", "，") if preview_label else None
    story = ([Paragraph(escape(pdf_preview_label), styles["h2"]), Spacer(1, 7)] if pdf_preview_label else [])
    story.extend([Paragraph(escape(title.replace(" · ", "，").replace("·", "，")), styles["title"]), Spacer(1, 10)])
    list_number = 0
    for node in body:
        if node["type"] == "table":
            rows = node["children"]
            def cell_markup(cell: dict) -> str:
                parts: list[str] = []
                cell_number = 0
                for paragraph in cell["children"]:
                    markup = _pdf_node_markup(paragraph)
                    if paragraph.get("listStyleType") == "disc":
                        markup = "• " + markup
                        cell_number = 0
                    elif paragraph.get("listStyleType") == "decimal":
                        cell_number = paragraph.get("listStart", cell_number + 1)
                        markup = f"{cell_number}. " + markup
                    else:
                        cell_number = 0
                    parts.append(markup)
                return "<br/>".join(parts) or " "
            grid = [[Paragraph(cell_markup(cell), styles["table_cell"])
                     for cell in row["children"]] for row in rows]
            table = Table(grid, colWidths=(A4[0] - 104) / len(grid[0]), repeatRows=1 if all(cell["type"] == "th" for cell in rows[0]["children"]) else 0)
            table.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#CDD7E5")),
                                       ("VALIGN", (0, 0), (-1, -1), "TOP"),
                                       ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#F3F6FB"))]))
            story.extend([Spacer(1, 5), table, Spacer(1, 9)])
            list_number = 0
            continue
        markup = _pdf_node_markup(node)
        if node.get("listStyleType") == "disc":
            markup = "• " + markup
            list_number = 0
        elif node.get("listStyleType") == "decimal":
            list_number = node.get("listStart", list_number + 1)
            markup = f"{list_number}. " + markup
        else:
            list_number = 0
        story.append(Paragraph(markup or " ", styles[node["type"] if node["type"] in styles else "p"]))
    if basis_lines:
        story.extend([Spacer(1, 12), Paragraph("推演依据", styles["h2"])])
        for line in basis_lines:
            story.append(Paragraph(escape(line), styles["basis"]))
    if audit["facts"]:
        story.extend([Spacer(1, 12), Paragraph("事实依据", styles["h2"])])
        for fact in audit["facts"]:
            citation = f"{fact.get('label') or fact['key']}：{fact['value']}{fact['unit']}。{_fact_basis(fact, audit)}"
            story.append(Paragraph(escape(citation), styles["basis"]))
    def footer(canvas, doc):
        canvas.saveState()
        canvas.setFont("STSong-Light", 9)
        canvas.setFillColor(colors.HexColor("#8190A4"))
        canvas.drawRightString(A4[0] - 52, 32, f"第 {doc.page} 页")
        canvas.restoreState()
    SimpleDocTemplate(pdf_buffer, pagesize=A4, leftMargin=52, rightMargin=52, topMargin=56, bottomMargin=56).build(
        story, onFirstPage=footer, onLaterPages=footer)

    docx_bytes, pdf_bytes = docx_buffer.getvalue(), pdf_buffer.getvalue()
    audit = {**audit, "docx_sha256": hashlib.sha256(docx_bytes).hexdigest(),
             "pdf_sha256": hashlib.sha256(pdf_bytes).hexdigest()}
    bundle = io.BytesIO()
    with ZipFile(bundle, "w", ZIP_DEFLATED) as archive:
        archive.writestr("report.docx", docx_bytes)
        archive.writestr("report.pdf", pdf_bytes)
        archive.writestr("audit.json", json.dumps(audit, ensure_ascii=False, indent=2))
    return bundle.getvalue()
