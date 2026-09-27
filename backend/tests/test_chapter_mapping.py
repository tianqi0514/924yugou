"""Chapter mapping uses only project-owned, source-bound facts and immutable configs."""

import io
import os
import re
from pathlib import Path
from zipfile import ZipFile

os.environ["DATABASE_URL"] = "postgresql+psycopg://report:local_development_only@127.0.0.1:55432/report_platform_test"

import pytest
from docx import Document
from fastapi.testclient import TestClient
import pymupdf
from sqlalchemy import func, select

from app.db import AnalysisConfig, Base, ProjectCorpus, SessionLocal, engine
from app.main import app


REPORTS = Path(__file__).resolve().parents[2] / "test-fixtures/public-reports"


@pytest.fixture(autouse=True)
def isolated_database():
    assert engine.url.database == "report_platform_test", "Refusing to reset a non-test database"
    writable = [table for table in Base.metadata.sorted_tables if not table.name.startswith("corpus_")]
    Base.metadata.drop_all(engine, tables=writable)
    Base.metadata.create_all(engine, tables=writable)
    yield
    Base.metadata.drop_all(engine, tables=writable)


@pytest.fixture
def client(tmp_path, monkeypatch):
    import app.main as main
    monkeypatch.setattr(main, "STORAGE", tmp_path)
    with TestClient(app) as value:
        yield value


def setup_project(client, lines):
    project = client.post("/api/projects", json={"name": "映射测试"}).json()
    base = f"/api/projects/{project['id']}"
    doc = Document()
    for line in lines:
        doc.add_paragraph(line)
    stream = io.BytesIO(); doc.save(stream)
    uploaded = client.post(base + "/documents", files={"file": (
        "项目原文.docx", stream.getvalue(),
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document")})
    assert uploaded.status_code == 201, uploaded.text
    document_id = uploaded.json()["id"]
    segments = client.get(base + f"/documents/{document_id}").json()["segments"]
    report = client.post(base + "/reports", json={
        "title": "可研工作稿", "report_type": "government_feasibility"})
    assert report.status_code == 201, report.text
    return base, document_id, segments, report.json()["id"]


def add_fact(client, base, doc_id, segment, *, key, label, value, data_type="text", unit=""):
    created = client.post(base + "/facts", json={
        "key": key, "label": label, "data_type": data_type,
        "value": value, "unit": unit, "source": "项目原文"})
    assert created.status_code == 201, created.text
    bound = client.post(base + f"/facts/{key}/evidence/bind", json={
        "document_id": doc_id, "source_refs": [segment["ref"]]})
    assert bound.status_code == 200, bound.text


def test_narrative_mapping_preview_cancel_commit_and_draft(client):
    statement = "本项目计划建设交通枢纽并与轨道交通衔接。"
    base, doc, segments, report = setup_project(client, [statement])
    add_fact(client, base, doc, segments[0], key="station_plan", label="枢纽建设计划",
             value=statement)
    path = base + f"/reports/{report}/chapter-mapping/necessity"
    options = client.get(path)
    assert options.status_code == 200, options.text
    assert options.json()["facts"][0]["source"]["verified"] is True
    choice = {"mode": "narrative", "fact_keys": ["station_plan"], "result_keys": []}
    preview = client.post(path + "/preview", json=choice)
    assert preview.status_code == 200, preview.text
    with SessionLocal() as session:
        assert session.scalar(select(func.count()).select_from(AnalysisConfig)) == 0
    # Closing the preview is a read-only action; the same token is still valid.
    committed = client.post(path + "/commit", json={**choice, **{
        key: preview.json()[key] for key in ("preview_digest", "preview_token", "expires_at")}})
    assert committed.status_code == 200, committed.text
    assert committed.json()["status"] == "PUBLISHED"
    repeated = client.post(path + "/commit", json={**choice, **{
        key: preview.json()[key] for key in ("preview_digest", "preview_token", "expires_at")}})
    assert repeated.status_code == 200 and repeated.json()["id"] == committed.json()["id"]
    chapters = client.get(base + f"/reports/{report}/chapters").json()["chapters"]
    assert next(row for row in chapters if row["section_id"] == "necessity")["status"] == "可起草"
    candidate = client.post(base + f"/analysis/reports/{report}/draft/preview", json={
        "config_id": committed.json()["id"], "section_id": "necessity", "mode": "excerpt"})
    assert candidate.status_code == 200, candidate.text
    assert statement in str(candidate.json()["content"])


def test_mapping_rejects_stale_and_cross_project(client):
    statement = "本项目计划建设交通枢纽。"
    base, doc, segments, report = setup_project(client, [statement])
    add_fact(client, base, doc, segments[0], key="station_plan", label="建设计划",
             value=statement)
    path = base + f"/reports/{report}/chapter-mapping/necessity"
    choice = {"mode": "narrative", "fact_keys": ["station_plan"], "result_keys": []}
    preview = client.post(path + "/preview", json=choice).json()
    changed = client.post(base + "/changes/preview", json={
        "fact_key": "station_plan", "value": "本项目拟建设交通枢纽。",
        "source": "项目原文", "reason": "修订"}).json()
    assert client.post(base + "/changes/commit", json={
        "fact_key": "station_plan", "value": "本项目拟建设交通枢纽。",
        "source": "项目原文", "reason": "修订",
        "base_version": changed["base_version"],
        "preview_token": changed["preview_token"]}).status_code == 200
    stale = client.post(path + "/commit", json={**choice, **{
        key: preview[key] for key in ("preview_digest", "preview_token", "expires_at")}})
    assert stale.status_code == 409
    other = client.post("/api/projects", json={"name": "其他项目"}).json()["id"]
    assert client.get(f"/api/projects/{other}/reports/{report}/chapter-mapping/necessity").status_code == 404
    with SessionLocal() as session:
        assert session.scalar(select(func.count()).select_from(AnalysisConfig)) == 0


def test_same_section_id_is_scoped_to_report_type(client):
    statement = "建设资金由本项目财政预算安排。"
    base, doc, segments, government_id = setup_project(client, [statement])
    add_fact(client, base, doc, segments[0], key="funding_source", label="资金来源",
             value=statement)
    enterprise = client.post(base + "/reports", json={
        "title": "企业可研", "report_type": "enterprise_feasibility"})
    assert enterprise.status_code == 201, enterprise.text
    enterprise_id = enterprise.json()["id"]
    choice = {"mode": "narrative", "fact_keys": ["funding_source"], "result_keys": []}
    government_path = base + f"/reports/{government_id}/chapter-mapping/investment"
    enterprise_path = base + f"/reports/{enterprise_id}/chapter-mapping/investment"
    preview = client.post(government_path + "/preview", json=choice)
    assert preview.status_code == 200, preview.text
    government_config = client.post(government_path + "/commit", json={
        **choice, **{key: preview.json()[key] for key in
                   ("preview_digest", "preview_token", "expires_at")}})
    assert government_config.status_code == 200, government_config.text
    assert government_config.json()["sections"][0]["report_type"] == "government_feasibility"
    assert client.get(enterprise_path).json()["current_config"] is None
    enterprise_chapters = client.get(base + f"/reports/{enterprise_id}/chapters").json()["chapters"]
    assert next(row for row in enterprise_chapters if row["section_id"] == "investment")["status"] == "待配置"
    wrong_draft = client.post(base + f"/analysis/reports/{enterprise_id}/draft/preview", json={
        "config_id": government_config.json()["id"], "section_id": "investment", "mode": "excerpt"})
    assert wrong_draft.status_code == 409
    scenario = client.post(base + "/analysis/scenarios", json={
        "name": "政府配置运行", "source": "config", "config_id": government_config.json()["id"],
        "use_project_facts": True})
    assert scenario.status_code == 201, scenario.text
    run = client.post(base + f"/analysis/scenarios/{scenario.json()['id']}/runs", json={
        "scenario_revision": scenario.json()["revision"], "request_key": "cross-report-type-v1"})
    assert run.status_code == 201, run.text
    wrong_run = client.post(base + f"/analysis/reports/{enterprise_id}/select", json={
        "run_id": run.json()["id"], "base_version": 0})
    assert wrong_run.status_code == 409
    enterprise_preview = client.post(enterprise_path + "/preview", json=choice)
    assert enterprise_preview.status_code == 200, enterprise_preview.text
    enterprise_config = client.post(enterprise_path + "/commit", json={
        **choice, **{key: enterprise_preview.json()[key] for key in
                   ("preview_digest", "preview_token", "expires_at")}})
    assert enterprise_config.status_code == 200, enterprise_config.text
    assert enterprise_config.json()["id"] != government_config.json()["id"]
    assert enterprise_config.json()["sections"][0]["report_type"] == "enterprise_feasibility"
    candidate = client.post(base + f"/analysis/reports/{enterprise_id}/draft/preview", json={
        "config_id": enterprise_config.json()["id"], "section_id": "investment", "mode": "excerpt"})
    assert candidate.status_code == 200, candidate.text
    assert statement in str(candidate.json()["content"])


def test_calculation_mapping_uses_project_rule_and_keeps_zero(client):
    base, doc, segments, report = setup_project(client, [
        "首年需求为0套。", "合格能力为254016套。"])
    add_fact(client, base, doc, segments[0], key="first_year_demand", label="首年需求",
             value="0", data_type="integer", unit="套")
    add_fact(client, base, doc, segments[1], key="qualified_capacity", label="合格能力",
             value="254016", data_type="integer", unit="套")
    created = client.post(base + "/facts", json={
        "key": "planned_sales", "label": "计划销售量", "data_type": "integer", "unit": "套"})
    assert created.status_code == 201, created.text
    rule = client.post(base + "/rules", json={"name": "销售取小", "target_key": "planned_sales",
                                                "expression": "min(first_year_demand, qualified_capacity)"})
    assert rule.status_code == 201, rule.text
    path = base + f"/reports/{report}/chapter-mapping/demand"
    choice = {"mode": "calculation", "fact_keys": ["planned_sales"],
              "result_keys": ["planned_sales"]}
    preview = client.post(path + "/preview", json=choice)
    assert preview.status_code == 200, preview.text
    assert preview.json()["results"][0]["value"] == "0"
    committed = client.post(path + "/commit", json={**choice, **{
        key: preview.json()[key] for key in ("preview_digest", "preview_token", "expires_at")}})
    assert committed.status_code == 200, committed.text
    assert {item["key"] for item in committed.json()["definitions"]} == {
        "first_year_demand", "qualified_capacity", "planned_sales"}
    assert committed.json()["rules"][0]["expression"] == "min(first_year_demand, qualified_capacity)"
    scenario = client.post(base + "/analysis/scenarios", json={
        "name": "零需求方案", "source": "config", "config_id": committed.json()["id"],
        "use_project_facts": True})
    assert scenario.status_code == 201, scenario.text
    run = client.post(base + f"/analysis/scenarios/{scenario.json()['id']}/runs", json={
        "scenario_revision": scenario.json()["revision"], "request_key": "mapping-zero-v1"})
    assert run.status_code == 201, run.text
    assert run.json()["snapshot"]["results"]["planned_sales"]["value"] == "0"
    selected = client.post(base + f"/analysis/reports/{report}/select", json={
        "run_id": run.json()["id"], "base_version": 0})
    assert selected.status_code == 200, selected.text
    candidate = client.post(base + f"/analysis/reports/{report}/draft/preview", json={
        "run_id": run.json()["id"], "section_id": "demand", "mode": "computed"})
    assert candidate.status_code == 200, candidate.text
    assert "计划销售量0套" in str(candidate.json()["content"])
    adopted = client.post(base + f"/analysis/reports/{report}/draft/commit", json={
        **{key: value for key, value in candidate.json().items()
           if key not in {"issues", "preserved_blocks"}}, "replace_section": True})
    assert adopted.status_code == 200, adopted.text
    add_fact(client, base, doc, segments[0], key="demand_note", label="首年需求原文表述",
             value="首年需求为0套。")
    narrative_path = base + f"/reports/{report}/chapter-mapping/necessity"
    narrative_choice = {"mode": "narrative", "fact_keys": ["demand_note"], "result_keys": []}
    narrative_preview = client.post(narrative_path + "/preview", json=narrative_choice)
    assert narrative_preview.status_code == 200, narrative_preview.text
    narrative_config = client.post(narrative_path + "/commit", json={
        **narrative_choice, **{key: narrative_preview.json()[key] for key in
                             ("preview_digest", "preview_token", "expires_at")}})
    assert narrative_config.status_code == 200, narrative_config.text
    narrative = client.post(base + f"/analysis/reports/{report}/draft/preview", json={
        "config_id": narrative_config.json()["id"], "section_id": "necessity",
        "mode": "excerpt"})
    assert narrative.status_code == 200, narrative.text
    assert client.post(base + f"/analysis/reports/{report}/draft/commit", json={
        **{key: value for key, value in narrative.json().items()
           if key not in {"issues", "preserved_blocks"}}, "replace_section": True}).status_code == 200
    exported = client.get(base + f"/reports/{report}/export?level=preview")
    assert exported.status_code == 200, exported.text[:300]
    with ZipFile(io.BytesIO(exported.content)) as archive:
        docx = Document(io.BytesIO(archive.read("report.docx")))
        with pymupdf.open(stream=archive.read("report.pdf"), filetype="pdf") as pdf:
            pdf_text = "".join(page.get_text() for page in pdf)
        audit = archive.read("audit.json").decode()
    docx_text = "\n".join(paragraph.text for paragraph in docx.paragraphs)
    assert "计划销售量0套" in docx_text and "首年需求为0套" in docx_text
    assert "计划销售量0套" in pdf_text and "首年需求为0套" in pdf_text
    assert run.json()["id"] in audit and "preview_only" in audit


def test_two_public_feasibility_reports_keep_own_source_and_chapter(client):
    cases = [
        ("01_beijing_feasibility.pdf", 17, "necessity", "公交接驳计划",
         lambda text: re.search(r"马昌营公共交通枢纽紧邻.{0,85}?承接地铁客流", text).group()),
        ("05_fengxian_feasibility.pdf", 7, "investment", "财政资金来源",
         lambda text: "资金由区财政解决"),
    ]
    paths = []
    for filename, page, section, label, extract in cases:
        project = client.post("/api/projects", json={"name": filename}).json()
        base = f"/api/projects/{project['id']}"
        uploaded = client.post(base + "/documents", files={"file": (
            filename, (REPORTS / filename).read_bytes(), "application/pdf")})
        assert uploaded.status_code == 201, uploaded.text
        doc_id = uploaded.json()["id"]
        segments = client.get(base + f"/documents/{doc_id}?page={page}").json()["segments"]
        marker = "马昌营公共交通枢纽紧邻" if section == "necessity" else "资金由区财政解决"
        segment = next(item for item in segments if marker in item["text"])
        statement = extract(segment["text"])
        key = "transit_plan" if section == "necessity" else "funding_source"
        add_fact(client, base, doc_id, segment, key=key, label=label, value=statement)
        report = client.post(base + "/reports", json={
            "title": "本项目可研", "report_type": "government_feasibility"}).json()
        path = base + f"/reports/{report['id']}/chapter-mapping/{section}"
        preview = client.post(path + "/preview", json={
            "mode": "narrative", "fact_keys": [key], "result_keys": []})
        assert preview.status_code == 200, preview.text
        committed = client.post(path + "/commit", json={
            "mode": "narrative", "fact_keys": [key], "result_keys": [],
            **{name: preview.json()[name] for name in
               ("preview_digest", "preview_token", "expires_at")}})
        assert committed.status_code == 200, committed.text
        candidate = client.post(base + f"/analysis/reports/{report['id']}/draft/preview", json={
            "config_id": committed.json()["id"], "section_id": section, "mode": "excerpt"})
        assert candidate.status_code == 200, candidate.text
        assert statement in str(candidate.json()["content"])
        assert "禾进装备" not in str(candidate.json()["content"])
        paths.append((base, report["id"], section, key))
    first = client.get(paths[0][0] + f"/reports/{paths[0][1]}/chapter-mapping/{paths[0][2]}").json()
    second = client.get(paths[1][0] + f"/reports/{paths[1][1]}/chapter-mapping/{paths[1][2]}").json()
    assert [row["key"] for row in first["facts"]] == [paths[0][3]]
    assert [row["key"] for row in second["facts"]] == [paths[1][3]]


def test_evidence_only_fact_maps_and_drafts_from_current_original(client):
    statement = "本项目拟建设公交换乘枢纽，并与轨道交通衔接。"
    base, doc_id, segments, report_id = setup_project(client, [statement])
    created = client.post(base + "/facts", json={
        "key": "transit_plan", "label": "换乘方案", "data_type": "text",
        "value": statement, "source": "项目原文"})
    assert created.status_code == 201, created.text
    first_payload = {"document_id": doc_id, "source_refs": [segments[0]["ref"]],
                     "statement": statement, "label": "原文方案", "fact_key": "transit_plan"}
    first = client.post(base + "/evidence", json=first_payload)
    assert first.status_code == 201, first.text
    path = base + f"/reports/{report_id}/chapter-mapping/necessity"
    option = client.get(path).json()["facts"][0]
    assert option["source"]["verified"] is True
    assert option["source"]["location"]["evidence_id"] == first.json()["id"]
    choice = {"mode": "narrative", "fact_keys": ["transit_plan"], "result_keys": []}
    preview = client.post(path + "/preview", json=choice)
    assert preview.status_code == 200, preview.text
    config = client.post(path + "/commit", json={
        **choice, **{key: preview.json()[key] for key in
                   ("preview_digest", "preview_token", "expires_at")}})
    assert config.status_code == 200, config.text
    draft = base + f"/analysis/reports/{report_id}/draft/preview"
    candidate = client.post(draft, json={
        "config_id": config.json()["id"], "section_id": "necessity", "mode": "excerpt"})
    assert candidate.status_code == 200, candidate.text
    assert candidate.json()["content"][1]["project_evidence_refs"][0]["evidence_id"] == first.json()["id"]
    adopted = client.post(base + f"/analysis/reports/{report_id}/draft/commit", json={
        **{key: value for key, value in candidate.json().items()
           if key not in {"issues", "preserved_blocks"}}, "replace_section": True})
    assert adopted.status_code == 200, adopted.text
    exported = client.get(base + f"/reports/{report_id}/export?level=preview")
    assert exported.status_code == 200, exported.text[:300]
    with ZipFile(io.BytesIO(exported.content)) as archive:
        docx = Document(io.BytesIO(archive.read("report.docx")))
        with pymupdf.open(stream=archive.read("report.pdf"), filetype="pdf") as pdf:
            pdf_text = "".join(page.get_text() for page in pdf)
        audit = archive.read("audit.json").decode()
    assert statement in "\n".join(paragraph.text for paragraph in docx.paragraphs)
    assert statement in pdf_text
    assert first.json()["id"] in audit

    supplement = Document(); supplement.add_paragraph("补充来源：" + statement)
    stream = io.BytesIO(); supplement.save(stream)
    uploaded = client.post(base + "/documents", files={"file": (
        "补充来源.docx", stream.getvalue(),
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document")})
    assert uploaded.status_code == 201, uploaded.text
    second_id = uploaded.json()["id"]
    second_ref = client.get(base + f"/documents/{second_id}").json()["segments"][0]["ref"]
    second_payload = {**first_payload, "document_id": second_id, "source_refs": [second_ref]}
    second = client.post(base + "/evidence", json=second_payload)
    assert second.status_code == 201, second.text
    both = client.post(draft, json={
        "config_id": config.json()["id"], "section_id": "necessity", "mode": "excerpt",
        "evidence_ids": [first.json()["id"], second.json()["id"]]})
    assert both.status_code == 200, both.text
    assert {row["evidence_id"] for row in both.json()["content"][1]["project_evidence_refs"]} == {
        first.json()["id"], second.json()["id"]}

    replacement_preview = client.post(base + f"/evidence/{first.json()['id']}/replacement/preview",
                                      json=second_payload)
    assert replacement_preview.status_code == 200, replacement_preview.text
    replacement = client.post(base + f"/evidence/{first.json()['id']}/replacement/commit", json={
        **second_payload, **{key: replacement_preview.json()[key] for key in
                           ("base_report_versions", "expires_at", "preview_token")}})
    assert replacement.status_code == 201, replacement.text
    assert client.post(base + "/evidence", json=first_payload).status_code == 409
    repeated_current = client.post(base + "/evidence", json=second_payload)
    assert repeated_current.status_code == 201 and repeated_current.json()["id"] in {
        second.json()["id"], replacement.json()["id"]}
    listed = client.get(base + "/evidence?fact_key=transit_plan").json()["items"]
    assert next(row for row in listed if row["id"] == first.json()["id"])["superseded_by_id"] == replacement.json()["id"]
    assert next(row for row in listed if row["id"] == second.json()["id"])["superseded_by_id"] is None
    materials = client.get(base + f"/reports/{report_id}/materials?section_id=necessity").json()
    states = {row["id"]: row["source_state"] for row in materials["project_evidence"]}
    assert states[first.json()["id"]] == "已更新"
    assert states[second.json()["id"]] == "当前原文位置"
    stale = client.post(draft, json={
        "config_id": config.json()["id"], "section_id": "necessity", "mode": "excerpt",
        "evidence_ids": [first.json()["id"]]})
    assert stale.status_code == 409
    fresh = client.post(draft, json={
        "config_id": config.json()["id"], "section_id": "necessity", "mode": "excerpt"})
    assert fresh.status_code == 200, fresh.text
    assert fresh.json()["content"][1]["project_evidence_refs"][0]["evidence_id"] != first.json()["id"]


def test_secondary_replacement_does_not_keep_original_gate_open(client):
    statement = "本项目拟建设公交换乘枢纽。"
    base, doc_id, segments, report_id = setup_project(client, [statement])
    assert client.post(base + "/facts", json={
        "key": "transit_plan", "label": "换乘方案", "data_type": "text",
        "value": statement, "source": "项目原文"}).status_code == 201
    payload = {"document_id": doc_id, "source_refs": [segments[0]["ref"]],
               "statement": statement, "label": "原文方案", "fact_key": "transit_plan"}
    original = client.post(base + "/evidence", json=payload)
    assert original.status_code == 201, original.text
    path = base + f"/reports/{report_id}/chapter-mapping/necessity"
    choice = {"mode": "narrative", "fact_keys": ["transit_plan"], "result_keys": []}
    mapping_preview = client.post(path + "/preview", json=choice)
    assert mapping_preview.status_code == 200, mapping_preview.text
    replacement_body = {**payload, "source_type": "secondary"}
    preview = client.post(base + f"/evidence/{original.json()['id']}/replacement/preview",
                          json=replacement_body)
    assert preview.status_code == 200, preview.text
    replaced = client.post(base + f"/evidence/{original.json()['id']}/replacement/commit", json={
        **replacement_body, **{key: preview.json()[key] for key in
                              ("base_report_versions", "expires_at", "preview_token")}})
    assert replaced.status_code == 201, replaced.text
    assert client.post(base + "/evidence", json=payload).status_code == 409
    secondary_retry = client.post(base + "/evidence", json=replacement_body)
    assert secondary_retry.status_code == 201 and secondary_retry.json()["id"] == replaced.json()["id"]
    source = client.get(path).json()["facts"][0]["source"]
    assert source == {"verified": False, "location": None}
    stale = client.post(path + "/commit", json={
        **choice, **{key: mapping_preview.json()[key] for key in
                   ("preview_digest", "preview_token", "expires_at")}})
    assert stale.status_code == 409
    assert client.post(path + "/preview", json={
        **choice}).status_code == 409


def test_historical_project_cannot_create_or_replace_project_evidence(client):
    statement = "历史项目原文仅供查阅。"
    base, doc_id, segments, _ = setup_project(client, [statement])
    payload = {"document_id": doc_id, "source_refs": [segments[0]["ref"]],
               "statement": statement, "label": "历史来源"}
    original = client.post(base + "/evidence", json=payload)
    assert original.status_code == 201, original.text
    project_id = base.removeprefix("/api/projects/")
    with SessionLocal.begin() as session:
        session.add(ProjectCorpus(project_id=project_id, corpus_id="qa-frozen-history"))
    assert client.post(base + "/evidence", json=payload).status_code == 403
    assert client.post(base + f"/evidence/{original.json()['id']}/replacement/preview",
                       json=payload).status_code == 403
    assert client.get(base + f"/evidence/{original.json()['id']}").status_code == 200
