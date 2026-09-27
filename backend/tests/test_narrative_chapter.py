"""NEXT-28: a source-backed narrative chapter needs no calculation run."""

import io
import os
from uuid import uuid4

os.environ["DATABASE_URL"] = "postgresql+psycopg://report:local_development_only@127.0.0.1:55432/report_platform_test"

import pytest
from docx import Document
from fastapi.testclient import TestClient

from app.db import Base, SessionLocal, ReportDraft, SourceDocument, SourceParseRevision, engine
from app.main import app
from app.model_settings import ChatResult


@pytest.fixture(autouse=True)
def isolated_database():
    assert engine.url.database == "report_platform_test", "Refusing to reset a non-test database"
    writable = [table for table in Base.metadata.sorted_tables if not table.name.startswith("corpus_")]
    Base.metadata.drop_all(engine, tables=writable)
    Base.metadata.create_all(engine, tables=writable)
    yield
    Base.metadata.drop_all(engine, tables=writable)


@pytest.fixture
def client():
    with TestClient(app) as value:
        yield value


def setup_narrative(client, monkeypatch, tmp_path):
    import app.main as main
    monkeypatch.setattr(main, "STORAGE", tmp_path)
    project = client.post("/api/projects", json={"name": "文字章 QA"}).json()
    base = f"/api/projects/{project['id']}"
    word = Document()
    statement = "马昌营公共交通枢纽站计划与地铁同期开通并承接地铁客流，尚非已签订单。"
    word.add_paragraph(statement)
    stream = io.BytesIO(); word.save(stream)
    uploaded = client.post(base + "/documents", files={"file": (
        "文字章原文.docx", stream.getvalue(),
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document")})
    assert uploaded.status_code == 201, uploaded.text
    doc_id = uploaded.json()["id"]
    segment = client.get(base + f"/documents/{doc_id}").json()["segments"][0]["ref"]
    assert client.post(base + "/facts", json={"key": "transit_plan", "label": "接驳计划",
        "data_type": "text", "source": "项目建议书"}).status_code == 201
    changed = client.post(base + "/changes/preview", json={"fact_key": "transit_plan",
        "value": statement, "source": "项目建议书", "reason": "核对原文"}).json()
    committed = client.post(base + "/changes/commit", json={"fact_key": "transit_plan",
        "value": statement, "source": "项目建议书", "reason": "核对原文",
        "base_version": changed["base_version"], "preview_token": changed["preview_token"]})
    assert committed.status_code == 200, committed.text
    bound = client.post(base + "/facts/transit_plan/evidence/bind", json={
        "document_id": doc_id, "source_refs": [segment]})
    assert bound.status_code == 200, bound.text
    config = client.post(base + "/analysis/configs", json={"name": "项目背景"}).json()
    updated = client.put(base + f"/analysis/configs/{config['id']}", json={
        "revision": config["revision"], "name": "项目背景",
        "definitions": [{"key": "transit_plan", "label": "接驳计划", "data_type": "text",
                         "unit": "", "group": "背景", "computed": False}],
        "rules": [], "sections": [{"id": "background", "title": "项目背景",
                                 "kind": "narrative", "result_keys": [],
                                 "evidence_keys": ["transit_plan"],
                                 "forbidden_terms": ["资金已落实"]}]})
    assert updated.status_code == 200, updated.text
    config = updated.json()
    trial = client.post(base + f"/analysis/configs/{config['id']}/test", json={
        "revision": config["revision"], "sample_inputs": {"transit_plan": statement}})
    assert trial.status_code == 200 and trial.json()["test_token"], trial.text
    published = client.post(base + f"/analysis/configs/{config['id']}/publish", json={
        "revision": config["revision"], "test_token": trial.json()["test_token"]})
    assert published.status_code == 200, published.text
    report = client.post(base + "/reports", json={"title": "项目背景报告"}).json()
    return base, config, report, statement, doc_id


def test_narrative_excerpt_cancel_commit_review_and_stale(client, monkeypatch, tmp_path):
    base, config, report, statement, doc_id = setup_narrative(client, monkeypatch, tmp_path)
    path = base + f"/analysis/reports/{report['id']}/draft"
    request = {"config_id": config["id"], "section_id": "background", "mode": "excerpt"}
    preview = client.post(path + "/preview", json=request)
    assert preview.status_code == 200, preview.text
    candidate = preview.json()
    assert "run_id" not in candidate
    assert candidate["content"][1]["fact_keys"] == ["transit_plan"]
    assert client.get(base + f"/reports/{report['id']}").json()["version"] == 0
    other = client.post("/api/projects", json={"name": "另一项目"}).json()
    assert client.post(f"/api/projects/{other['id']}/analysis/reports/{report['id']}/draft/preview",
                       json=request).status_code == 404
    accepted = client.post(path + "/commit", json={**{key: value for key, value in candidate.items()
        if key not in {"issues", "preserved_blocks"}}, "replace_section": False})
    assert accepted.status_code == 200, accepted.text
    assert accepted.json()["report"]["version"] == 1
    assert client.get(base + "/facts/transit_plan/source").json()["document_id"] == doc_id
    assert client.post(base + f"/reports/{report['id']}/review").status_code == 200
    stale = client.post(path + "/commit", json={**{key: value for key, value in candidate.items()
        if key not in {"issues", "preserved_blocks"}}, "replace_section": True})
    assert stale.status_code == 409
    changed = client.post(base + "/changes/preview", json={"fact_key": "transit_plan",
        "value": "计划调整", "source": "新版原件", "reason": "更新"}).json()
    assert client.post(base + "/changes/commit", json={"fact_key": "transit_plan",
        "value": "计划调整", "source": "新版原件", "reason": "更新",
        "base_version": changed["base_version"], "preview_token": changed["preview_token"]}).status_code == 200
    after = client.get(base + f"/reports/{report['id']}").json()
    assert after["reviewed"] is False and any(issue["code"] == "FACT_CHANGED" for issue in after["issues"])
    assert client.post(base + f"/reports/{report['id']}/review").status_code != 200
    assert client.post(path + "/preview", json=request).status_code == 409


def test_narrative_model_guard_and_no_run(client, monkeypatch, tmp_path):
    import app.analysis_writing as writing
    base, config, report, statement, _ = setup_narrative(client, monkeypatch, tmp_path)
    path = base + f"/analysis/reports/{report['id']}/draft/preview"
    request = {"config_id": config["id"], "section_id": "background", "mode": "model"}
    monkeypatch.setattr(writing, "chat_json", lambda *_args, **_kwargs: ChatResult({
        "paragraphs": [{"text": "该枢纽已完成建设。", "used_fact_keys": ["transit_plan"]}]},
        {"model": "qa-only"}))
    rejected = client.post(path, json=request)
    assert rejected.status_code == 422
    assert client.get(base + f"/reports/{report['id']}").json()["version"] == 0
    monkeypatch.setattr(writing, "chat_json", lambda *_args, **_kwargs: ChatResult({
        "paragraphs": [{"text": "原文记载，马昌营公共交通枢纽站计划与地铁同期开通。",
                        "used_fact_keys": ["transit_plan"]}]}, {"model": "qa-only"}))
    valid = client.post(path, json=request)
    assert valid.status_code == 200, valid.text
    assert valid.json()["content"][1]["origin"] == "model"
    assert valid.json()["model_audit"]["model_call"]["model"] == "qa-only"


def test_negated_order_in_source_cannot_become_signed_order(client, monkeypatch, tmp_path):
    import app.analysis_writing as writing
    base, config, report, _, _ = setup_narrative(client, monkeypatch, tmp_path)
    monkeypatch.setattr(writing, "chat_json", lambda *_args, **_kwargs: ChatResult({
        "paragraphs": [{"text": "该项目计划与地铁同期开通，已签订单。", "used_fact_keys": ["transit_plan"]}]}, {}))
    response = client.post(base + f"/analysis/reports/{report['id']}/draft/preview", json={
        "config_id": config["id"], "section_id": "background", "mode": "model"})
    assert response.status_code == 422
    assert "拟议事项写成已完成" in response.json()["detail"]


def test_narrative_config_requires_evidence_and_no_results(client):
    base = f"/api/projects/{client.post('/api/projects', json={'name': '配置边界'}).json()['id']}"
    config = client.post(base + "/analysis/configs", json={"name": "文字章"}).json()
    body = {"revision": 1, "name": "文字章", "definitions": [{"key": "background", "label": "背景",
        "data_type": "text", "unit": "", "computed": False}], "rules": [], "sections": [
        {"id": "background", "title": "背景", "kind": "narrative", "result_keys": [], "evidence_keys": []}]}
    path = base + f"/analysis/configs/{config['id']}"
    assert client.put(path, json=body).status_code == 400
    body["sections"][0]["evidence_keys"] = ["background"]
    body["sections"][0]["result_keys"] = ["background"]
    assert client.put(path, json=body).status_code == 400


def test_two_project_sources_and_immutable_replacement(client, monkeypatch, tmp_path):
    base, config, report, statement, first_doc = setup_narrative(client, monkeypatch, tmp_path)
    another = Document(); another.add_paragraph("补充说明：" + statement)
    stream = io.BytesIO(); another.save(stream)
    second = client.post(base + "/documents", files={"file": (
        "补充原件.docx", stream.getvalue(),
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document")})
    assert second.status_code == 201, second.text
    second_doc = second.json()["id"]
    first_statement = {"document_id": first_doc, "source_refs": ["d-p1"],
                       "statement": statement, "label": "接驳计划", "fact_key": "transit_plan"}
    original = client.post(base + "/evidence", json=first_statement)
    assert original.status_code == 201, original.text
    second_statement = {**first_statement, "document_id": second_doc}
    supporting = client.post(base + "/evidence", json=second_statement)
    assert supporting.status_code == 201, supporting.text
    first_id, second_id = original.json()["id"], supporting.json()["id"]
    listed = client.get(base + "/evidence?fact_key=transit_plan").json()["items"]
    assert {row["id"] for row in listed} == {first_id, second_id}
    assert all(row["independent_verification"] == "not_recorded" for row in listed)
    other = client.post("/api/projects", json={"name": "跨项目证据 QA"}).json()
    other_base = f"/api/projects/{other['id']}"
    assert client.get(other_base + f"/evidence/{first_id}").status_code == 404
    draft = base + f"/analysis/reports/{report['id']}/draft"
    cross = client.post(draft + "/preview", json={"config_id": config["id"],
        "section_id": "background", "mode": "excerpt", "evidence_ids": [first_id, "foreign"]})
    assert cross.status_code == 404
    candidate = client.post(draft + "/preview", json={"config_id": config["id"],
        "section_id": "background", "mode": "excerpt", "evidence_ids": [first_id, second_id]})
    assert candidate.status_code == 200, candidate.text
    body = candidate.json()
    assert {row["evidence_id"] for row in body["content"][1]["project_evidence_refs"]} == {first_id, second_id}
    adopted = client.post(draft + "/commit", json={**{key: value for key, value in body.items()
        if key not in {"issues", "preserved_blocks"}}, "replace_section": False})
    assert adopted.status_code == 200, adopted.text
    assert client.post(base + f"/reports/{report['id']}/review").status_code == 200
    preview = client.post(base + f"/evidence/{first_id}/replacement/preview", json=second_statement)
    assert preview.status_code == 200, preview.text
    assert len(preview.json()["impacts"]) == 1
    assert client.get(base + f"/evidence/{first_id}").json()["supersedes_id"] is None
    signed = {**second_statement, **{key: preview.json()[key] for key in
        ("base_report_versions", "expires_at", "preview_token")}}
    replaced = client.post(base + f"/evidence/{first_id}/replacement/commit", json=signed)
    assert replaced.status_code == 201, replaced.text
    assert replaced.json()["supersedes_id"] == first_id
    repeated = client.post(base + f"/evidence/{first_id}/replacement/commit", json=signed)
    assert repeated.status_code == 201 and repeated.json()["id"] == replaced.json()["id"]
    old = client.get(base + f"/evidence/{first_id}").json()
    assert old["document_id"] == first_doc and old["excerpt"] == statement
    assert old["original_url"].endswith(f"/documents/{first_doc}/original")
    report_after = client.get(base + f"/reports/{report['id']}").json()
    assert report_after["reviewed"] is False
    assert any(issue["code"] == "PROJECT_EVIDENCE_SUPERSEDED" for issue in report_after["issues"])
    assert client.post(base + f"/reports/{report['id']}/review").status_code != 200


def test_old_evidence_keeps_parse_snapshot_after_new_parse(client, monkeypatch, tmp_path):
    base, config, report, statement, document_id = setup_narrative(client, monkeypatch, tmp_path)
    registered = client.post(base + "/evidence", json={"document_id": document_id,
        "source_refs": ["d-p1"], "statement": statement, "label": "接驳计划",
        "fact_key": "transit_plan"})
    assert registered.status_code == 201, registered.text
    old = registered.json()
    assert old["parse_revision"] == 1 and old["parse_revision_id"]
    candidate = client.post(base + f"/analysis/reports/{report['id']}/draft/preview", json={
        "config_id": config["id"], "section_id": "background", "mode": "excerpt",
        "evidence_ids": [old["id"]]}).json()
    assert candidate["content"][1]["project_evidence_refs"][0]["parse_revision_id"] == old["parse_revision_id"]
    assert client.post(base + f"/analysis/reports/{report['id']}/draft/commit", json={
        **{key: value for key, value in candidate.items() if key not in {"issues", "preserved_blocks"}},
        "replace_section": False}).status_code == 200
    with SessionLocal.begin() as session:
        document = session.get(SourceDocument, document_id)
        document.parse_revision = 2
        document.segments = [{"ref": "d-p9", "page": 1, "text": statement}]
        session.add(SourceParseRevision(document_id=document_id, document_sha256=document.sha256,
                                        revision=2, segments=document.segments, status="READY"))
    revisions = client.get(base + f"/documents/{document_id}/parse-revisions").json()
    assert [row["revision"] for row in revisions] == [1, 2]
    first = client.get(base + f"/documents/{document_id}/parse-revisions/{old['parse_revision_id']}").json()
    assert first["segments"][0]["ref"] == "d-p1"
    assert client.get(base + f"/evidence/{old['id']}").json()["excerpt"] == statement
    assert not any(issue["code"] == "PROJECT_EVIDENCE_LOCATION_CHANGED"
                   for issue in client.get(base + f"/reports/{report['id']}").json()["issues"])
    assert client.post(base + f"/reports/{report['id']}/review").status_code == 200


def test_chapter_material_pack_has_30_honest_roles_and_stales_candidate(client, monkeypatch, tmp_path):
    base, config, report, statement, document_id = setup_narrative(client, monkeypatch, tmp_path)
    pack_url = base + f"/reports/{report['id']}/materials?section_id=background"
    pack = client.get(pack_url)
    assert pack.status_code == 200, pack.text
    before = pack.json()
    assert len(before["categories"]) == 30
    assert before["categories"][10]["status"] == "未生成"
    assert before["categories"][29]["status"] == "部分未生成"
    assert all(not row["used_records"] for row in before["categories"])
    assert before["facts"][0]["key"] == "transit_plan"
    assert not before["project_evidence"]
    assert client.get(base + f"/reports/{report['id']}/materials?section_id=unknown").status_code == 404
    other = client.post("/api/projects", json={"name": "他人项目"}).json()
    assert client.get(f"/api/projects/{other['id']}/reports/{report['id']}/materials?section_id=background").status_code == 404
    draft_url = base + f"/analysis/reports/{report['id']}/draft"
    candidate = client.post(draft_url + "/preview", json={"config_id": config["id"],
        "section_id": "background", "mode": "excerpt"}).json()
    assert candidate["model_audit"]["writing_pack_sha256"] == before["watermark_sha256"]
    added = client.post(base + "/evidence", json={"document_id": document_id,
        "source_refs": ["d-p1"], "statement": statement, "label": "原文证明",
        "fact_key": "transit_plan"})
    assert added.status_code == 201, added.text
    after = client.get(pack_url).json()
    assert len(after["project_evidence"]) == 1
    assert after["watermark_sha256"] != before["watermark_sha256"]
    stale = client.post(draft_url + "/commit", json={**{key: value for key, value in candidate.items()
        if key not in {"issues", "preserved_blocks"}}, "replace_section": False})
    assert stale.status_code == 409
    assert client.get(base + f"/reports/{report['id']}").json()["version"] == 0


def test_material_policy_version_and_project_source_boundary(client, monkeypatch, tmp_path):
    base, config, report, _, _ = setup_narrative(client, monkeypatch, tmp_path)
    ref = client.get(base + "/writing/reference").json()
    chosen = client.put(base + "/writing/reference", json={
        "corpus_id": ref["corpus_id"], "corpus_version": ref["corpus_version"],
        "target_report_type": "feasibility", "accepted": True})
    assert chosen.status_code == 200, chosen.text
    url = base + f"/reports/{report['id']}/materials"
    before = client.get(url + "?section_id=background").json()
    assert before["corpus_id"] == ref["corpus_id"]
    draft_url = base + f"/analysis/reports/{report['id']}/draft"
    candidate = client.post(draft_url + "/preview", json={
        "config_id": config["id"], "section_id": "background", "mode": "excerpt"}).json()
    body = {"section_id": "background", "base_revision": before["policy_revision"],
            "settings": [{"category_id": "CAT-09", "mode": "exclude"}]}
    saved = client.put(url, json=body)
    assert saved.status_code == 200, saved.text
    assert saved.json()["policy_revision"] == 1
    assert saved.json()["categories"][8]["mode"] == "exclude"
    assert client.put(url, json=body).status_code == 409
    stale = client.post(draft_url + "/commit", json={**{key: value for key, value in candidate.items()
        if key not in {"issues", "preserved_blocks"}}, "replace_section": False})
    assert stale.status_code == 409
    # CAT-09 is historical corpus text. Excluding it cannot hide this project's
    # independently registered source requirement or make an unlinked draft valid.
    fresh = client.post(draft_url + "/preview", json={
        "config_id": config["id"], "section_id": "background", "mode": "excerpt"})
    assert fresh.status_code == 200, fresh.text


def test_conflicting_sources_block_model_and_formal_review_until_new_resolution(client, monkeypatch, tmp_path):
    base, config, report, statement, first_doc = setup_narrative(client, monkeypatch, tmp_path)
    def upload(name, text):
        word = Document(); word.add_paragraph(text)
        stream = io.BytesIO(); word.save(stream)
        response = client.post(base + "/documents", files={"file": (
            name, stream.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")})
        assert response.status_code == 201, response.text
        return response.json()["id"]
    second_doc = upload("相反来源.docx", "另一来源称枢纽开通时间尚未确定。")
    first = client.post(base + "/evidence", json={"document_id": first_doc,
        "source_refs": ["d-p1"], "statement": statement, "label": "计划同期"}).json()
    second = client.post(base + "/evidence", json={"document_id": second_doc,
        "source_refs": ["d-p1"], "statement": "另一来源称枢纽开通时间尚未确定。",
        "label": "时间未定"}).json()
    issue_body = {"kind": "conflict", "title": "开通时点冲突", "statement": "两份来源对开通时点不一致",
                  "section_ids": ["background"], "support_evidence_ids": [first["id"]],
                  "opposing_evidence_ids": [second["id"]]}
    invalid = client.post(base + "/issues", json={**issue_body,
        "opposing_evidence_ids": [first["id"]]})
    assert invalid.status_code == 400
    issue = client.post(base + "/issues", json=issue_body)
    assert issue.status_code == 201, issue.text
    issue_id = issue.json()["id"]
    assert {row["statement"] for row in issue.json()["support"] + issue.json()["opposing"]} == {
        statement, "另一来源称枢纽开通时间尚未确定。"}
    model = client.post(base + f"/analysis/reports/{report['id']}/draft/preview", json={
        "config_id": config["id"], "section_id": "background", "mode": "model"})
    assert model.status_code == 409
    candidate = client.post(base + f"/analysis/reports/{report['id']}/draft/preview", json={
        "config_id": config["id"], "section_id": "background", "mode": "excerpt"}).json()
    adopted = client.post(base + f"/analysis/reports/{report['id']}/draft/commit", json={
        **{key: value for key, value in candidate.items() if key not in {"issues", "preserved_blocks"}},
        "replace_section": False})
    assert adopted.status_code == 200, adopted.text
    state = client.get(base + f"/reports/{report['id']}").json()
    assert any(row["code"] == "PROJECT_CONFLICT_OPEN" for row in state["issues"])
    assert client.post(base + f"/reports/{report['id']}/review").status_code == 400
    assert client.get(base + f"/reports/{report['id']}/export?level=formal").status_code == 409
    preview = client.get(base + f"/reports/{report['id']}/export?level=preview")
    assert preview.status_code == 200, preview.text[:300]
    illegal_resolution = client.post(base + f"/issues/{issue_id}/decision", json={
        "revision": 1, "status": "RESOLVED", "decision": "选来源A",
        "resolution_evidence_id": first["id"]})
    assert illegal_resolution.status_code == 409
    third_doc = upload("核对结果.docx", "经重新核对，枢纽仍按与地铁同期的方案推进，实际开通待确认。")
    third = client.post(base + "/evidence", json={"document_id": third_doc,
        "source_refs": ["d-p1"], "statement": "经重新核对，枢纽仍按与地铁同期的方案推进，实际开通待确认。",
        "label": "核对结果"}).json()
    resolved = client.post(base + f"/issues/{issue_id}/decision", json={
        "revision": 1, "status": "RESOLVED", "decision": "经重新核对，枢纽仍按与地铁同期的方案推进，实际开通待确认。",
        "resolution_evidence_id": third["id"]})
    assert resolved.status_code == 200, resolved.text
    assert resolved.json()["independent_verification"] == "not_recorded"
    assert client.post(base + f"/issues/{issue_id}/decision", json={
        "revision": 1, "status": "OPEN", "decision": "旧请求"}).status_code == 409
    assert client.post(base + f"/reports/{report['id']}/review").status_code == 200
    assert client.get(base + f"/reports/{report['id']}/export?level=formal").status_code == 409
    other = client.post("/api/projects", json={"name": "跨项目问题"}).json()
    assert client.get(f"/api/projects/{other['id']}/issues/{issue_id}").status_code == 404


def test_candidate_is_separate_restorable_declined_and_idempotently_adopted(client, monkeypatch, tmp_path):
    base, config, report, _, _ = setup_narrative(client, monkeypatch, tmp_path)
    path = base + f"/analysis/reports/{report['id']}/draft"
    request = {"config_id": config["id"], "section_id": "background", "mode": "excerpt"}
    first = client.post(path + "/preview", json=request).json()
    assert first["candidate_id"]
    listed = client.get(path + "/candidates?section_id=background").json()
    assert listed[0]["candidate_id"] == first["candidate_id"] and listed[0]["status"] == "READY"
    restored = client.get(path + "/candidates/" + first["candidate_id"])
    assert restored.status_code == 200 and restored.json()["content"] == first["content"]
    assert restored.json()["issues"] == first["issues"]
    assert restored.json()["preserved_blocks"] == first["preserved_blocks"]
    other = client.post("/api/projects", json={"name": "候选归属隔离"}).json()
    assert client.get(f"/api/projects/{other['id']}/analysis/reports/{report['id']}/draft/candidates/{first['candidate_id']}").status_code == 404
    cancelled = client.post(path + f"/candidates/{first['candidate_id']}/decline")
    assert cancelled.status_code == 200 and cancelled.json()["status"] == "DECLINED"
    assert client.post(path + f"/candidates/{first['candidate_id']}/decline").status_code == 200
    assert client.get(path + "/candidates/" + first["candidate_id"]).status_code == 409
    assert client.get(base + f"/reports/{report['id']}").json()["version"] == 0
    second = client.post(path + "/preview", json=request).json()
    body = {**{key: value for key, value in second.items()
               if key not in {"issues", "preserved_blocks"}}, "replace_section": False}
    adopted = client.post(path + "/commit", json=body)
    assert adopted.status_code == 200, adopted.text
    assert adopted.json()["report"]["version"] == 1
    again = client.post(path + "/commit", json=body)
    assert again.status_code == 200 and again.json()["idempotent"] is True
    assert again.json()["report"]["version"] == 1
    tampered = client.post(path + "/commit", json={**body,
        "content": [{**body["content"][0], "children": [{"text": "被篡改"}]}, *body["content"][1:]]})
    assert tampered.status_code == 409


def test_editing_sourced_paragraph_requires_explicit_local_source_review(client, monkeypatch, tmp_path):
    base, config, report, statement, document_id = setup_narrative(client, monkeypatch, tmp_path)
    evidence = client.post(base + "/evidence", json={"document_id": document_id,
        "source_refs": ["d-p1"], "statement": statement, "label": "接驳计划",
        "fact_key": "transit_plan"})
    assert evidence.status_code == 201, evidence.text
    path = base + f"/analysis/reports/{report['id']}/draft"
    candidate = client.post(path + "/preview", json={
        "config_id": config["id"], "section_id": "background", "mode": "excerpt",
        "evidence_ids": [evidence.json()["id"]]}).json()
    committed = client.post(path + "/commit", json={
        **{key: value for key, value in candidate.items()
           if key not in {"issues", "preserved_blocks"}}, "replace_section": False})
    assert committed.status_code == 200, committed.text
    before = committed.json()["report"]
    content = [dict(block) for block in before["content"]]
    sourced = next(block for block in content if block.get("project_evidence_refs"))
    sourced["children"] = [{"text": "据原报告记载，相关客户预测需求仍须核实。"}]
    preview_url = base + f"/reports/{report['id']}/preview"
    save_url = base + f"/reports/{report['id']}"
    rejected = client.post(preview_url, json={"base_version": before["version"], "content": content})
    assert rejected.status_code == 409
    assert client.get(save_url).json()["version"] == before["version"]
    sourced["source_review_required"] = True
    preview = client.post(preview_url, json={"base_version": before["version"], "content": content})
    assert preview.status_code == 200, preview.text
    saved = client.put(save_url, json={"base_version": before["version"], "content": content,
        "preview_token": preview.json()["preview_token"]})
    assert saved.status_code == 200, saved.text
    assert any(issue["code"] == "SOURCE_REVIEW_REQUIRED" for issue in saved.json()["report"]["issues"])
    assert client.post(save_url + "/review").status_code == 400
    version = saved.json()["report"]["version"]
    reviewed = client.post(save_url + f"/blocks/{sourced['id']}/source-review",
        json={"base_version": version})
    assert reviewed.status_code == 200, reviewed.text
    assert reviewed.json()["report"]["version"] == version + 1
    assert not any(issue["code"] == "SOURCE_REVIEW_REQUIRED" for issue in reviewed.json()["report"]["issues"])
    assert client.post(save_url + f"/blocks/{sourced['id']}/source-review",
        json={"base_version": version}).status_code == 409


def test_durable_task_cancel_retry_candidate_and_export(client, monkeypatch, tmp_path):
    import app.analysis_writing as writing
    import app.work_tasks as tasks
    base, config, report, _, _ = setup_narrative(client, monkeypatch, tmp_path)
    monkeypatch.setattr(tasks.executor, "submit", lambda *_args, **_kwargs: None)
    monkeypatch.setattr(writing, "chat_json", lambda *_args, **_kwargs: ChatResult({
        "paragraphs": [{"text": "原文记载，马昌营公共交通枢纽站计划与地铁同期开通。",
                        "used_fact_keys": ["transit_plan"]}]}, {"model": "qa-only"}))
    endpoint = base + "/tasks"
    payload = {"kind": "model_draft", "report_id": report["id"],
               "request_key": "model-chapter-0001", "config_id": config["id"],
               "section_id": "background"}
    queued = client.post(endpoint, json=payload)
    assert queued.status_code == 201, queued.text
    task_id = queued.json()["id"]
    assert client.post(endpoint, json=payload).json()["id"] == task_id
    assert client.post(endpoint, json={**payload, "section_id": "other"}).status_code == 404
    assert client.post(endpoint + f"/{task_id}/cancel").json()["status"] == "CANCELLED"
    assert tasks.process_one(task_id) is False
    assert client.get(base + f"/reports/{report['id']}").json()["version"] == 0
    assert client.post(endpoint + f"/{task_id}/retry").json()["status"] == "PENDING"
    assert tasks.process_one(task_id) is True
    complete = client.get(endpoint + f"/{task_id}").json()
    assert complete["status"] == "COMPLETED" and complete["result"]["candidate_id"]
    assert client.get(base + f"/analysis/reports/{report['id']}/draft/candidates/" +
                      complete["result"]["candidate_id"]).status_code == 200
    assert client.get(base + f"/reports/{report['id']}").json()["version"] == 0
    assert client.post(endpoint + f"/{task_id}/retry").status_code == 409
    other = client.post('/api/projects', json={'name': '任务隔离'}).json()
    assert client.get(f"/api/projects/{other['id']}/tasks/{task_id}").status_code == 404

    candidate = client.get(base + f"/analysis/reports/{report['id']}/draft/candidates/" +
                           complete["result"]["candidate_id"]).json()
    adopted = client.post(base + f"/analysis/reports/{report['id']}/draft/commit", json={
        **{key: value for key, value in candidate.items() if key not in {"issues", "preserved_blocks"}},
        "replace_section": False})
    assert adopted.status_code == 200, adopted.text
    export_task = client.post(endpoint, json={"kind": "export", "report_id": report["id"],
        "request_key": "preview-export-0001", "level": "preview"})
    assert export_task.status_code == 201, export_task.text
    assert tasks.process_one(export_task.json()["id"]) is True
    finished = client.get(endpoint + f"/{export_task.json()['id']}").json()
    assert finished["status"] == "COMPLETED" and finished["result"]["export_id"]
    download = client.get(base + f"/reports/{report['id']}/exports/{finished['result']['export_id']}")
    assert download.status_code == 200 and download.headers["x-archive-sha256"] == finished["result"]["sha256"]


def test_cancelled_running_model_cannot_create_late_candidate(client, monkeypatch, tmp_path):
    import threading
    import app.analysis_writing as writing
    import app.work_tasks as tasks
    base, config, report, _, _ = setup_narrative(client, monkeypatch, tmp_path)
    monkeypatch.setattr(tasks.executor, "submit", lambda *_args, **_kwargs: None)
    started, proceed = threading.Event(), threading.Event()

    def delayed_model(*_args, **_kwargs):
        started.set()
        assert proceed.wait(5)
        return ChatResult({"paragraphs": [{
            "text": "原文记载，马昌营公共交通枢纽站计划与地铁同期开通。",
            "used_fact_keys": ["transit_plan"]}]}, {"model": "qa-only"})

    monkeypatch.setattr(writing, "chat_json", delayed_model)
    endpoint = base + "/tasks"
    created = client.post(endpoint, json={"kind": "model_draft", "report_id": report["id"],
        "request_key": "cancel-running-0001", "config_id": config["id"],
        "section_id": "background"}).json()
    thread = threading.Thread(target=tasks.process_one, args=(created["id"],))
    thread.start()
    try:
        assert started.wait(5)
        assert client.post(endpoint + f"/{created['id']}/cancel").json()["status"] == "CANCELLED"
    finally:
        proceed.set()
        thread.join(8)
    assert not thread.is_alive()
    assert client.get(endpoint + f"/{created['id']}").json()["status"] == "CANCELLED"
    assert client.get(base + f"/analysis/reports/{report['id']}/draft/candidates").json() == []
    assert client.get(base + f"/reports/{report['id']}").json()["version"] == 0


def test_expired_worker_is_uncertain_and_old_attempt_cannot_commit(client, monkeypatch, tmp_path):
    from datetime import timedelta
    import app.work_tasks as tasks
    import app.analysis_writing as writing
    from app.db import WorkTask, utcnow
    base, config, report, _, _ = setup_narrative(client, monkeypatch, tmp_path)
    monkeypatch.setattr(tasks.executor, "submit", lambda *_args, **_kwargs: None)
    endpoint = base + "/tasks"
    created = client.post(endpoint, json={"kind": "model_draft", "report_id": report["id"],
        "request_key": "expired-attempt-0001", "config_id": config["id"],
        "section_id": "background"}).json()
    with SessionLocal.begin() as session:
        row = session.get(WorkTask, created["id"])
        row.status, row.stage = "RUNNING", "CALLING_MODEL"
        row.generation = "old-attempt"
        row.lease_expires_at = utcnow() - timedelta(seconds=1)
    assert tasks.recover_expired() == 1
    assert client.get(endpoint + f"/{created['id']}").json()["status"] == "UNCERTAIN"
    monkeypatch.setattr(writing, "chat_json", lambda *_args, **_kwargs: ChatResult({
        "paragraphs": [{"text": "原文记载，马昌营公共交通枢纽站计划与地铁同期开通。",
                        "used_fact_keys": ["transit_plan"]}]}, {"model": "qa-only"}))
    with pytest.raises(Exception):
        writing._draft_preview(base.rsplit("/", 1)[-1], report["id"],
            writing.DraftRequest(section_id="background", mode="model", config_id=config["id"]),
            (created["id"], "old-attempt"))
    assert client.get(base + f"/analysis/reports/{report['id']}/draft/candidates").json() == []
    assert client.post(endpoint + f"/{created['id']}/retry").json()["status"] == "PENDING"
    assert tasks.process_one(created["id"]) is True
    assert client.get(endpoint + f"/{created['id']}").json()["status"] == "COMPLETED"


def test_cancelled_running_export_does_not_publish_archive(client, monkeypatch, tmp_path):
    import threading
    import app.main as main
    import app.work_tasks as tasks
    base, config, report, _, _ = setup_narrative(client, monkeypatch, tmp_path)
    draft_path = base + f"/analysis/reports/{report['id']}/draft"
    candidate = client.post(draft_path + "/preview", json={
        "config_id": config["id"], "section_id": "background", "mode": "excerpt"}).json()
    assert client.post(draft_path + "/commit", json={
        **{key: value for key, value in candidate.items() if key not in {"issues", "preserved_blocks"}},
        "replace_section": False}).status_code == 200
    monkeypatch.setattr(tasks.executor, "submit", lambda *_args, **_kwargs: None)
    started, proceed = threading.Event(), threading.Event()
    original_render = main.export_bundle

    def delayed_render(*args, **kwargs):
        started.set()
        assert proceed.wait(10)
        return original_render(*args, **kwargs)

    monkeypatch.setattr(main, "export_bundle", delayed_render)
    endpoint = base + "/tasks"
    created = client.post(endpoint, json={"kind": "export", "report_id": report["id"],
        "request_key": "cancel-export-0001", "level": "preview"}).json()
    worker = threading.Thread(target=tasks.process_one, args=(created["id"],))
    worker.start()
    try:
        assert started.wait(5)
        assert client.post(endpoint + f"/{created['id']}/cancel").json()["status"] == "CANCELLED"
    finally:
        proceed.set()
        worker.join(10)
    assert not worker.is_alive()
    assert client.get(endpoint + f"/{created['id']}").json()["status"] == "CANCELLED"
    assert client.get(base + f"/reports/{report['id']}/exports").json() == []


def test_living_task_renews_lease_and_cancel_stops_heartbeat(client, monkeypatch, tmp_path):
    from datetime import timedelta
    from threading import Event, Thread
    import time
    import app.work_tasks as tasks
    from app.db import WorkTask, utcnow
    base, config, report, _, _ = setup_narrative(client, monkeypatch, tmp_path)
    monkeypatch.setattr(tasks.executor, "submit", lambda *_args, **_kwargs: None)
    monkeypatch.setattr(tasks, "HEARTBEAT_SECONDS", 0.02)
    endpoint = base + "/tasks"
    row = client.post(endpoint, json={"kind": "model_draft", "report_id": report["id"],
        "request_key": "heartbeat-0001", "config_id": config["id"],
        "section_id": "background"}).json()
    with SessionLocal.begin() as session:
        stored = session.get(WorkTask, row["id"])
        stored.status, stored.generation = "RUNNING", "live-attempt"
        stored.lease_expires_at = utcnow() + timedelta(seconds=1)
    stop = Event()
    worker = Thread(target=tasks._heartbeat, args=(row["id"], "live-attempt", stop))
    worker.start()
    try:
        deadline = time.monotonic() + 2
        while time.monotonic() < deadline:
            with SessionLocal() as session:
                expires = session.get(WorkTask, row["id"]).lease_expires_at
            if expires > utcnow() + timedelta(seconds=30):
                break
            time.sleep(0.02)
        else:
            pytest.fail("活跃任务没有续租")
        assert client.post(endpoint + f"/{row['id']}/cancel").json()["status"] == "CANCELLED"
    finally:
        stop.set()
        worker.join(2)
    assert not worker.is_alive()
