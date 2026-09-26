"""Approved Plate nodes must survive validation, comparison and both exports."""

from io import BytesIO
from types import SimpleNamespace
from zipfile import ZipFile

import pymupdf
import pytest
from docx import Document

from app.report_pipeline import change_impact, content_hash, export_bundle, gate, plain, report_fact_impacts, source_display, validate_content


def sample_content():
    return [
        {"type": "h1", "id": "heading-1", "children": [{"text": "项目报告"}]},
        {"type": "p", "id": "para-1", "listStyleType": "disc", "indent": 1,
         "children": [{"text": "已核实", "underline": True, "strikethrough": True}, {"type": "fact_ref", "fact_key": "land_area",
                                             "display": "建设用地面积 54088.771 平方米", "children": [{"text": ""}]}],
         "fact_keys": ["land_area"]},
        {"type": "p", "id": "para-2", "listStyleType": "decimal", "indent": 1,
         "children": [{"text": "参阅"}, {"type": "a", "url": "https://example.org/report", "children": [{"text": "来源报告", "underline": True}]}]},
        {"type": "table", "id": "table-1", "fact_keys": ["land_area"], "children": [
            {"type": "tr", "children": [
                {"type": "th", "children": [{"type": "p", "children": [{"text": "项目"}]}]},
                {"type": "th", "children": [{"type": "p", "children": [{"text": "数值"}]}]},
            ]},
            {"type": "tr", "children": [
                {"type": "td", "children": [{"type": "p", "children": [{"text": "建设用地面积"}]}]},
                {"type": "td", "children": [{"type": "p", "children": [
                    {"type": "fact_ref", "fact_key": "land_area", "display": "54088.771 平方米", "children": [{"text": ""}]}
                ]}]},
            ]},
        ]},
    ]


def test_plate_nodes_roundtrip_through_exports():
    content = sample_content()
    validate_content(content)
    fact = SimpleNamespace(value_text="54088.771", revision=1, value_status="PROVIDED", data_type="decimal")
    assert gate(content, content_hash(content), {"land_area": {"value": "54088.771", "revision": 1}}, {"land_area": fact}) == []
    assert plain(content[3]) == "项目\t数值\n建设用地面积\t54088.771 平方米"
    assert change_impact(content, [{**content[0], "children": [{"text": "新标题"}]}, *content[1:]])[0]["position"] == 1
    audit = {"facts": [{"key": "land_area", "label": "建设用地面积", "value": "54088.771", "unit": "平方米",
                        "source": "document:doc-id#p7-t1-r6+p7-s2 · example.pdf", "revision": 1}]}
    with ZipFile(BytesIO(export_bundle("项目报告", content, audit))) as bundle:
        word_bytes = bundle.read("report.docx")
        word = Document(BytesIO(word_bytes))
        assert any(paragraph.style.name == "List Bullet" and "54088.771" in paragraph.text for paragraph in word.paragraphs)
        assert any(paragraph.style.name == "List Number" for paragraph in word.paragraphs)
        assert len(word.tables) == 1
        assert word.tables[0].cell(1, 1).text == "54088.771 平方米"
        formatted = next(paragraph for paragraph in word.paragraphs if paragraph.text.startswith("已核实"))
        assert formatted.runs[0].underline is True and formatted.runs[0].font.strike is True
        assert any(rel.target_ref == "https://example.org/report" for rel in word.part.rels.values())
        with pymupdf.open(stream=bundle.read("report.pdf"), filetype="pdf") as pdf:
            pdf_text = "".join(page.get_text() for page in pdf)
            assert any(link.get("uri") == "https://example.org/report" for page in pdf for link in page.get_links())
        for value in ("建设用地面积", "54088.771", "来源报告", "项目"):
            assert value in pdf_text
        assert "example.pdf，第 7 页，表 1，第 6 行" in pdf_text
    assert source_display(audit["facts"][0]["source"]) == "example.pdf，第 7 页，表 1，第 6 行"


def test_feasibility_export_shows_cited_inputs_and_calculation_in_both_files():
    content = [{"type": "p", "id": "area-result", "children": [{"text": "分项面积17268.25㎡，差额0㎡。"}],
                "analysis_refs": [{"run_id": "run-area", "result_key": "area_difference", "value": "0", "unit": "㎡"}]}]
    audit = {"facts": [], "analysis_refs": [{"block_id": "area-result", "refs": content[0]["analysis_refs"]}],
             "analysis_run": {
                 "definitions": [{"key": key} for key in ("aboveground_area", "underground_area", "reported_total_area", "calculated_total_area", "area_difference", "total_investment")],
                 "inputs": {"aboveground_area": {"value": "15756.25", "origin": "scenario_assumption"},
                            "underground_area": {"value": "1512", "origin": "project_fact"},
                            "reported_total_area": {"value": "17268.25", "origin": "project_fact"},
                            "total_investment": {"value": "31255.57", "origin": "project_fact"}},
                 "results": {key: {"label": label, "value": value, "unit": unit} for key, label, value, unit in (
                     ("aboveground_area", "地上建筑面积", "15756.25", "㎡"),
                     ("underground_area", "地下建筑面积", "1512", "㎡"),
                     ("reported_total_area", "原文总建筑面积", "17268.25", "㎡"),
                     ("calculated_total_area", "分项计算建筑面积", "17268.25", "㎡"),
                     ("area_difference", "原文与分项差额", "0", "㎡"),
                     ("total_investment", "原文总投资", "31255.57", "万元"))},
                 "trace": [
                     {"target": "calculated_total_area", "status": "COMPUTED", "expression": "aboveground_area + underground_area", "inputs": {"aboveground_area": "15756.25", "underground_area": "1512"}},
                     {"target": "area_difference", "status": "COMPUTED", "expression": "reported_total_area - calculated_total_area", "inputs": {"reported_total_area": "17268.25", "calculated_total_area": "17268.25"}},
                 ]}}
    with ZipFile(BytesIO(export_bundle("可研面积核对", content, audit))) as bundle:
        word = Document(BytesIO(bundle.read("report.docx")))
        word_text = "\n".join(paragraph.text for paragraph in word.paragraphs)
        with pymupdf.open(stream=bundle.read("report.pdf"), filetype="pdf") as pdf:
            pdf_text = "\n".join(page.get_text() for page in pdf)
    for text in ("推演依据", "地上建筑面积 15756.25㎡（方案假设）", "地下建筑面积 1512㎡（项目事实快照）",
                 "计算：分项计算建筑面积 17268.25㎡", "规则：地上建筑面积 + 地下建筑面积", "计算：原文与分项差额 0㎡"):
        assert text in word_text
        assert text in pdf_text
    assert "原文总投资 31255.57" not in word_text


def test_inserted_blocks_do_not_make_unchanged_paragraphs_look_rewritten():
    before = sample_content()
    inserted = {"type": "p", "id": "new-paragraph", "children": [{"text": "新增段落"}]}
    after = [before[0], inserted, *before[1:]]
    changes = change_impact(before, after)
    assert changes == [{"position": 2, "before": "", "after": "新增段落", "fact_keys": []}]


def test_table_cell_lists_can_be_saved_and_exported():
    content = sample_content()
    content[3]["children"][1]["children"][0]["children"][0].update(
        listStyleType="disc", indent=1,
    )
    content[3]["children"][1]["children"][1]["children"][0].update(
        listStyleType="decimal", listStart=2, indent=1,
    )
    validate_content(content)
    audit = {"facts": [{"key": "land_area", "label": "建设用地面积", "value": "54088.771",
                        "unit": "平方米", "source": "人工录入", "revision": 1}]}
    with ZipFile(BytesIO(export_bundle("项目报告", content, audit))) as bundle:
        word = Document(BytesIO(bundle.read("report.docx")))
        assert word.tables[0].cell(1, 0).paragraphs[0].style.name == "List Bullet"
        assert word.tables[0].cell(1, 1).paragraphs[0].style.name == "List Number"
        with pymupdf.open(stream=bundle.read("report.pdf"), filetype="pdf") as pdf:
            rendered = "".join(page.get_text() for page in pdf)
            assert "2. 54088.771" in rendered


@pytest.mark.parametrize("change", [
    lambda c: c[2]["children"][1].update(url="javascript:alert(1)"),
    lambda c: c[1]["children"][1].update(children=[{"text": "可改写"}]),
    lambda c: c[3]["children"][1]["children"].pop(),
    lambda c: c[1].update(listStyleType="task"),
])
def test_invalid_or_unsupported_nodes_are_rejected(change):
    content = sample_content()
    change(content)
    with pytest.raises(ValueError):
        validate_content(content)


def test_model_draft_uses_configured_writing_route_and_rejects_new_numbers(monkeypatch):
    from app import report_pipeline

    observed = []
    def call(task, messages, max_tokens):
        observed.append((task, max_tokens, messages[1]["content"]))
        return {"paragraphs": [{"text": "本期产量为0套。", "fact_keys": ["capacity"]}]}
    monkeypatch.setattr(report_pipeline, "chat_json", call)
    facts = [{"key": "capacity", "label": "本期产量", "value": "0", "unit": "套", "source": "人工录入"}]
    drafted = report_pipeline.model_section("产量", facts)
    assert observed[0][:2] == ("writing", 1200)
    assert drafted[0]["origin"] == "model"
    monkeypatch.setattr(report_pipeline, "chat_json", lambda *args, **kwargs: {"paragraphs": [
        {"text": "本期产量为1套。", "fact_keys": ["capacity"]}]})
    with pytest.raises(ValueError, match="未提供的数字"):
        report_pipeline.model_section("产量", facts)


def test_model_fact_tokens_use_project_values_with_grouped_display(monkeypatch):
    from app import report_pipeline
    from app.writing_support import _model_fact_references

    facts = [{"key": "demand", "label": "首年需求", "value": "300000", "unit": "套", "source": "项目资料"},
             {"key": "capacity", "label": "合格能力", "value": "254016", "unit": "套", "source": "项目资料"},
             {"key": "sales", "label": "计划销售量", "value": "254016", "unit": "套", "source": "项目计算"}]
    text = "需求300,000套，能力254,016套；计划销售量254,016套。"
    monkeypatch.setattr(report_pipeline, "chat_json", lambda *args, **kwargs: {
        "paragraphs": [{"text": text, "fact_keys": ["demand", "capacity", "sales"]}]})
    assert report_pipeline.model_section("首年销售", facts)[0]["children"][0]["text"] == text
    state = {row["key"]: {"value": row["value"], "data_type": "integer"} for row in facts}
    assert _model_fact_references(text, ["demand", "capacity", "sales"], state) == [
        {"fact_key": "demand", "display": "300,000"},
        {"fact_key": "capacity", "display": "254,016"},
        {"fact_key": "sales", "display": "254,016"},
    ]
    with pytest.raises(ValueError, match="可定位值"):
        _model_fact_references("需求300,000套，能力254,016套。", ["demand", "capacity", "sales"], state)
    with pytest.raises(ValueError, match="可定位值"):
        _model_fact_references("需求180,000套。", ["demand"], state)


def test_changed_fact_requires_visible_token_update_before_review():
    content = [{"type": "p", "id": "fact-paragraph", "fact_keys": ["demand"], "children": [
        {"text": "首年需求为 "},
        {"type": "fact_ref", "fact_key": "demand", "display": "首年需求 300000套", "children": [{"text": ""}]},
        {"text": "。"},
    ]}]
    previous = {"demand": {"value": "300000", "revision": 1}}
    current = {"demand": SimpleNamespace(value_text="320000", revision=2, value_status="PROVIDED", data_type="integer")}
    impacts = report_fact_impacts(content, previous, current)
    assert [(item["position"], item["fact_key"], item["before"]["value"], item["after"]["value"])
            for item in impacts] == [(1, "demand", "300000", "320000")]
    assert {item["code"] for item in gate(content, content_hash(content), previous, current)} >= {"FACT_CHANGED", "FACT_TEXT_STALE"}
    updated = [{**content[0], "children": [content[0]["children"][0],
                 {**content[0]["children"][1], "display": "首年需求 320000套"}, content[0]["children"][2]]}]
    assert "FACT_TEXT_STALE" not in {item["code"] for item in gate(updated, content_hash(updated), previous, current)}
    current_snapshot = {"demand": {"value": "320000", "revision": 2}}
    assert gate(updated, content_hash(updated), current_snapshot, current) == []
