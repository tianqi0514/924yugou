"""Synthetic long document checks layout semantics without touching the database."""

import io
import json
from zipfile import ZipFile

import pymupdf
from docx import Document

from app.report_pipeline import export_bundle


def paragraph(kind: str, text: str) -> dict:
    return {"type": kind, "children": [{"text": text}]}


def test_multichapter_wide_table_has_directory_pages_and_repeating_header():
    headings = [paragraph("h2", "项目概况"), paragraph("h2", "建设方案")]
    headings[0]["children"].append({"text": ""})
    header = {"type": "tr", "children": [
        {"type": "th", "children": [paragraph("p", f"栏目{i}")]} for i in range(6)]}
    rows = [{"type": "tr", "children": [
        {"type": "td", "children": [paragraph("p", f"第{row}行第{column}列")]} for column in range(6)]}
        for row in range(72)]
    content = [paragraph("h1", "长文版式检验"), headings[0], paragraph("p", "本章概况。"),
               headings[1], {"type": "table", "children": [header, *rows]}]
    raw = export_bundle("长文版式检验", content, {"facts": []})
    with ZipFile(io.BytesIO(raw)) as archive:
        word_bytes = archive.read("report.docx")
        pdf_bytes = archive.read("report.pdf")
        audit = json.loads(archive.read("audit.json"))
    word = Document(io.BytesIO(word_bytes))
    assert word.sections[0].page_width > word.sections[0].page_height
    assert [entry.text for entry in word.paragraphs[:4]] == ["长文版式检验", "章节目录", "项目概况", "建设方案"]
    assert len(word.tables[0].rows) == 73
    with ZipFile(io.BytesIO(word_bytes)) as archive:
        xml = archive.read("word/document.xml").decode()
        footer = archive.read("word/footer1.xml").decode()
    assert "w:tblHeader" in xml
    assert 'w:instr="PAGE"' in footer
    pdf = pymupdf.open(stream=pdf_bytes, filetype="pdf")
    assert len(pdf) >= 2
    assert all(page.rect.width > page.rect.height for page in pdf)
    assert "章节目录" in pdf[0].get_text()
    assert "第71行第5列" in "".join(page.get_text() for page in pdf)
    assert all("栏目0" in page.get_text() for page in pdf if "第" in page.get_text() and "列" in page.get_text())
    assert audit["docx_sha256"] and audit["pdf_sha256"]
