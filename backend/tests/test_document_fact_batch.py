"""Document-side fact entry is one previewed, project-owned transaction."""

import io
import os
from pathlib import Path

os.environ["DATABASE_URL"] = "postgresql+psycopg://report:local_development_only@127.0.0.1:55432/report_platform_test"

import pytest
from docx import Document
from fastapi.testclient import TestClient

from app.db import AnalysisConfig, Base, SessionLocal, engine
from app.main import app


@pytest.fixture(autouse=True)
def isolated_database():
    assert engine.url.database == "report_platform_test"
    writable = [table for table in Base.metadata.sorted_tables if not table.name.startswith("corpus_")]
    Base.metadata.drop_all(engine, tables=writable)
    Base.metadata.create_all(engine, tables=writable)
    yield
    Base.metadata.drop_all(engine, tables=writable)


@pytest.fixture
def client(monkeypatch, tmp_path):
    import app.main as main
    monkeypatch.setattr(main, "STORAGE", tmp_path)
    with TestClient(app) as value:
        yield value


def _source(client, project_id):
    word = Document()
    word.add_paragraph("本项目首年需求为 0 套。")
    word.add_paragraph("本项目工期为 12 个月。")
    stream = io.BytesIO(); word.save(stream)
    response = client.post(f"/api/projects/{project_id}/documents", files={"file": (
        "输入.docx", stream.getvalue(),
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document")})
    assert response.status_code == 201, response.text
    return response.json()["id"]


def _batch(document_id):
    return {"document_id": document_id, "items": [
        {"key": "demand", "label": "首年需求", "data_type": "integer", "unit": "套",
         "value": "0", "source_ref": "d-p1"},
        {"key": "duration", "label": "工期", "data_type": "integer", "unit": "个月",
         "value": "12", "source_ref": "d-p2"},
    ]}


def test_preview_cancel_then_atomic_commit_and_source_roundtrip(client):
    project_id = client.post("/api/projects", json={"name": "原件旁录入"}).json()["id"]
    base = f"/api/projects/{project_id}"
    batch = _batch(_source(client, project_id))
    before = client.get(base + "/facts").json()
    preview = client.post(base + "/facts/from-document/preview", json=batch)
    assert preview.status_code == 200, preview.text
    result = preview.json()
    assert [(row["value"], row["status"]) for row in result["items"]] == [
        ("0", "原文值匹配，待确认"), ("12", "原文值匹配，待确认")]
    assert client.get(base + "/facts").json() == before  # Cancel: do not call commit.
    committed = client.post(base + "/facts/from-document/commit", json={
        **batch, "base_version": result["base_version"], "preview_token": result["preview_token"]})
    assert committed.status_code == 201, committed.text
    assert client.get(base + "/facts").json()["project_version"] == before["project_version"] + 1
    facts = {row["key"]: row for row in client.get(base + "/facts").json()["facts"]}
    assert facts["demand"]["value"] == "0" and facts["demand"]["evidence_status"] == "SOURCE_LOCATOR_REVIEWED"
    assert facts["duration"]["value"] == "12"
    source = client.get(base + "/facts/demand/source").json()
    assert source["document_id"] == batch["document_id"] and source["source_ref"] == "d-p1"
    evidence = client.get(base + "/evidence?fact_key=demand").json()["items"]
    assert len(evidence) == 1 and evidence[0]["independent_verification"] == "not_recorded"
    assert client.post(base + "/facts/from-document/commit", json={
        **batch, "base_version": result["base_version"], "preview_token": result["preview_token"]}).status_code == 409


def test_unsupported_item_and_stale_or_cross_project_source_never_half_commit(client):
    first = client.post("/api/projects", json={"name": "甲"}).json()["id"]
    second = client.post("/api/projects", json={"name": "乙"}).json()["id"]
    base = f"/api/projects/{first}"
    batch = _batch(_source(client, first))
    wrong = {**batch, "items": [batch["items"][0], {**batch["items"][1], "value": "99"}]}
    assert client.post(base + "/facts/from-document/preview", json=wrong).status_code == 409
    assert client.get(base + "/facts").json()["facts"] == []
    assert client.post(f"/api/projects/{second}/facts/from-document/preview", json=batch).status_code == 404
    preview = client.post(base + "/facts/from-document/preview", json=batch).json()
    assert client.post(base + "/facts", json={"key": "other", "label": "其他", "data_type": "integer"}).status_code == 201
    assert client.post(base + "/facts/from-document/commit", json={
        **batch, "base_version": preview["base_version"], "preview_token": preview["preview_token"]}).status_code == 409
    assert {row["key"] for row in client.get(base + "/facts").json()["facts"]} == {"other"}


def test_empty_value_creates_undefined_fact_without_false_source_review(client):
    project_id = client.post("/api/projects", json={"name": "缺值"}).json()["id"]
    base = f"/api/projects/{project_id}"
    batch = _batch(_source(client, project_id))
    batch["items"] = [{**batch["items"][0], "value": None}]
    preview = client.post(base + "/facts/from-document/preview", json=batch).json()
    assert preview["items"][0]["status"] == "未定义"
    committed = client.post(base + "/facts/from-document/commit", json={
        **batch, "base_version": preview["base_version"], "preview_token": preview["preview_token"]})
    assert committed.status_code == 201, committed.text
    fact = client.get(base + "/facts").json()["facts"][0]
    assert fact["value"] is None and fact["status"] == "UNDEFINED" and fact["evidence_status"] is None
    assert client.get(base + "/evidence?fact_key=demand").json()["items"] == []


def test_chapter_preparation_points_to_missing_fact_source_and_open_claim(client):
    project_id = client.post("/api/projects", json={"name": "章节准备"}).json()["id"]
    base = f"/api/projects/{project_id}"
    report = client.post(base + "/reports", json={
        "title": "政府可研", "report_type": "government_feasibility"}).json()
    endpoint = base + f"/reports/{report['id']}/chapters"
    empty = next(row for row in client.get(endpoint).json()["chapters"] if row["section_id"] == "necessity")
    assert empty["status"] == "待配置" and empty["next_actions"][0]["action"] == "config"
    with SessionLocal.begin() as session:
        session.add(AnalysisConfig(project_id=project_id, version=1, status="PUBLISHED", name="准备案例",
            definitions=[{"key": "demand", "label": "首年需求", "data_type": "integer", "unit": "套",
                          "group": "需求", "computed": False}], rules=[],
            sections=[{"id": "necessity", "title": "建设必要性", "kind": "narrative",
                       "result_keys": [], "evidence_keys": ["demand"]}]))
    missing = next(row for row in client.get(endpoint).json()["chapters"] if row["section_id"] == "necessity")
    assert missing["status"] == "缺资料"
    assert missing["next_actions"] == [{"kind": "fact_missing", "key": "demand",
                                        "label": "首年需求", "action": "new_fact"}]
    assert client.post(base + "/facts", json={"key": "demand", "label": "首年需求",
        "data_type": "integer", "unit": "套", "value": "0"}).status_code == 201
    unverified = next(row for row in client.get(endpoint).json()["chapters"] if row["section_id"] == "necessity")
    assert unverified["next_actions"][0]["action"] == "review_source"
    document_id = _source(client, project_id)
    batch = _batch(document_id); batch["items"] = [{**batch["items"][1], "key": "duration"}]
    preview = client.post(base + "/facts/from-document/preview", json=batch).json()
    assert client.post(base + "/facts/from-document/commit", json={
        **batch, "base_version": preview["base_version"],
        "preview_token": preview["preview_token"]}).status_code == 201
    assert client.post(base + "/facts/demand/evidence/bind", json={
        "document_id": document_id, "source_refs": ["d-p1"]}).status_code == 200
    ready = next(row for row in client.get(endpoint).json()["chapters"] if row["section_id"] == "necessity")
    assert ready["status"] == "可起草" and not ready["next_actions"]
    evidence = client.post(base + "/evidence", json={"document_id": document_id,
        "source_refs": ["d-p1"], "statement": "本项目首年需求为 0 套。",
        "label": "需求原文", "fact_key": "demand"})
    assert evidence.status_code == 201, evidence.text
    issue = client.post(base + "/issues", json={"kind": "claim", "title": "首年需求",
        "statement": "本项目首年需求为 0 套。", "section_ids": ["necessity"],
        "support_evidence_ids": [evidence.json()["id"]]})
    assert issue.status_code == 201, issue.text
    with_issue = next(row for row in client.get(endpoint).json()["chapters"] if row["section_id"] == "necessity")
    assert with_issue["open_issues"][0]["id"] == issue.json()["id"]
    assert with_issue["status"] == "可起草"  # Drafting is possible; review remains a separate decision.


def test_real_feasibility_pdf_fact_keeps_page_and_source(client):
    project_id = client.post("/api/projects", json={"name": "北京可研 PDF 录入"}).json()["id"]
    base = f"/api/projects/{project_id}"
    fixture = Path(__file__).resolve().parents[2] / "test-fixtures/public-reports/01_beijing_feasibility.pdf"
    with fixture.open("rb") as stream:
        upload = client.post(base + "/documents", files={"file": ("北京公开可研.pdf", stream, "application/pdf")})
    assert upload.status_code == 201, upload.text
    document_id = upload.json()["id"]
    page = client.get(base + f"/documents/{document_id}?page=7").json()
    segment = next(row for row in page["segments"] if "17268.25" in row["text"])
    batch = {"document_id": document_id, "items": [{
        "key": "reported_total_area", "label": "原文总建筑面积", "data_type": "decimal",
        "value": "17268.25", "unit": "㎡", "source_ref": segment["ref"],
    }]}
    preview = client.post(base + "/facts/from-document/preview", json=batch)
    assert preview.status_code == 200, preview.text
    commit = client.post(base + "/facts/from-document/commit", json={
        **batch, "base_version": preview.json()["base_version"],
        "preview_token": preview.json()["preview_token"]})
    assert commit.status_code == 201, commit.text
    source = client.get(base + "/facts/reported_total_area/source").json()
    assert source["document_id"] == document_id and source["source_ref"] == segment["ref"]
    assert source["page"] == 7
