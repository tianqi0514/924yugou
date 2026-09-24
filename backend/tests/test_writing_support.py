"""Database-backed writing flow, isolated from the user's PostgreSQL database."""

from __future__ import annotations

import copy
import json
import os
from io import BytesIO
from zipfile import ZipFile

import pymupdf
from docx import Document

os.environ["DATABASE_URL"] = "postgresql+psycopg://report:local_development_only@127.0.0.1:55432/report_platform_test"

import pytest
from fastapi.testclient import TestClient

from app.db import Base, engine
from app.main import app


@pytest.fixture(autouse=True)
def isolated_database():
    assert engine.url.database == "report_platform_test", "Refusing to reset a non-test database"
    writable = [table for table in Base.metadata.sorted_tables if not table.name.startswith("corpus_")]
    Base.metadata.drop_all(engine, tables=writable)
    Base.metadata.create_all(engine, tables=writable)
    yield
    Base.metadata.drop_all(engine, tables=writable)


@pytest.fixture
def client():
    with TestClient(app) as instance:
        yield instance


def ok(response):
    assert response.status_code in (200, 201), response.text
    return response.json()


def make_project(client):
    pid = ok(client.post("/api/projects", json={"name": "写作包独立验收"}))["id"]
    assert ok(client.get(f"/api/projects/{pid}/facts"))["facts"] == []
    assert client.get(f"/api/projects/{pid}/writing/sections").status_code == 409
    bad = {"corpus_id": "cy_tray_20260918", "corpus_version": "v2",
           "target_report_type": "accident_investigation", "accepted": True}
    assert client.put(f"/api/projects/{pid}/writing/reference", json=bad).status_code == 409
    good = {**bad, "target_report_type": "feasibility"}
    assert ok(client.put(f"/api/projects/{pid}/writing/reference", json=good))["selected"]
    assert ok(client.get(f"/api/projects/{pid}/facts"))["facts"] == []
    return pid


def setup(client, pid):
    assert ok(client.post(f"/api/projects/{pid}/writing/setup", json={}))["historical_values_copied"] == 0
    bindings = {"N017": "first_year_demand", "N034": "qualified_capacity",
                "N035": "planned_sales", "supplier_name": "supplier_name",
                "N080": "equipment_unit_price"}
    ok(client.put(f"/api/projects/{pid}/writing/bindings", json={"bindings": bindings}))


def set_fact(client, pid, key, value):
    change = {"fact_key": key, "value": value, "source": "本项目核对输入", "reason": "写作验证"}
    preview = ok(client.post(f"/api/projects/{pid}/changes/preview", json=change))
    ok(client.post(f"/api/projects/{pid}/changes/commit", json={
        **change, "base_version": preview["base_version"], "preview_token": preview["preview_token"]}))


def new_report(client, pid):
    return ok(client.post(f"/api/projects/{pid}/reports", json={"title": "写作支撑试验稿"}))["id"]


def candidate(client, pid, rid, section):
    pack = ok(client.get(f"/api/projects/{pid}/writing/packages/{section}"))
    approved = [item["item_id"] for item in pack["items"] if item["decision"] == "selectable"]
    return client.post(f"/api/projects/{pid}/reports/{rid}/writing/preview", json={
        "section_id": section, "approved_item_ids": approved, "mode": "guided"})


def commit(client, pid, rid, value):
    return ok(client.post(f"/api/projects/{pid}/reports/{rid}/writing/commit", json=value))


def test_reference_package_and_source_ids(client):
    pid = make_project(client)
    sections = ok(client.get(f"/api/projects/{pid}/writing/sections"))["sections"]
    assert len(sections) == 49
    assert {item["id"] for item in sections if item["guided_available"]} == {"S4", "S7.1", "S5.2"}
    s4 = ok(client.get(f"/api/projects/{pid}/writing/packages/S4"))
    ids = {(item["category_number"], item["semantic_id"]): item for item in s4["items"]}
    for category, semantic in ((9, "C0052"), (10, "C0052"), (13, "R006"),
                               (15, "L002"), (18, "EV-01")):
        item = ids[(category, semantic)]
        assert item["record_id"].startswith(f"{category:02d}:")
        assert item["artifact_id"].startswith("ART-")
        source = ok(client.get(f"/api/projects/{pid}/writing/sources/{category}/{semantic}"))
        assert source["source_ref"]["record_id"] == item["record_id"]
    excerpt = ok(client.get(f"/api/projects/{pid}/writing/sources/19/EX-01-01"))
    assert excerpt["source_ref"]["record_id"].startswith("19:")
    assert excerpt["payload"]["link"]["id"] == "EX-01-01"
    assert ids[(18, "EV-01")]["decision"] == "review_only"
    assert {(item["category_number"], item["semantic_id"]) for item in s4["items"] if item["required"]} == {
        (4, "S4"), (27, "S4"), (10, "C0052"), (13, "R006"), (15, "L002")}
    assert s4["required_item_ids"] == [item["item_id"] for item in s4["items"] if item["required"]]
    assert [issue["code"] for issue in s4["issues"]].count("HISTORICAL_EVIDENCE_INCOMPLETE") == 1
    assert s4["facts"] == []
    s52 = ok(client.get(f"/api/projects/{pid}/writing/packages/S5.2"))
    assert {"X001", "EV-06A", "EV-06B"}.issubset({item["semantic_id"] for item in s52["items"]})
    assert {(item["category_number"], item["semantic_id"]) for item in s52["items"] if item["required"]} == {
        (4, "S5.2"), (27, "S5.2")}
    assert not any(item["semantic_id"] == "EV-01" for item in s52["items"])
    assert sum(issue["code"] == "HISTORICAL_EVIDENCE_INCOMPLETE" for issue in s52["issues"]) <= 2


def test_guided_two_sections_preview_cancel_commit_and_preview_export(client):
    pid = make_project(client)
    setup(client, pid)
    rid = new_report(client, pid)
    assert candidate(client, pid, rid, "S4").status_code == 409
    set_fact(client, pid, "first_year_demand", "300000")
    set_fact(client, pid, "qualified_capacity", "254016")
    old = ok(client.get(f"/api/projects/{pid}/reports/{rid}"))
    s4 = ok(candidate(client, pid, rid, "S4"))
    assert ok(client.get(f"/api/projects/{pid}/reports/{rid}"))["version"] == old["version"]
    body = json.dumps(s4["paragraphs"], ensure_ascii=False)
    assert "300,000" in body and "254,016" in body and "需求低于能力" not in body
    assert "180,000" not in body
    assert {"09:000051", "10:000051", "13:000005", "15:000001", "18:000001"}.issubset(set(s4["used_items"]))
    assert any(child.get("type") == "fact_ref" for block in s4["paragraphs"] for child in block.get("children", []))
    saved = commit(client, pid, rid, s4)
    assert saved["report"]["version"] == old["version"] + 1
    assert client.post(f"/api/projects/{pid}/reports/{rid}/writing/commit", json=s4).status_code == 409
    impact = ok(client.get(f"/api/projects/{pid}/writing/impact?record_id=13%3A000005"))["impacts"]
    assert impact and impact[0]["block_id"]

    set_fact(client, pid, "supplier_name", "新拓设备")
    set_fact(client, pid, "equipment_unit_price", "1234.50")
    s7 = ok(candidate(client, pid, rid, "S7.1"))
    body = json.dumps(s7["paragraphs"], ensure_ascii=False)
    assert "新拓设备" in body and "禾进装备" not in body
    assert "12,345,000" in body and "1,234.50" in body
    commit(client, pid, rid, s7)
    current = ok(client.get(f"/api/projects/{pid}/reports/{rid}"))
    assert any(issue["code"] == "PROJECT_EVIDENCE_UNVERIFIED" for issue in current["issues"])
    assert client.post(f"/api/projects/{pid}/reports/{rid}/review", json={}).status_code == 400
    assert client.get(f"/api/projects/{pid}/reports/{rid}/export").status_code == 409
    exported = client.get(f"/api/projects/{pid}/reports/{rid}/export?level=preview")
    assert exported.status_code == 200, exported.text
    with ZipFile(BytesIO(exported.content)) as archive:
        audit = json.loads(archive.read("audit.json"))
        docx = Document(BytesIO(archive.read("report.docx")))
        with pymupdf.open(stream=archive.read("report.pdf"), filetype="pdf") as pdf:
            pdf_text = "".join(page.get_text() for page in pdf)
    assert audit["delivery_status"] == "preview_only"
    assert "预审稿" in docx.paragraphs[0].text
    assert docx.paragraphs[0].text == "预审稿 · 来源待核 · 不构成正式结论"
    assert "预审稿，来源待核，不构成正式结论" in pdf_text
    assert any(issue["code"] == "PROJECT_EVIDENCE_UNVERIFIED" for issue in audit["writing_issues"])
    assert any(ref["record_id"] == "13:000005" for item in audit["source_refs"] for ref in item["refs"])
    project_rules = [rule for item in audit["source_refs"] for rule in item["project_rule_refs"]]
    assert len(project_rules) == 1
    assert project_rules[0]["expression"] == "min(first_year_demand, qualified_capacity)"
    assert project_rules[0]["input_fact_revisions"]
    computed_basis = next(paragraph.text for paragraph in docx.paragraphs if paragraph.text.startswith("首年计划销售量："))
    assert "本项目计算规则：min(first_year_demand, qualified_capacity)" in computed_basis
    assert "首年客户需求 r1" in computed_basis and "合格能力 r1" in computed_basis
    assert "输入证据待核对" in computed_basis and "来源：" not in computed_basis
    assert "min(first_year_demand, qualified_capacity)" in pdf_text
    calculated = next(entry for entry in audit["evidence_bindings"] if entry["fact_key"] == "planned_sales")
    assert calculated["status"] == "derived_inputs_unverified"
    assert calculated["document_id"] is None
    assert calculated["project_rule"]["rule_id"] == project_rules[0]["rule_id"]
    assert {entry["fact_key"] for entry in calculated["input_fact_refs"]} == {
        "first_year_demand", "qualified_capacity"}


def test_conflict_disclosure_and_removed_source_link(client):
    pid = make_project(client)
    rid = new_report(client, pid)
    s52 = ok(candidate(client, pid, rid, "S5.2"))
    body = " ".join(str(child.get("text", child.get("display", "")))
                    for block in s52["paragraphs"] for child in block.get("children", []))
    assert "暂不作配电适配结论" in body and "800" not in body and "650" not in body
    assert "23:000000" in s52["used_items"]
    saved = commit(client, pid, rid, s52)["report"]
    assert any(issue["code"] == "POWER_CONFLICT_UNRESOLVED" for issue in saved["issues"])
    assert client.post(f"/api/projects/{pid}/reports/{rid}/review", json={}).status_code == 400
    assert client.get(f"/api/projects/{pid}/reports/{rid}/export").status_code == 409
    content = [node for node in saved["content"] if not any(
        ref.get("semantic_id") == "X001" for ref in node.get("source_refs", []))]
    proposal = ok(client.post(f"/api/projects/{pid}/reports/{rid}/preview", json={
        "content": content, "base_version": saved["version"]}))
    changed = ok(client.put(f"/api/projects/{pid}/reports/{rid}", json={
        "content": content, "base_version": saved["version"],
        "preview_token": proposal["preview_token"]}))["report"]
    assert not any(issue["code"] == "POWER_CONFLICT_UNRESOLVED" for issue in changed["issues"])


def test_source_ref_tampering_blocks_review(client):
    pid = make_project(client)
    setup(client, pid)
    set_fact(client, pid, "first_year_demand", "300000")
    set_fact(client, pid, "qualified_capacity", "254016")
    rid = new_report(client, pid)
    original = commit(client, pid, rid, ok(candidate(client, pid, rid, "S4")))["report"]
    edited = copy.deepcopy(original["content"])
    target = next(node for node in edited if any(
        ref.get("semantic_id") == "R006" for ref in node.get("source_refs", [])))
    target["source_refs"] = []
    proposal = ok(client.post(f"/api/projects/{pid}/reports/{rid}/preview", json={
        "content": edited, "base_version": original["version"]}))
    changed = ok(client.put(f"/api/projects/{pid}/reports/{rid}", json={
        "content": edited, "base_version": original["version"],
        "preview_token": proposal["preview_token"]}))["report"]
    assert any(issue["code"] == "SOURCE_LINK_CHANGED" for issue in changed["issues"])
    assert client.post(f"/api/projects/{pid}/reports/{rid}/review", json={}).status_code == 400


def test_stale_project_rule_provenance_cannot_rebind_updated_fact(client):
    pid = make_project(client)
    setup(client, pid)
    set_fact(client, pid, "first_year_demand", "300000")
    set_fact(client, pid, "qualified_capacity", "254016")
    rid = new_report(client, pid)
    original = commit(client, pid, rid, ok(candidate(client, pid, rid, "S4")))["report"]
    set_fact(client, pid, "first_year_demand", "200000")
    edited = copy.deepcopy(original["content"])
    calculated = next(node for node in edited if node.get("project_rule_refs"))
    for child in calculated["children"]:
        if child.get("fact_key") in {"first_year_demand", "planned_sales"}:
            child["display"] = "200,000"
    proposed = ok(client.post(f"/api/projects/{pid}/reports/{rid}/preview", json={
        "content": edited, "base_version": original["version"]}))
    saved = ok(client.put(f"/api/projects/{pid}/reports/{rid}", json={
        "content": edited, "base_version": original["version"],
        "preview_token": proposed["preview_token"]}))["report"]
    codes = {issue["code"] for issue in saved["issues"]}
    assert "PROJECT_RULE_INPUT_CHANGED" in codes
    assert "FACT_CHANGED" in codes
    assert saved["bound_facts"]["first_year_demand"] == original["bound_facts"]["first_year_demand"]
    assert saved["bound_facts"]["planned_sales"] == original["bound_facts"]["planned_sales"]
    assert client.post(f"/api/projects/{pid}/reports/{rid}/review", json={}).status_code == 400


def test_project_original_binding_allows_formal_export_and_stales_on_revision(client, monkeypatch, tmp_path):
    import app.main as main_module

    monkeypatch.setattr(main_module, "STORAGE", tmp_path)
    pid = make_project(client)
    setup(client, pid)
    for key, value in (("first_year_demand", "300000"), ("qualified_capacity", "254016"),
                       ("supplier_name", "新拓设备"), ("equipment_unit_price", "1234.50")):
        set_fact(client, pid, key, value)
    rid = new_report(client, pid)
    commit(client, pid, rid, ok(candidate(client, pid, rid, "S4")))
    commit(client, pid, rid, ok(candidate(client, pid, rid, "S7.1")))

    original = Document()
    original.add_paragraph("首年客户需求 300000 套，来自本项目预测资料。")
    original.add_paragraph("本项目合格能力为 254016 套。")
    original.add_paragraph("供应商新拓设备报自动焊接线 1234.50 万元/条。")
    stream = BytesIO()
    original.save(stream)
    uploaded = ok(client.post(f"/api/projects/{pid}/documents", files={
        "file": ("本项目核对原件.docx", stream.getvalue(),
                 "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}))
    doc_id = uploaded["id"]
    document = ok(client.get(f"/api/projects/{pid}/documents/{doc_id}"))
    assert len(document["segments"]) == 3
    mapping = {"first_year_demand": "d-p1", "qualified_capacity": "d-p2",
               "supplier_name": "d-p3", "equipment_unit_price": "d-p3"}
    invalid = client.post(f"/api/projects/{pid}/facts/equipment_unit_price/evidence/bind", json={
        "document_id": doc_id, "source_refs": ["d-p1"]})
    assert invalid.status_code == 409
    for key, source_ref in mapping.items():
        response = ok(client.post(f"/api/projects/{pid}/facts/{key}/evidence/bind", json={
            "document_id": doc_id, "source_refs": [source_ref]}))
        assert response["status"] == "SOURCE_LOCATOR_REVIEWED"
    assert ok(client.get(f"/api/projects/{pid}/facts/first_year_demand/evidence"))["status"] == "SOURCE_LOCATOR_REVIEWED"
    direct_source = ok(client.get(f"/api/projects/{pid}/facts/first_year_demand/source"))
    assert direct_source["kind"] == "document" and direct_source["review_status"] == "SOURCE_LOCATOR_REVIEWED"
    assert direct_source["filename"] == "本项目核对原件.docx" and direct_source["source_ref"] == "d-p1"
    assert "300000" in direct_source["excerpt"] and direct_source["original_url"].endswith("/original")
    calculated_source = ok(client.get(f"/api/projects/{pid}/facts/planned_sales/source"))
    assert calculated_source["kind"] == "rule"
    assert calculated_source["review_status"] == "DERIVED_FROM_REVIEWED_INPUTS"
    assert calculated_source["rule"]["expression"] == "min(first_year_demand, qualified_capacity)"
    assert {entry["key"] for entry in calculated_source["input_facts"]} == {
        "first_year_demand", "qualified_capacity"}
    assert all(entry["review_status"] == "SOURCE_LOCATOR_REVIEWED"
               and entry["source_url"].endswith(f"/facts/{entry['key']}/source")
               for entry in calculated_source["input_facts"])
    reviewed = ok(client.post(f"/api/projects/{pid}/reports/{rid}/review", json={}))
    assert reviewed["reviewed"]
    exported = client.get(f"/api/projects/{pid}/reports/{rid}/export")
    assert exported.status_code == 200, exported.text
    with ZipFile(BytesIO(exported.content)) as archive:
        audit = json.loads(archive.read("audit.json"))
        word = Document(BytesIO(archive.read("report.docx")))
        with pymupdf.open(stream=archive.read("report.pdf"), filetype="pdf") as pdf:
            pdf_text = "".join(page.get_text() for page in pdf)
    assert audit["delivery_status"] == "source_locator_reviewed"
    assert any(entry["document_sha256"] == uploaded["sha256"] for entry in audit["evidence_bindings"])
    assert any(entry["source_refs"] == ["d-p3"] for entry in audit["evidence_bindings"])
    assert any(entry["document_filename"] == "本项目核对原件.docx"
               and entry["location_labels"] == ["段落 1"] for entry in audit["evidence_bindings"])
    assert any("本项目原件：本项目核对原件.docx，段落 1" in paragraph.text
               for paragraph in word.paragraphs)
    assert "本项目原件：本项目核对原件.docx，段落 1" in pdf_text
    assert not any("来源：本项目核对输入" in paragraph.text for paragraph in word.paragraphs)
    calculated = next(entry for entry in audit["evidence_bindings"] if entry["fact_key"] == "planned_sales")
    assert calculated["status"] == "derived_from_reviewed_inputs"
    assert calculated["document_id"] is None
    assert calculated["project_rule"]["expression"] == "min(first_year_demand, qualified_capacity)"
    assert {entry["fact_key"] for entry in calculated["input_fact_refs"]} == {
        "first_year_demand", "qualified_capacity"}
    assert all(entry["status"] == "source_locator_reviewed" for entry in calculated["input_fact_refs"])
    assert all(entry["document_sha256"] == uploaded["sha256"] for entry in calculated["input_fact_refs"])
    computed_basis = next(paragraph.text for paragraph in word.paragraphs if paragraph.text.startswith("首年计划销售量："))
    assert "本项目计算规则：min(first_year_demand, qualified_capacity)" in computed_basis
    assert "输入原文位置已由操作者核对" in computed_basis and "来源：" not in computed_basis
    assert "min(first_year_demand, qualified_capacity)" in pdf_text

    set_fact(client, pid, "equipment_unit_price", "999")
    assert ok(client.get(f"/api/projects/{pid}/facts/equipment_unit_price/evidence"))["status"] == "UNVERIFIED"
    assert client.post(f"/api/projects/{pid}/facts/equipment_unit_price/evidence/bind", json={
        "document_id": doc_id, "source_refs": ["d-p3"]}).status_code == 409
    assert client.get(f"/api/projects/{pid}/reports/{rid}/export").status_code == 409
