"""对真实 PostgreSQL 测试库与内置语料执行关键验收。"""

from __future__ import annotations

import os
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile
from decimal import Decimal

os.environ["DATABASE_URL"] = "postgresql+psycopg://report:local_development_only@127.0.0.1:55432/report_platform_test"

import pytest
from fastapi.testclient import TestClient

from app.corpus import CorpusRepository
from app.db import Base, engine
from app.main import app
from app.rules import EvalValue, RuleError, evaluate, parse_expression, unit_dimension


FIXTURES = Path(__file__).resolve().parents[2] / "test-fixtures" / "public-reports"


@pytest.fixture(autouse=True)
def isolated_database():
    assert engine.url.database == "report_platform_test", "Refusing to reset a non-test database"
    # 历史语料只读且导入较重；每个用例只重建项目可写表，仍保持项目数据隔离。
    writable_tables = [table for table in Base.metadata.sorted_tables if not table.name.startswith("corpus_")]
    Base.metadata.drop_all(engine, tables=writable_tables)
    Base.metadata.create_all(engine, tables=writable_tables)
    yield
    Base.metadata.drop_all(engine, tables=writable_tables)


@pytest.fixture
def client():
    with TestClient(app) as test_client:
        yield test_client


def create_project(client: TestClient) -> str:
    response = client.post("/api/projects", json={"name": "验收项目"})
    assert response.status_code == 201, response.text
    project_id = response.json()["id"]
    assert client.get(f"/api/projects/{project_id}/facts").json()["facts"] == []
    return project_id


def test_demo_documents_are_seeded_once_without_facts(client: TestClient, tmp_path):
    from app.db import SessionLocal
    from app.demo_seed import DEMO_DOCUMENTS, seed_demo_documents

    assert engine.url.database == "report_platform_test"
    first = seed_demo_documents(storage_dir=tmp_path, session_factory=SessionLocal)
    second = seed_demo_documents(storage_dir=tmp_path, session_factory=SessionLocal)
    checked = seed_demo_documents(storage_dir=tmp_path, session_factory=SessionLocal, verify_only=True)
    assert len(first) == len(second) == len(checked) == 4
    assert {row["project_id"] for row in first} == {spec.project_id for spec in DEMO_DOCUMENTS}
    assert all(row["seed_state"] == "已预置" and row["fact_state"] == "事实台账为空" for row in first)
    assert all(row["seed_state"] == "已存在" for row in second + checked)
    for row, spec in zip(first, DEMO_DOCUMENTS):
        documents = client.get(f"/api/projects/{spec.project_id}/documents").json()
        assert len(documents) == 1
        assert documents[0]["id"] == row["document_id"] and documents[0]["sha256"] == spec.sha256
        assert documents[0]["pages"] == spec.pages
        assert client.get(f"/api/projects/{spec.project_id}/facts").json()["facts"] == []
        assert (tmp_path / spec.storage_name).is_file()
    assert first[2]["status"] == "OCR_REQUIRED"


def add_fact(client: TestClient, project_id: str, key: str, unit: str = "套") -> None:
    response = client.post(f"/api/projects/{project_id}/facts", json={
        "key": key, "label": key, "data_type": "integer", "unit": unit,
        "caliber": "首年", "source": "", "as_of": "",
    })
    assert response.status_code == 201, response.text


def preview(client: TestClient, project_id: str, key: str, value: str | None):
    change = {"fact_key": key, "value": value, "source": "本项目资料", "reason": "验收输入"}
    result = client.post(f"/api/projects/{project_id}/changes/preview", json=change)
    assert result.status_code == 200, result.text
    return change, result.json()


def commit(client: TestClient, project_id: str, change: dict, result: dict):
    response = client.post(f"/api/projects/{project_id}/changes/commit", json={
        **change, "base_version": result["base_version"], "preview_token": result["preview_token"],
    })
    assert response.status_code == 200, response.text
    return response.json()


def bind_project_original(client: TestClient, project_id: str, key: str, value: str) -> None:
    from docx import Document

    original = Document()
    original.add_paragraph(f"本项目 {key} 数量为 {value} 套。")
    stream = BytesIO()
    original.save(stream)
    uploaded = client.post(f"/api/projects/{project_id}/documents", files={
        "file": (f"{key}-{value}.docx", stream.getvalue(),
                 "application/vnd.openxmlformats-officedocument.wordprocessingml.document")})
    assert uploaded.status_code == 201, uploaded.text
    bound = client.post(f"/api/projects/{project_id}/facts/{key}/evidence/bind", json={
        "document_id": uploaded.json()["id"], "source_refs": ["d-p1"]})
    assert bound.status_code == 200, bound.text
    assert bound.json()["status"] == "SOURCE_LOCATOR_REVIEWED"


def test_preview_cancel_commit_zero_and_missing(client: TestClient):
    project_id = create_project(client)
    for key in ("first_year_demand", "qualified_capacity", "planned_sales"):
        add_fact(client, project_id, key)
    rule = client.post(f"/api/projects/{project_id}/rules", json={
        "name": "计划销售量", "target_key": "planned_sales",
        "expression": "min(first_year_demand, qualified_capacity)",
    })
    assert rule.status_code == 201, rule.text
    initial = {x["key"]: x for x in client.get(f"/api/projects/{project_id}/facts").json()["facts"]}
    assert initial["planned_sales"]["status"] == "UNEVALUABLE"
    assert initial["planned_sales"]["value"] is None

    change, result = preview(client, project_id, "first_year_demand", "300000")
    assert any(x["key"] == "first_year_demand" for x in result["changes"])
    # 只预览后取消：不调用提交 API，版本和值均保持不变。
    assert {x["key"]: x for x in client.get(f"/api/projects/{project_id}/facts").json()["facts"]}["first_year_demand"]["value"] is None
    commit(client, project_id, change, result)
    assert client.get(f"/api/projects/{project_id}/facts").json()["project_version"] == result["base_version"] + 1

    change, result = preview(client, project_id, "qualified_capacity", "254016")
    assert next(x for x in result["changes"] if x["key"] == "planned_sales")["after"]["value"] == "254016"
    commit(client, project_id, change, result)
    facts = {x["key"]: x for x in client.get(f"/api/projects/{project_id}/facts").json()["facts"]}
    assert facts["planned_sales"]["value"] == "254016"
    assert facts["planned_sales"]["status"] == "COMPUTED"
    assert len(client.get(f"/api/projects/{project_id}/revisions").json()) >= 3

    change, result = preview(client, project_id, "qualified_capacity", "0")
    assert next(x for x in result["changes"] if x["key"] == "planned_sales")["after"]["value"] == "0"
    commit(client, project_id, change, result)
    facts = {x["key"]: x for x in client.get(f"/api/projects/{project_id}/facts").json()["facts"]}
    assert facts["qualified_capacity"]["value"] == "0"
    assert facts["planned_sales"]["value"] == "0"

    change, result = preview(client, project_id, "qualified_capacity", None)
    commit(client, project_id, change, result)
    facts = {x["key"]: x for x in client.get(f"/api/projects/{project_id}/facts").json()["facts"]}
    assert facts["qualified_capacity"]["status"] == "UNDEFINED"
    assert facts["planned_sales"]["status"] == "UNEVALUABLE"


def test_rule_validation_and_atomic_failures(client: TestClient):
    project_id = create_project(client)
    for key, unit in (("a", "套"), ("b", "套"), ("wrong_unit", "kW"), ("result", "套")):
        add_fact(client, project_id, key, unit)
    for key, value in (("a", "10"), ("b", "2"), ("wrong_unit", "3")):
        change, result = preview(client, project_id, key, value)
        commit(client, project_id, change, result)
    before = client.get(f"/api/projects/{project_id}/facts").json()
    for expression in ("a / (b - b)", "a + wrong_unit"):
        response = client.post(f"/api/projects/{project_id}/rules", json={
            "name": "无效规则", "target_key": "result", "expression": expression,
        })
        assert response.status_code == 400
        assert client.get(f"/api/projects/{project_id}/facts").json() == before
        assert client.get(f"/api/projects/{project_id}/rules").json()["rules"] == []
    with pytest.raises(RuleError):
        parse_expression("__import__('os').system('echo unsafe')")
    with pytest.raises(RuleError, match="单位不一致"):
        evaluate("yuan + ten_thousand_yuan", {
            "yuan": EvalValue(Decimal("1"), unit_dimension("元"), "元"),
            "ten_thousand_yuan": EvalValue(Decimal("1"), unit_dimension("万元"), "万元"),
        })

    # 依赖回路应在新增第二条规则时拒绝，且第一条仍保持原样。
    first = client.post(f"/api/projects/{project_id}/rules", json={"name": "a 来源 b", "target_key": "a", "expression": "b"})
    assert first.status_code == 200 or first.status_code == 201, first.text
    second = client.post(f"/api/projects/{project_id}/rules", json={"name": "b 来源 a", "target_key": "b", "expression": "a"})
    assert second.status_code == 400
    assert len(client.get(f"/api/projects/{project_id}/rules").json()["rules"]) == 1


def test_stale_preview_cannot_commit(client: TestClient):
    project_id = create_project(client)
    add_fact(client, project_id, "amount")
    old_change, old_preview = preview(client, project_id, "amount", "10")
    change, fresh_preview = preview(client, project_id, "amount", "11")
    commit(client, project_id, change, fresh_preview)
    response = client.post(f"/api/projects/{project_id}/changes/commit", json={
        **old_change, "base_version": old_preview["base_version"], "preview_token": old_preview["preview_token"],
    })
    assert response.status_code == 409
    assert client.get(f"/api/projects/{project_id}/facts").json()["facts"][0]["value"] == "11"


def test_project_file_review_is_source_bound_and_isolated(client: TestClient, monkeypatch, tmp_path):
    import pymupdf
    import app.main as main_module

    monkeypatch.setattr(main_module, "STORAGE", tmp_path)
    first, second = create_project(client), create_project(client)
    pdf = pymupdf.open()
    page = pdf.new_page()
    page.insert_text((50, 50), "Demand 300000 units. Capacity 254016 units.")
    data = pdf.tobytes()
    uploaded = client.post(f"/api/projects/{first}/documents", files={"file": ("source.pdf", data, "application/pdf")})
    assert uploaded.status_code == 201, uploaded.text
    document_id = uploaded.json()["id"]
    assert client.get(f"/api/projects/{second}/documents/{document_id}").status_code == 404
    assert client.get(f"/api/projects/{first}/documents/{document_id}/original").status_code == 200
    assert client.get(f"/api/projects/{first}/facts").json()["facts"] == []
    assert client.post(f"/api/projects/{first}/documents", files={"file": ("copy.pdf", data, "application/pdf")}).status_code == 409

    detail = client.get(f"/api/projects/{first}/documents/{document_id}").json()
    ref = detail["segments"][0]["ref"]
    pages = client.get(f"/api/projects/{first}/documents/{document_id}/pages")
    assert pages.status_code == 200
    assert pages.json()["pages"][0]["extractable"] is True
    assert pages.json()["pages"][0]["segment_count"] == len(detail["segments"])
    assert client.get(f"/api/projects/{second}/documents/{document_id}/pages").status_code == 404
    assert client.post(f"/api/projects/{first}/documents/{document_id}/extract?page=2").status_code == 404
    monkeypatch.setattr(main_module, "model_candidates", lambda segments: [
        {"label": "需求", "value_text": "300000", "unit": "套", "data_type": "integer", "source_ref": ref, "source_valid": True},
        {"label": "错引", "value_text": "999", "unit": "套", "data_type": "integer", "source_ref": ref, "source_valid": False},
        {"label": "截断数值", "value_text": "30000", "unit": "套", "data_type": "integer", "source_ref": ref, "source_valid": True},
    ])
    extracted = client.post(f"/api/projects/{first}/documents/{document_id}/extract?page=1")
    assert extracted.status_code == 200 and extracted.json()["created"] == 3
    repeated = client.post(f"/api/projects/{first}/documents/{document_id}/extract?page=1")
    assert repeated.status_code == 200 and repeated.json()["created"] == 0
    assert client.get(f"/api/projects/{first}/documents/{document_id}/pages").json()["pages"][0]["pending_count"] == 3
    assert client.get(f"/api/projects/{first}/facts").json()["facts"] == []
    candidates = client.get(f"/api/projects/{first}/documents/{document_id}").json()["candidates"]
    valid = next(item for item in candidates if item["label"] == "需求")
    invalid = next(item for item in candidates if item["label"] == "错引")
    truncated = next(item for item in candidates if item["label"] == "截断数值")
    assert truncated["source_valid"] is False
    # Older stored flags are advisory; reads recheck against the immutable source segment.
    from app.db import ExtractionCandidate, SessionLocal
    with SessionLocal.begin() as session:
        session.get(ExtractionCandidate, truncated["id"]).source_valid = True
    assert next(item for item in client.get(f"/api/projects/{first}/documents/{document_id}").json()["candidates"] if item["id"] == truncated["id"])["source_valid"] is False
    assert client.get(f"/api/projects/{first}/documents/{document_id}/pages").json()["pages"][0]["invalid_count"] == 2
    body = {"label": "首年需求", "value": "300000", "unit": "套", "data_type": "integer", "source_ref": ref, "key": "demand"}
    assert client.post(f"/api/projects/{second}/documents/{document_id}/candidates/{valid['id']}/approve", json=body).status_code == 404
    assert client.post(f"/api/projects/{first}/documents/{document_id}/candidates/{invalid['id']}/approve", json={**body, "value": "999"}).status_code == 400
    assert client.post(f"/api/projects/{first}/documents/{document_id}/candidates/{truncated['id']}/approve", json={**body, "value": "30000"}).status_code == 400
    approved = client.post(f"/api/projects/{first}/documents/{document_id}/candidates/{valid['id']}/approve", json=body)
    assert approved.status_code == 200, approved.text
    assert approved.json()["fact"]["value"] == "300000"
    assert next(item for item in client.get(f"/api/projects/{first}/documents/{document_id}").json()["candidates"] if item["id"] == valid["id"])["approved_fact"]["value"] == "300000"
    source = client.get(f"/api/projects/{first}/facts/demand/source")
    assert source.status_code == 200
    assert source.json()["kind"] == "document"
    assert "300000" in source.json()["excerpt"]
    assert source.json()["source_ref"] == ref
    assert source.json()["original_url"].endswith("#page=1")
    assert client.get(f"/api/projects/{second}/facts/demand/source").status_code == 404
    assert client.post(f"/api/projects/{first}/documents/{document_id}/candidates/{valid['id']}/approve", json=body).status_code == 409
    assert len(client.get(f"/api/projects/{first}/revisions").json()) == 1
    assert client.get(f"/api/projects/{second}/facts").json()["facts"] == []
    corrected = client.post(f"/api/projects/{first}/documents/{document_id}/candidates/{truncated['id']}/approve", json={**body, "key": "corrected_demand"})
    assert corrected.status_code == 200
    corrected_candidate = next(item for item in client.get(f"/api/projects/{first}/documents/{document_id}").json()["candidates"] if item["id"] == truncated["id"])
    assert corrected_candidate["value"] == "30000"  # retain raw model suggestion
    assert corrected_candidate["approved_fact"]["value"] == "300000"


def test_table_candidates_and_multi_segment_evidence_are_project_bound(client: TestClient, monkeypatch, tmp_path):
    """真实公开表格可生成候选，确认后能保留两处同页证据。"""
    import pymupdf
    import app.main as main_module

    monkeypatch.setattr(main_module, "STORAGE", tmp_path)
    monkeypatch.setattr(main_module, "model_candidates", lambda segments: [
        {"label": "容积率", "value_text": "0.291", "unit": "", "data_type": "decimal", "source_ref": "p1-t1-r8"}
    ])
    project_id = create_project(client)
    with pymupdf.open(FIXTURES / "01_beijing_feasibility.pdf") as source:
        pdf = pymupdf.open()
        pdf.insert_pdf(source, from_page=6, to_page=6)
        pdf.new_page().insert_text((50, 50), "Other page 0.291")
        data = pdf.tobytes()
    uploaded = client.post(f"/api/projects/{project_id}/documents", files={"file": ("table.pdf", data, "application/pdf")})
    assert uploaded.status_code == 201, uploaded.text
    document_id = uploaded.json()["id"]
    before = client.get(f"/api/projects/{project_id}/documents/{document_id}?page=1").json()
    old_refs = {part["ref"] for part in before["segments"]}
    extracted = client.post(f"/api/projects/{project_id}/documents/{document_id}/extract?page=1")
    assert extracted.status_code == 200, extracted.text
    assert extracted.json()["table_returned"] == 18
    assert extracted.json()["created"] == 18
    detail = client.get(f"/api/projects/{project_id}/documents/{document_id}?page=1").json()
    assert old_refs <= {part["ref"] for part in detail["segments"]}
    ratio = next(item for item in detail["candidates"] if item["label"] == "容积率")
    assert ratio["value"] == "0.291" and ratio["source_ref"] == "p1-t1-r8" and ratio["source_valid"]
    run = client.get(f"/api/projects/{project_id}/documents/{document_id}/extraction-runs?page=1").json()[0]
    assert ratio["extraction_origin"] == "TABLE" and ratio["extraction_run_id"] == run["id"]
    assert ratio["id"] in run["candidate_ids"]
    assert not any(item["label"] == "地下机动车位停车位" for item in detail["candidates"])
    from app.db import SessionLocal, SourceDocument
    with SessionLocal.begin() as session:
        stored = session.get(SourceDocument, document_id)
        stored.segments = stored.segments + [{"ref": "p1-t1-r0", "page": 1, "kind": "table_row", "text": "旧版误识别表头"}]
    repeated = client.post(f"/api/projects/{project_id}/documents/{document_id}/extract?page=1")
    assert repeated.status_code == 200 and repeated.json()["created"] == 0
    repeated_segments = client.get(f"/api/projects/{project_id}/documents/{document_id}?page=1").json()["segments"]
    assert len(repeated_segments) == len(detail["segments"])
    assert "p1-t1-r0" not in {part["ref"] for part in repeated_segments}

    primary = "p1-t1-r8"
    supplement = next(part["ref"] for part in detail["segments"] if part["ref"] in old_refs)
    page_two = client.get(f"/api/projects/{project_id}/documents/{document_id}?page=2").json()["segments"][0]["ref"]
    body = {"label": "容积率", "value": "0.291", "unit": "", "data_type": "decimal", "source_ref": primary,
            "support_refs": [supplement], "key": "floor_ratio"}
    bad = client.post(f"/api/projects/{project_id}/documents/{document_id}/candidates/{ratio['id']}/approve",
                      json={**body, "support_refs": [page_two]})
    assert bad.status_code == 400
    assert client.get(f"/api/projects/{project_id}/facts").json()["facts"] == []
    approved = client.post(f"/api/projects/{project_id}/documents/{document_id}/candidates/{ratio['id']}/approve", json=body)
    assert approved.status_code == 200, approved.text
    source_view = client.get(f"/api/projects/{project_id}/facts/floor_ratio/source").json()
    assert source_view["source_refs"] == [primary, supplement]
    assert "容积率" in source_view["excerpt"] and "0.291" in source_view["excerpt"]
    assert source_view["original_url"].endswith("#page=1")


def test_docx_text_and_table_are_locatable(client: TestClient, monkeypatch, tmp_path):
    import app.main as main_module
    from docx import Document

    monkeypatch.setattr(main_module, "STORAGE", tmp_path)
    project_id = create_project(client)
    doc = Document()
    doc.add_paragraph("项目概况")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "需求"
    table.cell(0, 1).text = "300000"
    stream = BytesIO()
    doc.save(stream)
    uploaded = client.post(f"/api/projects/{project_id}/documents", files={"file": ("input.docx", stream.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")})
    assert uploaded.status_code == 201, uploaded.text
    detail = client.get(f"/api/projects/{project_id}/documents/{uploaded.json()['id']}").json()
    assert any(item["ref"].startswith("d-p") for item in detail["segments"])
    assert any(item["ref"].startswith("d-t") and "300000" in item["text"] for item in detail["segments"])


def test_report_preview_plate_save_gate_and_export(client: TestClient, monkeypatch, tmp_path):
    import app.main as main_module

    monkeypatch.setattr(main_module, "STORAGE", tmp_path)
    project_id = create_project(client)
    add_fact(client, project_id, "demand")
    change, result = preview(client, project_id, "demand", "300000")
    commit(client, project_id, change, result)
    report = client.post(f"/api/projects/{project_id}/reports", json={"title": "项目报告"}).json()
    report_id = report["id"]
    content = [{"type": "h1", "children": [{"text": "项目报告"}]},
               {"type": "p", "children": [{"text": "首年需求为300000套。"}], "fact_keys": ["demand"]}]
    preview_response = client.post(f"/api/projects/{project_id}/reports/{report_id}/preview", json={"content": content, "base_version": 0})
    assert preview_response.status_code == 200, preview_response.text
    proposal = preview_response.json()
    assert proposal["fact_keys_affected"] == ["demand"]
    assert client.get(f"/api/projects/{project_id}/reports/{report_id}").json()["version"] == 0
    saved = client.put(f"/api/projects/{project_id}/reports/{report_id}", json={"content": content, "base_version": 0, "preview_token": proposal["preview_token"]})
    assert saved.status_code == 200, saved.text
    assert saved.json()["report"]["version"] == 1
    versions = client.get(f"/api/projects/{project_id}/reports/{report_id}/versions").json()
    assert [item["version"] for item in versions] == [1, 0]
    compared = client.get(f"/api/projects/{project_id}/reports/{report_id}/compare?base=0").json()
    assert compared["current_version"] == 1
    assert any(item["fact_keys"] == ["demand"] for item in compared["changes"])
    assert client.get(f"/api/projects/{project_id}/reports/{report_id}/export").status_code == 409
    assert client.post(f"/api/projects/{project_id}/reports/{report_id}/review").status_code == 400
    preview_export = client.get(f"/api/projects/{project_id}/reports/{report_id}/export?level=preview")
    assert preview_export.status_code == 200
    with ZipFile(BytesIO(preview_export.content)) as bundle:
        from docx import Document
        preview_word = Document(BytesIO(bundle.read("report.docx")))
        assert any("项目原件位置待核对；原录入来源：本项目资料" in paragraph.text
                   for paragraph in preview_word.paragraphs)
    bind_project_original(client, project_id, "demand", "300000")
    source_view = client.get(f"/api/projects/{project_id}/facts/demand/source").json()
    assert source_view["kind"] == "document" and source_view["review_status"] == "SOURCE_LOCATOR_REVIEWED"
    assert source_view["filename"] == "demand-300000.docx" and source_view["source_ref"] == "d-p1"
    assert "300000" in source_view["excerpt"] and source_view["original_source"] == "本项目资料"
    assert client.post(f"/api/projects/{project_id}/reports/{report_id}/review").status_code == 200
    exported = client.get(f"/api/projects/{project_id}/reports/{report_id}/export")
    assert exported.status_code == 200, exported.text
    with ZipFile(BytesIO(exported.content)) as archive:
        assert set(archive.namelist()) == {"report.docx", "report.pdf", "audit.json"}
        assert archive.read("report.pdf").startswith(b"%PDF")
        from docx import Document
        import pymupdf
        word = Document(BytesIO(archive.read("report.docx")))
        assert sum(paragraph.text == "项目报告" for paragraph in word.paragraphs) == 1
        assert any("demand：300000套。本项目原件：demand-300000.docx，段落 1" in paragraph.text for paragraph in word.paragraphs)
        assert not any("来源：本项目资料" in paragraph.text for paragraph in word.paragraphs)
        with pymupdf.open(stream=archive.read("report.pdf"), filetype="pdf") as pdf:
            pdf_text = "".join(page.get_text() for page in pdf)
        assert pdf_text.count("项目报告") == 1
        assert "demand：300000套。本项目原件：demand-300000.docx，段落 1" in pdf_text
        audit = __import__("json").loads(archive.read("audit.json"))
        assert audit["facts"][0]["value"] == "300000"
        assert audit["facts"][0]["source"] == "本项目资料"

    change, result = preview(client, project_id, "demand", "0")
    assert result["report_impacts"] == [{"report_id": report_id, "report_title": "项目报告", "report_version": 1,
                                          "position": 2, "text": "首年需求为300000套。", "fact_key": "demand",
                                          "before": result["changes"][0]["before"], "after": result["changes"][0]["after"]}]
    assert client.get(f"/api/projects/{project_id}/reports/{report_id}").json()["fact_impacts"] == []
    commit(client, project_id, change, result)
    after = client.get(f"/api/projects/{project_id}/reports/{report_id}").json()
    assert any(issue["code"] == "FACT_CHANGED" for issue in after["issues"])
    assert after["fact_impacts"][0]["position"] == 2
    assert client.get(f"/api/projects/{project_id}/reports/{report_id}/export").status_code == 409
    assert client.post(f"/api/projects/{project_id}/reports/{report_id}/review").status_code == 400
    revised = [{"type": "h1", "children": [{"text": "项目报告"}]},
               {"type": "p", "children": [{"text": "首年需求为0套。"}], "fact_keys": ["demand"]}]
    next_preview = client.post(f"/api/projects/{project_id}/reports/{report_id}/preview", json={"content": revised, "base_version": 1}).json()
    assert client.put(f"/api/projects/{project_id}/reports/{report_id}", json={"content": revised, "base_version": 1, "preview_token": next_preview["preview_token"]}).status_code == 200
    assert client.post(f"/api/projects/{project_id}/reports/{report_id}/review").status_code == 400
    bind_project_original(client, project_id, "demand", "0")
    assert client.post(f"/api/projects/{project_id}/reports/{report_id}/review").status_code == 200
    assert client.get(f"/api/projects/{project_id}/reports/{report_id}/export").status_code == 200
    assert client.put(f"/api/projects/{project_id}/reports/{report_id}", json={"content": content, "base_version": 0, "preview_token": proposal["preview_token"]}).status_code == 409


def test_report_save_rebinds_only_current_visible_fact_revision(client: TestClient, monkeypatch, tmp_path):
    import app.main as main_module
    from app.db import ReportDraft, SessionLocal

    monkeypatch.setattr(main_module, "STORAGE", tmp_path)
    project_id = create_project(client)
    add_fact(client, project_id, "test_qty")
    change, result = preview(client, project_id, "test_qty", "10")
    commit(client, project_id, change, result)
    report_id = client.post(f"/api/projects/{project_id}/reports", json={"title": "数量核对"}).json()["id"]

    def save(content, version):
        proposed = client.post(f"/api/projects/{project_id}/reports/{report_id}/preview", json={
            "content": content, "base_version": version})
        assert proposed.status_code == 200, proposed.text
        saved = client.put(f"/api/projects/{project_id}/reports/{report_id}", json={
            "content": content, "base_version": version, "preview_token": proposed.json()["preview_token"]})
        assert saved.status_code == 200, saved.text
        return saved.json()["report"]

    old = [{"type": "h1", "children": [{"text": "数量核对"}]},
           {"type": "p", "fact_keys": ["test_qty"], "children": [{"text": "本次测试数量为10套。"}]}]
    first = save(old, 0)
    assert first["bound_facts"]["test_qty"] == {"value": "10", "revision": 1}
    change, result = preview(client, project_id, "test_qty", "12")
    commit(client, project_id, change, result)
    unchanged = save(old, 1)
    assert unchanged["version"] == 1
    assert unchanged["bound_facts"]["test_qty"] == {"value": "10", "revision": 1}
    assert any(issue["code"] == "FACT_CHANGED" for issue in unchanged["issues"])
    still_old = [{**old[0]}, {**old[1], "children": [{"text": "本次测试数量为10套！"}]}]
    stale = save(still_old, 1)
    assert stale["version"] == 2
    assert stale["bound_facts"]["test_qty"] == {"value": "10", "revision": 1}
    assert client.post(f"/api/projects/{project_id}/reports/{report_id}/review").status_code == 400
    current = [{**old[0]}, {**old[1], "children": [{"text": "本次测试数量为12套。"}]}]
    fixed = save(current, 2)
    assert fixed["version"] == 3
    assert fixed["bound_facts"]["test_qty"] == {"value": "12", "revision": 2}
    assert fixed["fact_impacts"] == []
    assert not any(issue["code"] in {"FACT_CHANGED", "FACT_TEXT_STALE"} for issue in fixed["issues"])

    # A report saved by the earlier implementation may already have current
    # prose but an old bound snapshot. A valid no-text-change save repairs it.
    with SessionLocal.begin() as session:
        report = session.get(ReportDraft, report_id)
        report.bound_facts = {"test_qty": {"value": "10", "revision": 1}}
    repaired = save(current, 3)
    assert repaired["version"] == 4 and repaired["fact_impacts"] == []
    assert repaired["bound_facts"]["test_qty"] == {"value": "12", "revision": 2}
    assert client.post(f"/api/projects/{project_id}/reports/{report_id}/review").status_code == 400
    bind_project_original(client, project_id, "test_qty", "12")
    assert client.post(f"/api/projects/{project_id}/reports/{report_id}/review").status_code == 200


def test_review_checks_each_cited_paragraph_and_uncited_numbers(client: TestClient):
    project_id = create_project(client)
    add_fact(client, project_id, "capacity")
    change, result = preview(client, project_id, "capacity", "10")
    commit(client, project_id, change, result)
    report_id = client.post(f"/api/projects/{project_id}/reports", json={"title": "核对测试"}).json()["id"]
    content = [{"type": "h1", "children": [{"text": "核对测试"}]},
               {"type": "p", "children": [{"text": "本期能力为10套。"}], "fact_keys": ["capacity"]},
               {"type": "p", "children": [{"text": "历史能力为8套。"}], "fact_keys": ["capacity"]},
               {"type": "p", "children": [{"text": "尚有3套待确认。"}]}]
    proposed = client.post(f"/api/projects/{project_id}/reports/{report_id}/preview", json={"content": content, "base_version": 0}).json()
    saved = client.put(f"/api/projects/{project_id}/reports/{report_id}", json={"content": content, "base_version": 0, "preview_token": proposed["preview_token"]})
    assert saved.status_code == 200
    issues = saved.json()["report"]["issues"]
    assert any(issue["code"] == "FACT_TEXT_STALE" and issue["position"] == 3 for issue in issues)
    assert any(issue["code"] == "UNCITED_NUMBER" and issue["position"] == 4 for issue in issues)
    assert any(issue["code"] == "UNSUPPORTED_NUMBER" and issue["position"] == 3 for issue in issues)
    assert client.post(f"/api/projects/{project_id}/reports/{report_id}/review").status_code == 400
    assert client.get(f"/api/projects/{project_id}/reports/{report_id}/export").status_code == 409


def test_export_preserves_plate_marks_and_readable_document_citation():
    from docx import Document
    from app.report_pipeline import export_bundle, source_display

    source = "document:doc-id#p3-s2 · example.pdf"
    assert source_display(source) == "example.pdf，第 3 页，第 2 段"
    content = [{"type": "h1", "children": [{"text": "测试报告"}]},
               {"type": "p", "children": [{"text": "加粗", "bold": True}, {"text": "与斜体", "italic": True}], "fact_keys": ["amount"]},
               {"type": "blockquote", "children": [{"text": "引文说明"}]}]
    audit = {"facts": [{"key": "amount", "label": "金额", "value": "10", "unit": "万元", "source": source, "revision": 1}]}
    with ZipFile(BytesIO(export_bundle("测试报告", content, audit))) as archive:
        word = Document(BytesIO(archive.read("report.docx")))
        paragraph = next(paragraph for paragraph in word.paragraphs if paragraph.text == "加粗与斜体")
        assert paragraph.runs[0].bold is True
        assert paragraph.runs[1].italic is True
        quote = next(paragraph for paragraph in word.paragraphs if paragraph.text == "引文说明")
        assert quote.paragraph_format.left_indent is not None
        assert any("example.pdf，第 3 页，第 2 段" in paragraph.text for paragraph in word.paragraphs)


def test_model_draft_uses_only_project_facts(client: TestClient, monkeypatch):
    import app.main as main_module

    project_id = create_project(client)
    add_fact(client, project_id, "capacity")
    change, result = preview(client, project_id, "capacity", "254016")
    commit(client, project_id, change, result)
    report_id = client.post(f"/api/projects/{project_id}/reports", json={"title": "能力说明"}).json()["id"]
    monkeypatch.setattr(main_module, "model_section", lambda title, facts: [{"type": "p", "children": [{"text": "合格能力为254016套。"}], "fact_keys": ["capacity"], "origin": "model"}])
    path = f"/api/projects/{project_id}/reports/{report_id}/generate"
    candidate = client.post(f"{path}/preview", json={"title": "生产能力", "fact_keys": ["capacity"]})
    assert candidate.status_code == 200, candidate.text
    proposed = candidate.json()
    assert proposed["paragraphs"][0]["fact_keys"] == ["capacity"]
    assert proposed["base_version"] == 0
    # 关闭候选相当于不请求 commit：正文和版本保持原样。
    assert client.get(f"/api/projects/{project_id}/reports/{report_id}").json()["version"] == 0
    assert len(client.get(f"/api/projects/{project_id}/reports/{report_id}/versions").json()) == 1
    tampered = client.post(f"{path}/commit", json={**proposed, "paragraphs": [
        {"type": "p", "children": [{"text": "合格能力为999999套。"}], "fact_keys": ["capacity"], "origin": "model"}]})
    assert tampered.status_code == 409
    assert client.get(f"/api/projects/{project_id}/reports/{report_id}").json()["version"] == 0
    generated = client.post(f"{path}/commit", json=proposed)
    assert generated.status_code == 200, generated.text
    assert generated.json()["version"] == 1
    assert generated.json()["content"][-1]["fact_keys"] == ["capacity"]
    assert any(issue["code"] == "UNREVIEWED" for issue in generated.json()["issues"])
    assert client.get(f"/api/projects/{project_id}/reports/{report_id}").json() == generated.json()
    assert client.post(f"{path}/commit", json=proposed).status_code == 409
    assert client.post(f"{path}/preview", json={"title": "未知", "fact_keys": ["not_in_project"]}).status_code == 400
    stale = client.post(f"{path}/preview", json={"title": "第二节", "fact_keys": ["capacity"]}).json()
    change, result = preview(client, project_id, "capacity", "254017")
    commit(client, project_id, change, result)
    assert client.post(f"{path}/commit", json=stale).status_code == 409
    assert client.get(f"/api/projects/{project_id}/reports/{report_id}").json()["version"] == 1
    assert client.post(path, json={"title": "旧接口", "fact_keys": ["capacity"]}).status_code == 404


def test_all_categories_and_graph_are_grounded(client: TestClient):
    corpus = CorpusRepository()
    builtin = next(item for item in client.get("/api/projects").json() if item["is_builtin"])
    project_id = builtin["id"]
    assert builtin["has_corpus"] is True
    assert client.get("/api/corpus/summary").status_code == 404
    summary = client.get(f"/api/projects/{project_id}/corpus/summary").json()
    assert summary["project_id"] == project_id
    assert summary["integrity"]["valid"]
    assert summary["integrity"]["total_files"] == 117
    assert sum(group["count"] for group in summary["groups"]) == 30
    categories = client.get(f"/api/projects/{project_id}/corpus/categories").json()
    assert len(categories) == 30
    assert sum(category["count"] for category in categories) == 114
    for number in range(1, 31):
        result = client.get(f"/api/projects/{project_id}/corpus/categories/{number}")
        assert result.status_code == 200, (number, result.text)
        assert result.json()["records"], number
    assert corpus.parsed("chunks/embeddings.parquet")["rows"] == 0
    entities, edges, _ = corpus.graph()
    assert len(corpus.parsed("graph/relations.yaml")) == 1666
    assert len(edges) > 1666
    assert all(edge["source"] in entities and edge["target"] in entities for edge in edges)
    edge_pairs = {(edge["source"], edge["target"], edge["type"]) for edge in edges}
    assert ("N013", "R001", "RULE_INPUT") in edge_pairs
    assert ("R001", "N017", "RULE_OUTPUT") in edge_pairs
    assert ("N017", "L002", "SUPPORTS") in edge_pairs
    assert ("EV-01", "L002", "MUST_CITE") in edge_pairs
    assert corpus.parsed("evidence/evidence.yaml")[1]["excerpts"][0]["path"] == "evidence/excerpts/EX-01-01.md"
    conflict = corpus.category(23)["records"][0]
    assert {(str(node["value"]), node["unit"]) for node in conflict["linked_nodes"]} == {("800", "kW"), ("650", "kW")}
    assert corpus.category(16)["records"][0]["check"] is not None
    assert client.get(f"/api/projects/{project_id}/corpus/graph?focus=X001&depth=2").status_code == 200
    auxiliary = client.get(f"/api/projects/{project_id}/corpus/raw?path=validate_package.py")
    assert auxiliary.status_code == 200 and auxiliary.content.startswith(b'#!/')


def test_builtin_corpus_is_project_scoped_and_read_only(client: TestClient):
    builtin = next(item for item in client.get("/api/projects").json() if item["is_builtin"])
    project_id = builtin["id"]
    blank_id = create_project(client)
    assert client.get(f"/api/projects/{blank_id}/corpus/summary").status_code == 404
    assert client.get(f"/api/projects/{blank_id}/corpus/categories/12").status_code == 404
    assert client.get(f"/api/projects/{blank_id}/corpus/raw?path=manifest.yaml").status_code == 404
    assert client.get(f"/api/projects/{project_id}/corpus/categories/12").json()["project_id"] == project_id
    assert client.post(f"/api/projects/{project_id}/facts", json={
        "key": "new_value", "label": "新值", "data_type": "integer", "unit": "套",
    }).status_code == 403
    assert client.post(f"/api/projects/{project_id}/changes/preview", json={
        "fact_key": "N013", "value": "0", "source": "测试",
    }).status_code == 403
    assert client.get(f"/api/projects/{project_id}/facts").json()["facts"] == []


def writing_ready_project(client: TestClient) -> str:
    project_id = create_project(client)
    selected = client.put(f"/api/projects/{project_id}/writing/reference", json={
        "corpus_id": "cy_tray_20260918", "corpus_version": "v2",
        "target_report_type": "feasibility", "accepted": True,
    })
    assert selected.status_code == 200, selected.text
    result = client.post(f"/api/projects/{project_id}/writing/setup")
    assert result.status_code == 200, result.text
    assert result.json()["historical_values_copied"] == 0
    assert result.json()["bindings_created"] == 0
    facts = {item["key"]: item for item in client.get(f"/api/projects/{project_id}/facts").json()["facts"]}
    assert len(facts) == 5
    assert all(item["value"] is None for item in facts.values())
    assert facts["planned_sales"]["status"] == "UNEVALUABLE"
    bindings = {"N017": "first_year_demand", "N034": "qualified_capacity", "N035": "planned_sales",
                "supplier_name": "supplier_name", "N080": "equipment_unit_price"}
    saved = client.put(f"/api/projects/{project_id}/writing/bindings", json={"bindings": bindings})
    assert saved.status_code == 200, saved.text
    return project_id


def writing_commit(client: TestClient, project_id: str, key: str, value: str | None):
    change = {"fact_key": key, "value": value, "source": "本项目人工输入", "reason": "写作验证"}
    proposal = client.post(f"/api/projects/{project_id}/writing/preview", json=change)
    assert proposal.status_code == 200, proposal.text
    previewed = proposal.json()
    committed = client.post(f"/api/projects/{project_id}/changes/commit", json={
        **change, "base_version": previewed["base_version"], "preview_token": previewed["preview_token"],
    })
    assert committed.status_code == 200, committed.text
    return previewed


def test_writing_mapping_source_and_missing_values(client: TestClient):
    project_id = writing_ready_project(client)
    data = client.get(f"/api/projects/{project_id}/writing").json()
    assert data["project_id"] == project_id
    assert data["source_project_id"] != project_id
    assert [item["chunk_id"] for item in data["sources"]] == ["C0052", "C0083"]
    assert {item["mention_id"] for item in data["sources"][1]["mentions"]} == {"M0240", "M0241"}
    assert data["draft"]["blocked"] is True
    candidate = " ".join(block["text"] for block in data["draft"]["blocks"])
    assert "180,000" not in candidate and "1,480" not in candidate and "禾进装备" not in candidate
    assert client.post(f"/api/projects/{project_id}/writing/setup").status_code == 409
    before = client.get(f"/api/projects/{project_id}/writing").json()["bindings"]
    rejected = client.put(f"/api/projects/{project_id}/writing/bindings", json={
        "bindings": {"N017": "equipment_unit_price"},
    })
    assert rejected.status_code == 400
    assert client.get(f"/api/projects/{project_id}/writing").json()["bindings"] == before
    builtin_id = next(item["id"] for item in client.get("/api/projects").json() if item["is_builtin"])
    assert client.get(f"/api/projects/{builtin_id}/writing").status_code == 403


def test_writing_preview_condition_format_and_persistence(client: TestClient):
    project_id = writing_ready_project(client)
    first_preview = writing_commit(client, project_id, "first_year_demand", "300000")
    assert first_preview["affected"] == [{
        "chunk_id": "C0052",
        "before": first_preview["before"]["blocks"][0]["text"],
        "after": first_preview["after"]["blocks"][0]["text"],
        "issues_changed": True,
    }]
    before = client.get(f"/api/projects/{project_id}/facts").json()
    change = {"fact_key": "qualified_capacity", "value": "254016", "source": "本项目人工输入", "reason": "写作验证"}
    result = client.post(f"/api/projects/{project_id}/writing/preview", json=change)
    assert result.status_code == 200, result.text
    proposed = result.json()
    after = proposed["after"]
    paragraph = next(item for item in after["blocks"] if item["chunk_id"] == "C0052")
    assert "300,000套" in paragraph["text"]
    assert "254,016套" in paragraph["text"]
    assert "低于" not in paragraph["text"]
    assert {item["fact_key"] for item in paragraph["references"]} == {"first_year_demand", "qualified_capacity", "planned_sales"}
    assert any(item["code"] == "HISTORICAL_CONDITION_FALSE" for item in after["issues"])
    assert any(item["chunk_id"] == "C0052" for item in proposed["affected"])
    assert client.get(f"/api/projects/{project_id}/facts").json() == before  # 取消预览即不提交
    committed = client.post(f"/api/projects/{project_id}/changes/commit", json={
        **change, "base_version": proposed["base_version"], "preview_token": proposed["preview_token"],
    })
    assert committed.status_code == 200, committed.text
    reloaded = client.get(f"/api/projects/{project_id}/writing").json()
    assert reloaded["draft"]["blocks"][0]["text"] == paragraph["text"]
    assert next(item for item in reloaded["facts"] if item["key"] == "planned_sales")["value"] == "254016"

    writing_commit(client, project_id, "supplier_name", "新拓设备")
    price_preview = writing_commit(client, project_id, "equipment_unit_price", "1234.50")
    quote = next(item for item in price_preview["after"]["blocks"] if item["chunk_id"] == "C0083")
    assert "新拓设备" in quote["text"] and "禾进装备" not in quote["text"]
    assert "12,345,000元/条" in quote["text"] and "1,234.50万元/条" in quote["text"]
    assert {item["mention_id"] for item in quote["references"] if item["mention_id"]} == {"M0240", "M0241"}
    assert any(item["code"] == "QUOTE_SCOPE_UNVERIFIED" for item in price_preview["after"]["issues"])
    assert client.get(f"/api/projects/{project_id}/writing").json()["draft"]["blocks"][1]["text"] == quote["text"]


def test_writing_zero_and_format_loss_are_explicit(client: TestClient):
    project_id = writing_ready_project(client)
    writing_commit(client, project_id, "first_year_demand", "0")
    proposed = writing_commit(client, project_id, "qualified_capacity", "254016")
    assert "首年客户需求0套" in proposed["after"]["blocks"][0]["text"]
    assert "计划销售量为0套" in proposed["after"]["blocks"][0]["text"]
    writing_commit(client, project_id, "supplier_name", "新拓设备")
    change = {"fact_key": "equipment_unit_price", "value": "1234.567", "source": "本项目人工输入", "reason": "写作验证"}
    response = client.post(f"/api/projects/{project_id}/writing/preview", json=change)
    assert response.status_code == 200
    assert any(item["code"] == "FORMAT_LOSS" for item in response.json()["after"]["issues"])
    assert "1,234.57" not in response.json()["after"]["blocks"][1]["text"]
