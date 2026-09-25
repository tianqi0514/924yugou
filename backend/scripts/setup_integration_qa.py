"""Validate the isolated integration environment and write a synthetic DOCX fixture.

This script never clears tables or creates a project. The browser test creates a
unique QA project through the same UI as a user. Run it only with the explicit
test database and an independent storage directory.
"""

from __future__ import annotations

import argparse
import json
import os
from pathlib import Path
from urllib.parse import urlsplit


ROOT = Path(__file__).resolve().parents[2]
TEST_ROOT = (ROOT / "tmp" / "integration").resolve()


def checked_environment(fixture: Path) -> str:
    database_url = os.environ.get("DATABASE_URL", "")
    if urlsplit(database_url).path != "/report_platform_test":
        raise SystemExit("集成测试只允许连接 report_platform_test")
    storage_value = os.environ.get("DOCUMENT_STORAGE_DIR", "")
    if not storage_value:
        raise SystemExit("必须设置独立的 DOCUMENT_STORAGE_DIR")
    storage = Path(storage_value).resolve()
    model_settings = Path(os.environ.get("MODEL_SETTINGS_DIR", "")).resolve()
    try:
        storage.relative_to(TEST_ROOT)
        model_settings.relative_to(TEST_ROOT)
        fixture.resolve().relative_to(TEST_ROOT)
    except ValueError as exc:
        raise SystemExit("集成测试的文件必须位于 tmp/integration") from exc
    if storage == TEST_ROOT or model_settings == TEST_ROOT:
        raise SystemExit("文件存储目录必须单独设置")

    import psycopg

    # psycopg accepts PostgreSQL URLs, while SQLAlchemy's driver suffix must
    # be removed for this independent database-name check.
    with psycopg.connect(database_url.replace("postgresql+psycopg://", "postgresql://", 1)) as connection:
        with connection.cursor() as cursor:
            cursor.execute("SELECT current_database()")
            actual = cursor.fetchone()[0]
    if actual != "report_platform_test":
        raise SystemExit("数据库连接目标不是 report_platform_test")
    return actual


def write_fixture(destination: Path) -> None:
    from docx import Document

    destination.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    doc.add_heading("QA 合成原件 · 非真实业务资料", level=1)
    doc.add_paragraph("本项目首年客户预测需求为 300000 套，尚非已签订单。")
    doc.add_paragraph("本项目正常年合格能力为 254016 套。")
    doc.add_paragraph("压力测试情景下，本项目首年客户预测需求为 0 套；不能解释为已签订单。")
    doc.add_paragraph("本项目设备供应商为新拓设备。")
    doc.save(destination)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--fixture", type=Path, required=True)
    args = parser.parse_args()
    database_name = checked_environment(args.fixture)
    write_fixture(args.fixture)
    print(json.dumps({"database": database_name, "fixture": str(args.fixture)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
