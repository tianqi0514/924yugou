"""Inspect a saved scenario ZIP from a real browser/API writing journey."""

from __future__ import annotations

import argparse
import io
import json
import re
from zipfile import ZipFile

import pymupdf
from docx import Document


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("archive")
    parser.add_argument("--project", required=True)
    parser.add_argument("--report", required=True)
    parser.add_argument("--run", required=True)
    parser.add_argument("--text", action="append", default=[])
    parser.add_argument("--result", action="append", default=[])
    args = parser.parse_args()

    with ZipFile(args.archive) as archive:
        assert set(archive.namelist()) == {"report.docx", "report.pdf", "audit.json"}
        word = Document(io.BytesIO(archive.read("report.docx")))
        docx_text = "\n".join([
            *(paragraph.text for paragraph in word.paragraphs),
            *(cell.text for table in word.tables for row in table.rows for cell in row.cells),
        ])
        with pymupdf.open(stream=archive.read("report.pdf"), filetype="pdf") as pdf:
            pdf_text = "\n".join(page.get_text() for page in pdf)
            pdf_pages = len(pdf)
        audit = json.loads(archive.read("audit.json"))

    assert audit["project_id"] == args.project
    assert audit["report_id"] == args.report
    assert audit["analysis_run_id"] == args.run
    assert audit["delivery_status"] == "scenario_assumptions_reviewed"
    assert audit["analysis_refs"], "正文缺少运行引用"
    assert audit["content_sha256"]
    assert audit["export_watermark_sha256"]
    for expected in args.text:
        normalized = re.sub(r"\s+", "", expected)
        assert normalized in re.sub(r"\s+", "", docx_text), f"DOCX 缺少：{expected}"
        assert normalized in re.sub(r"\s+", "", pdf_text), f"PDF 缺少：{expected}"
    for entry in args.result:
        key, value = entry.split("=", 1)
        assert audit["analysis_run"]["results"][key]["value"] == value, entry
    print(json.dumps({"report_version": audit["report_version"], "run_id": args.run,
                      "docx_tables": len(word.tables), "pdf_pages": pdf_pages,
                      "checked_text": args.text, "checked_results": args.result}, ensure_ascii=False))


if __name__ == "__main__":
    main()
