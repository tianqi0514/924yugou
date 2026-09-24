"""Read-only audit of the 117 stored corpus files and their per-file purposes.

Run from ``backend`` with the project's virtual environment.  DATABASE_URL may
point at either the application database or a separate test database: this
script starts a read-only transaction and does not import or migrate data.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import sys
from pathlib import Path
from urllib.parse import quote
from urllib.request import urlopen

import yaml
from openpyxl import load_workbook
from sqlalchemy import text
from sqlalchemy.exc import SQLAlchemyError

from app.db import engine


CORPUS_ROOT = Path(__file__).resolve().parents[2] / "demo-data" / "澄岳精密" / "解析产物"
EXPLANATION = CORPUS_ROOT.parent / "澄岳精密_30类产物文件位置与意义说明.xlsx"


def expected_files() -> dict[str, tuple[int | None, str, str]]:
    """Read the 117 authoritative per-file descriptions from the supplied sheet."""
    workbook = load_workbook(EXPLANATION, read_only=True, data_only=True)
    try:
        sheet = workbook["逐文件明细"]
        expected: dict[str, tuple[int | None, str, str]] = {}
        for row in sheet.iter_rows(min_row=7, values_only=True):
            number, path, purpose, boundary = row[1], row[4], row[8], row[9]
            if not path:
                continue  # The sheet ends with a blank row and a grand total.
            if path in expected:
                raise AssertionError(f"说明表重复路径: {path}")
            if not isinstance(number, int) or not purpose or not boundary:
                raise AssertionError(f"说明表缺少类别、用途或边界: {path}")
            expected[path] = (None if number == 0 else number, str(purpose).strip(), str(boundary).strip())
    finally:
        workbook.close()
    return expected


def fetch_json(url: str) -> object:
    with urlopen(url, timeout=15) as response:
        return json.load(response)


def check_api(base_url: str, project_id: str, expected: dict, actual: dict, problems: list[str]) -> None:
    """Optionally confirm the running API exposes every stored file and meaning."""
    prefix = f"{base_url.rstrip('/')}/api/projects/{quote(project_id, safe='')}/corpus"
    try:
        categories = fetch_json(f"{prefix}/categories")
        summary = fetch_json(f"{prefix}/summary")
    except Exception as exc:
        problems.append(f"无法读取语料目录 API: {exc}")
        return
    catalog: dict[str, dict] = {}
    for category in categories:
        for artifact in category.get("artifacts", []):
            catalog[artifact["artifact_id"]] = artifact
    for artifact in summary.get("auxiliary_assets", []):
        catalog[artifact["artifact_id"]] = artifact
    stored_ids = {row["artifact_id"] for row in actual.values()}
    if set(catalog) != stored_ids:
        problems.append(f"API 目录资产 ID 集合不符: 缺少={len(stored_ids - set(catalog))}，多出={len(set(catalog) - stored_ids)}")

    for path, row in actual.items():
        if path not in expected:
            continue
        category_number, purpose, boundary = expected[path]
        artifact_id = row["artifact_id"]
        card = catalog.get(artifact_id)
        if card is not None:
            if str(card.get("purpose") or "").strip() != purpose or str(card.get("boundary") or "").strip() != boundary:
                problems.append(f"API 目录用途或边界不符: {path}")
        try:
            detail = fetch_json(f"{prefix}/artifacts/{quote(artifact_id, safe='')}")
        except Exception as exc:
            problems.append(f"API 文件详情不可读: {path}: {exc}")
            continue
        if (detail.get("artifact_id") != artifact_id or detail.get("path") != path
                or detail.get("category_number") != category_number
                or str(detail.get("purpose") or "").strip() != purpose
                or str(detail.get("boundary") or "").strip() != boundary):
            problems.append(f"API 文件详情身份、用途或边界不符: {path}")


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--base-url", help="可选：运行中 API 的根地址，例如 http://127.0.0.1:8000")
    parser.add_argument("--project-id", help="可选：绑定了澄岳语料的项目 ID，须与 --base-url 同时填写")
    args = parser.parse_args()
    if bool(args.base_url) != bool(args.project_id):
        parser.error("--base-url 与 --project-id 必须同时填写")

    manifest = yaml.safe_load((CORPUS_ROOT / "manifest.yaml").read_text(encoding="utf-8"))
    corpus_id = manifest["corpus_id"]
    manifest_paths = {
        member["path"]
        for category in manifest["files"]
        for member in (category.get("files") or [category])
    }
    expected = expected_files()
    problems: list[str] = []
    if len(expected) != 117 or len(manifest_paths) != 114:
        problems.append(f"源清单数量异常: Excel={len(expected)}，manifest={len(manifest_paths)}")
    if {path for path, (number, _, _) in expected.items() if number is not None} != manifest_paths:
        problems.append("Excel 的 114 个产物路径与 manifest 不一致")

    try:
        with engine.connect() as connection:
            with connection.begin():
                connection.execute(text("SET TRANSACTION READ ONLY"))
                rows = connection.execute(
                    text("SELECT artifact_id, path, category_number, sha256, content, purpose, boundary, purpose_source, writing_use "
                         "FROM corpus_artifacts WHERE corpus_id = :corpus_id"),
                    {"corpus_id": corpus_id},
                ).mappings().all()
                source = connection.execute(
                    text("SELECT sha256, content FROM corpus_source_assets "
                         "WHERE corpus_id = :corpus_id AND name = :name"),
                    {"corpus_id": corpus_id, "name": EXPLANATION.name},
                ).mappings().first()
    except SQLAlchemyError as exc:
        print(f"FAIL: 语料数据库不可读或尚未迁移逐文件说明字段（{type(exc.orig if hasattr(exc, 'orig') else exc).__name__}）", file=sys.stderr)
        return 1

    actual = {row["path"]: row for row in rows}
    if len(rows) != len(actual):
        problems.append("数据库存在重复文件路径")
    if set(actual) != set(expected):
        problems.append(f"数据库文件集合不符: 缺少={sorted(set(expected)-set(actual))}，多出={sorted(set(actual)-set(expected))}")
    if source is None or hashlib.sha256(bytes(source["content"])).hexdigest() != source["sha256"] \
            or bytes(source["content"]) != EXPLANATION.read_bytes():
        problems.append("入库的文件意义说明表与源表不一致")

    checksum = actual.get("SHA256SUMS")
    checksums: dict[str, str] = {}
    if checksum is None:
        problems.append("数据库缺少 SHA256SUMS")
    else:
        for line in bytes(checksum["content"]).decode("utf-8").splitlines():
            digest, path = line.split("  ", 1)
            if path in checksums:
                problems.append(f"SHA256SUMS 重复路径: {path}")
            checksums[path] = digest
        if set(checksums) != set(actual) - {"SHA256SUMS"} or len(checksums) != 116:
            problems.append("SHA256SUMS 未完整覆盖其余 116 个入库文件")

    for path, row in actual.items():
        if path not in expected:
            continue
        number, purpose, boundary = expected[path]
        if row["category_number"] != number:
            problems.append(f"类别错误: {path}")
        if str(row["purpose"] or "").strip() != purpose:
            problems.append(f"用途与说明表不一致: {path}")
        if str(row["boundary"] or "").strip() != boundary:
            problems.append(f"边界与说明表不一致: {path}")
        if not str(row["purpose_source"] or "").strip():
            problems.append(f"用途缺少来源标记: {path}")
        if number is not None and not str(row["writing_use"] or "").strip():
            problems.append(f"用途缺少写作阶段说明: {path}")
        digest = hashlib.sha256(bytes(row["content"])).hexdigest()
        if digest != row["sha256"] or (path in checksums and digest != checksums[path]):
            problems.append(f"文件内容摘要不一致: {path}")
        if digest != hashlib.sha256((CORPUS_ROOT / path).read_bytes()).hexdigest():
            problems.append(f"入库内容与语料复制件不一致: {path}")

    if args.base_url and args.project_id:
        check_api(args.base_url, args.project_id, expected, actual, problems)

    if problems:
        print(f"FAIL: {len(problems)} 项不一致（显示前 30 项）", file=sys.stderr)
        for problem in problems[:30]:
            print(f"- {problem}", file=sys.stderr)
        return 1
    print(f"PASS: corpus={corpus_id}; 30 类; 114 个产物 + 3 个辅助文件; "
          "117 条逐文件用途/边界与说明表一致; 116 条 SHA256SUMS 和 117 个入库内容摘要一致")
    if args.base_url:
        print("PASS: API 目录和 117 个文件详情的身份、用途、边界与数据库/说明表一致")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
