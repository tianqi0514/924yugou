"""Inspect a real report export produced by the browser integration test."""

from __future__ import annotations

import argparse
import io
import json
import re
from zipfile import ZipFile

import pymupdf
from docx import Document


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("archive")
    parser.add_argument("--project", required=True)
    parser.add_argument("--report", required=True)
    parser.add_argument("--level", choices=("formal", "preview"), required=True)
    parser.add_argument("--text", action="append", default=[])
    parser.add_argument("--fact", action="append", default=[])
    args = parser.parse_args()

    with ZipFile(args.archive) as archive:
        assert set(archive.namelist()) == {"report.docx", "report.pdf", "audit.json"}
        word = Document(io.BytesIO(archive.read("report.docx")))
        docx_text = "\n".join([
            *(paragraph.text for paragraph in word.paragraphs),
            *(cell.text for table in word.tables for row in table.rows for cell in row.cells),
        ])
        with pymupdf.open(stream=archive.read("report.pdf"), filetype="pdf") as pdf:
            pdf_text = "\n".join(page.get_text() for page in pdf)
            pdf_pages = len(pdf)
        audit = json.loads(archive.read("audit.json"))

    assert audit["project_id"] == args.project
    assert audit["report_id"] == args.report
    assert audit["delivery_status"] == ("source_locator_reviewed" if args.level == "formal" else "preview_only")
    normalized_docx = re.sub(r"\s+", "", docx_text)
    normalized_pdf = re.sub(r"\s+", "", pdf_text)
    for text in args.text:
        normalized_expected = re.sub(r"\s+", "", text)
        assert normalized_expected in normalized_docx, f"DOCX 缺少正文：{text}"
        assert normalized_expected in normalized_pdf, f"PDF 缺少正文：{text}"
    assert "300000" not in normalized_docx and "300,000" not in normalized_docx, "DOCX 残留旧需求值"
    assert "300000" not in normalized_pdf and "300,000" not in normalized_pdf, "PDF 残留旧需求值"
    facts = {fact["key"]: fact for fact in audit["facts"]}
    evidence = {row["fact_key"]: row for row in audit["evidence_bindings"]}
    for pair in args.fact:
        key, expected = pair.split("=", 1)
        assert facts[key]["value"] == expected, f"审计事实值不一致：{key}"
        assert evidence[key]["fact_revision"] == facts[key]["revision"]
        if args.level == "formal":
            assert evidence[key]["status"] in {"source_locator_reviewed", "derived_from_reviewed_inputs"}
    assert audit["source_refs"], "审计缺少章节语料来源"
    assert audit["content_sha256"]
    rule_refs = [ref for block in audit["source_refs"] for ref in block.get("project_rule_refs", [])]
    assert rule_refs, "审计缺少本项目计算规则"
    for ref in rule_refs:
        assert facts[ref["target_key"]]["revision"] == ref["target_fact_revision"]
        for item in ref["input_fact_revisions"]:
            assert facts[item["fact_key"]]["revision"] == item["revision"]
            assert facts[item["fact_key"]]["value"] == item["value"]
    print(json.dumps({
        "delivery_status": audit["delivery_status"],
        "report_version": audit["report_version"],
        "docx_tables": len(word.tables),
        "pdf_pages": pdf_pages,
        "facts_checked": sorted(pair.split("=", 1)[0] for pair in args.fact),
        "source_blocks": len(audit["source_refs"]),
        "rule_refs": len(rule_refs),
    }, ensure_ascii=False))


if __name__ == "__main__":
    main()
